import json
import sys
import time
from pathlib import Path
from uuid import uuid4

import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = PROJECT_ROOT / "backend"
sys.path.insert(0, str(BACKEND_ROOT))

from app.services.chat_database import (  # noqa: E402
    chat_database,
)
from app.services.orchestrator import SessionOrchestrator  # noqa: E402
from app.services.plugin_registry import PluginManifest, plugin_registry  # noqa: E402


def test_plugin_lifecycle_permissions_and_builtin_actions(client):
    session_id = str(uuid4())
    plugins = client.get("/plugins/")
    assert plugins.status_code == 200
    by_id = {plugin["id"]: plugin for plugin in plugins.json()}
    assert {
        "desktop-control",
        "comfyui",
        "browser-links",
        "windows-inspector",
        "documents",
        "pdf",
        "spreadsheets",
        "presentations",
        "github-readonly",
        "visualize",
        "qa-diagnostics",
    } <= set(by_id)
    assert all(plugin["enabled"] for plugin in by_id.values())

    health = client.post("/plugins/qa-diagnostics/invoke", json={
        "action": "workspace_health",
        "session_id": session_id,
        "arguments": {},
    })
    assert health.status_code == 200
    assert health.json()["result"]["healthy"] is True
    assert health.json()["result"]["artifacts"] >= 3

    blocked_scheme = client.post("/plugins/browser-links/invoke", json={
        "action": "open_url",
        "session_id": session_id,
        "arguments": {"url": "file:///C:/Windows/System32"},
    })
    assert blocked_scheme.status_code == 400

    safe_link = client.post("/plugins/browser-links/invoke", json={
        "action": "open_url",
        "session_id": session_id,
        "arguments": {"url": "https://example.com/docs"},
    })
    assert safe_link.status_code == 200
    assert safe_link.json()["result"] == {
        "kind": "open_external",
        "url": "https://example.com/docs",
        "browser": "default",
    }

    invalid_browser = client.post("/plugins/browser-links/invoke", json={
        "action": "open_url",
        "session_id": session_id,
        "arguments": {"url": "https://example.com", "browser": "untrusted"},
    })
    assert invalid_browser.status_code == 400

    disabled = client.patch("/plugins/visualize", json={"enabled": False})
    assert disabled.status_code == 200
    assert disabled.json()["enabled"] is False
    disabled_run = client.post("/plugins/visualize/invoke", json={
        "action": "task_chart",
        "session_id": session_id,
        "arguments": {},
    })
    assert disabled_run.status_code == 403
    assert client.patch("/plugins/visualize", json={"enabled": True}).status_code == 200

    custom_manifest = {
        "id": "local-note-helper",
        "name": "本地备注助手",
        "version": "1.0.0",
        "description": "测试声明式插件",
        "author": "test",
        "category": "Productivity",
        "permissions": [],
        "actions": [{
            "id": "echo_note",
            "label": "回显备注",
            "description": "不执行代码",
            "handler": "template",
            "input_schema": {
                "type": "object",
                "required": ["text"],
                "properties": {"text": {"type": "string"}},
            },
            "output_template": "备注：${text}",
        }],
    }
    installed = client.post("/plugins/", json=custom_manifest)
    assert installed.status_code == 200
    assert installed.json()["builtin"] is False
    custom_run = client.post("/plugins/local-note-helper/invoke", json={
        "action": "echo_note",
        "session_id": session_id,
        "arguments": {"text": "安全模板"},
    })
    assert custom_run.status_code == 200
    assert custom_run.json()["result"]["text"] == "备注：安全模板"

    tool_names = {tool.name for tool in SessionOrchestrator(session_id)._available_tools()}
    assert "plugin__qa_diagnostics__workspace_health" in tool_names
    assert "plugin__local_note_helper__echo_note" in tool_names
    assert "plugin__desktop_control__capabilities" in tool_names
    assert "plugin__desktop_control__get_state" in tool_names
    assert "plugin__desktop_control__click" in tool_names
    assert "plugin__desktop_control__type_text" in tool_names
    assert "plugin__desktop_control__read_document" in tool_names
    assert "plugin__desktop_control__write_document" in tool_names
    assert "plugin__documents__export_to_local" in tool_names
    assert "plugin__comfyui__status" in tool_names
    assert "plugin__comfyui__generate_image" in tool_names
    assert "plugin__comfyui__generate_video" in tool_names
    assert "plugin__comfyui__run_workflow" in tool_names

    example_path = PROJECT_ROOT / "plugins" / "local-note-helper" / ".aether-plugin" / "plugin.json"
    PluginManifest.model_validate_json(example_path.read_text(encoding="utf-8"))

    unsafe_manifest = {
        **custom_manifest,
        "id": "unsafe-runner",
        "actions": [{**custom_manifest["actions"][0], "handler": "run_command"}],
    }
    assert client.post("/plugins/", json=unsafe_manifest).status_code == 400
    assert client.delete("/plugins/browser-links").status_code == 403
    assert client.delete("/plugins/local-note-helper").status_code == 200


@pytest.mark.asyncio
async def test_workspace_document_can_be_read_and_exported_to_local_path(tmp_path, monkeypatch):
    session_id = str(uuid4())
    await plugin_registry.artifacts.write(
        session_id,
        "articles/example.md",
        "workspace document content",
        "test document bridge",
        "test",
    )

    read_result = await plugin_registry.invoke_tool(
        "plugin__desktop_control__read_document",
        session_id,
        "model:test",
        {"path": "articles/example.md"},
    )
    assert read_result["result"]["source"] == "session_workspace"
    assert read_result["result"]["data"]["text"] == "workspace document content"

    monkeypatch.setenv("AETHERCHAT_TEST_AUTO_APPROVE_COMPUTER_USE", "1")
    destination = tmp_path / "exported.md"
    export_result = await plugin_registry.invoke_tool(
        "plugin__documents__export_to_local",
        session_id,
        "model:test",
        {"path": "articles/example.md", "destination_path": str(destination)},
    )
    assert destination.read_text(encoding="utf-8") == "workspace document content"
    assert export_result["result"]["destination_path"] == str(destination.resolve())


def test_comfyui_graphs_cover_local_image_and_video_pipelines():
    from app.services.comfyui_service import ComfyUIService

    service = ComfyUIService()
    checkpoint = service._checkpoint_image_graph(
        "test image", "bad image", "realvisXL_V50.safetensors",
        768, 768, 6, 2.0, 123, "zhlanchangfu_v2.safetensors", 0.8,
    )
    assert checkpoint["1"]["inputs"]["ckpt_name"] == "realvisXL_V50.safetensors"
    assert checkpoint["2"]["class_type"] == "LoraLoader"
    assert checkpoint["8"]["class_type"] == "SaveImage"

    qwen = service._qwen_image_graph(
        "中文招牌", "模糊", "qwen_image_2512_fp8_e4m3fn.safetensors",
        768, 768, 4, 1.5, 456, None, 1.0,
    )
    assert qwen["2"]["inputs"]["type"] == "qwen_image"
    assert qwen["12"]["class_type"] == "SaveImage"

    wan = service._wan_i2v_graph("slow camera move", "input.png", 480, 832, 17, 16, 4, 3.0, 789)
    assert wan["1"]["inputs"]["unet_name"].startswith("wan2.1_i2v")
    assert wan["12"]["class_type"] == "SaveWEBM"

    remix = service._wan22_remix_i2v_graph(
        "slow camera move", "input.png", 480, 832, 33, 16, 8, 1.0, 789,
    )
    assert "high_lighting" in remix["1"]["inputs"]["unet_name"]
    assert "low_lighting" in remix["2"]["inputs"]["unet_name"]
    assert remix["11"]["inputs"]["end_at_step"] == 4
    assert remix["12"]["inputs"]["start_at_step"] == 4
    assert remix["15"]["class_type"] == "SaveVideo"

    ltx = service._ltx_t2v_graph("cinematic motion", 768, 512, 9, 16, 4, 3.0, 101)
    assert ltx["1"]["inputs"]["ckpt_name"].startswith("ltx-2.3")
    assert ltx["19"]["class_type"] == "SaveVideo"

    h3 = service._minimax_h3_graph("city rooftop", None, 608, 352, 39, 8, 42)
    assert h3["1"]["inputs"]["unet_name"].startswith("minimax_h3")
    assert h3["2"]["inputs"]["type"] == "minimax"
    assert h3["5"]["class_type"] == "MiniMaxH3ImageToVideo"
    assert h3["14"]["class_type"] == "SaveVideo"
    assert "first_frame" not in h3["5"]["inputs"]
    h3_i2v = service._minimax_h3_graph("city rooftop", "start.png", 608, 352, 39, 8, 42)
    assert h3_i2v["5"]["inputs"]["first_frame"] == ["15", 0]

    rewired = service._rewire_bypassed_nodes({"nodes": [
        {
            "id": 1, "type": "Bypassed", "mode": 4,
            "inputs": [{"name": "samples", "type": "LATENT", "link": 10}],
            "outputs": [{"name": "LATENT", "type": "LATENT", "links": [20]}],
        },
        {
            "id": 2, "type": "Consumer", "mode": 0,
            "inputs": [{"name": "samples", "type": "LATENT", "link": 20}],
            "outputs": [],
        },
    ]})
    assert rewired["nodes"][1]["inputs"][0]["link"] == 10


@pytest.mark.asyncio
async def test_comfyui_generate_image_auto_selects_available_default(monkeypatch):
    from app.services.comfyui_service import ComfyUIService

    service = ComfyUIService()
    captured = {}

    async def fake_inventory(**_kwargs):
        return {
            "checkpoints": [
                "other-model.safetensors",
                "RealVisXL_V5.0_Lightning_fp16.safetensors",
            ],
            "diffusion_models": ["qwen_image_2512_fp8_e4m3fn.safetensors"],
            "loras": [],
        }

    async def fake_submit(graph, timeout, metadata):
        captured["graph"] = graph
        captured["timeout"] = timeout
        captured["metadata"] = metadata
        return {"completed": True, "files": []}

    monkeypatch.setattr(service, "inventory", fake_inventory)
    monkeypatch.setattr(service, "_submit_and_wait", fake_submit)

    result = await service.generate_image(
        "a blue ceramic mug",
        width=256,
        height=256,
        steps=1,
    )

    assert result["completed"] is True
    assert captured["graph"]["1"]["inputs"]["ckpt_name"] == (
        "RealVisXL_V5.0_Lightning_fp16.safetensors"
    )
    assert captured["metadata"]["parameters"]["model"] == (
        "RealVisXL_V5.0_Lightning_fp16.safetensors"
    )


@pytest.mark.asyncio
async def test_comfyui_generate_image_uses_bigasp_v2_sampling_defaults(monkeypatch):
    from app.services.comfyui_service import ComfyUIService

    service = ComfyUIService()
    captured = {}

    async def fake_inventory(**_kwargs):
        return {
            "checkpoints": ["bigaspv2.safetensors"],
            "diffusion_models": [],
            "loras": [],
        }

    async def fake_submit(graph, timeout, metadata):
        captured["graph"] = graph
        captured["timeout"] = timeout
        captured["metadata"] = metadata
        return {"completed": True, "files": []}

    monkeypatch.setattr(service, "inventory", fake_inventory)
    monkeypatch.setattr(service, "_submit_and_wait", fake_submit)

    result = await service.generate_image(
        "cinematic rainy street",
        model="bigaspv2",
        width=512,
        height=512,
    )

    sampler = captured["graph"]["6"]["inputs"]
    assert result["completed"] is True
    assert sampler["steps"] == 40
    assert sampler["cfg"] == 3.4
    assert captured["metadata"]["parameters"]["steps"] == 40
    assert captured["metadata"]["parameters"]["cfg"] == 3.4


def test_comfyui_resolve_installed_model_accepts_unique_prefixes():
    from app.services.comfyui_service import (
        WAN22_REMIX_HIGH_MODEL,
        WAN22_REMIX_LOW_MODEL,
        WAN_I2V_MODEL,
        ComfyUIService,
    )

    qwen = "qwen_image_2512_fp8_e4m3fn.safetensors"
    assert ComfyUIService._resolve_installed_model(
        "qwen_image_2512_fp8",
        [qwen, "RealVisXL_V5.0_Lightning_fp16.safetensors"],
    ) == qwen
    assert ComfyUIService._resolve_installed_model(
        "Wan2.2_Remix_NSFW_i2v_14b_high_lighting_fp8_e4m3fn_v",
        [WAN_I2V_MODEL, WAN22_REMIX_HIGH_MODEL],
    ) == WAN22_REMIX_HIGH_MODEL
    assert ComfyUIService._resolve_installed_model(
        "Wan2.2_Remix",
        [WAN_I2V_MODEL, WAN22_REMIX_HIGH_MODEL, WAN22_REMIX_LOW_MODEL],
        group_of=lambda name: (
            WAN22_REMIX_HIGH_MODEL
            if "remix" in name.casefold()
            else name
        ),
    ) == WAN22_REMIX_HIGH_MODEL
    assert ComfyUIService._resolve_installed_model(
        "realvis",
        [
            "RealVisXL_V5.0_Lightning_fp16.safetensors",
            "realvisXL_V50.safetensors",
        ],
    ) is None


@pytest.mark.asyncio
async def test_comfyui_generate_image_resolves_truncated_qwen_name(monkeypatch):
    from app.services.comfyui_service import ComfyUIService

    service = ComfyUIService()
    captured = {}
    qwen = "qwen_image_2512_fp8_e4m3fn.safetensors"

    async def fake_inventory(**_kwargs):
        return {
            "checkpoints": ["RealVisXL_V5.0_Lightning_fp16.safetensors"],
            "diffusion_models": [qwen],
            "loras": [],
        }

    async def fake_submit(graph, timeout, metadata):
        captured["graph"] = graph
        captured["metadata"] = metadata
        return {"completed": True, "files": []}

    monkeypatch.setattr(service, "inventory", fake_inventory)
    monkeypatch.setattr(service, "_submit_and_wait", fake_submit)

    result = await service.generate_image(
        "中文招牌",
        model="qwen_image_2512_fp8",
        width=768,
        height=768,
    )

    assert result["completed"] is True
    assert captured["graph"]["2"]["inputs"]["type"] == "qwen_image"
    assert captured["graph"]["1"]["inputs"]["unet_name"] == qwen
    assert captured["metadata"]["parameters"]["model"] == qwen


@pytest.mark.asyncio
async def test_comfyui_status_keeps_workflow_inventory_compact(monkeypatch):
    from app.services.comfyui_service import ComfyUIService

    service = ComfyUIService()
    captured = {}

    async def fake_ready():
        return {"ready": True}

    async def fake_get_json(_path, _timeout):
        return {"system": {}, "devices": []}

    async def fake_inventory(**kwargs):
        captured.update(kwargs)
        return {"workflows": [{"name": "sample", "path": "/sample.json"}]}

    monkeypatch.setattr(service, "ensure_ready", fake_ready)
    monkeypatch.setattr(service, "_get_json", fake_get_json)
    monkeypatch.setattr(service, "inventory", fake_inventory)

    result = await service.status()

    assert result["ready"] is True
    assert result["generation"]["video_pipelines"][1]["available"] is False
    assert captured == {
        "include_workflow_details": False,
        "_ensure_runtime": False,
    }


@pytest.mark.asyncio
async def test_comfyui_model_catalog_lists_supported_models_and_blocks_ltx_when_memory_is_low(
    monkeypatch,
):
    from app.services.comfyui_service import (
        LTX_MIN_FREE_MEMORY,
        WAN22_REMIX_HIGH_MODEL,
        WAN22_REMIX_LOW_MODEL,
        WAN22_REMIX_TEXT_ENCODER,
        ComfyUIService,
    )

    service = ComfyUIService()

    async def fake_inventory(**_kwargs):
        return {
            "checkpoints": [
                "RealVisXL_V5.0_Lightning_fp16.safetensors",
                "JoyAI-Echo-release.safetensors",
                "ltx-2.3-22b-dev.safetensors",
                "nested/model.safetensors",
            ],
            "diffusion_models": [
                "qwen_image_2512_fp8_e4m3fn.safetensors",
                "wan2.1_i2v_480p_14B_fp8_scaled.safetensors",
                WAN22_REMIX_HIGH_MODEL,
                WAN22_REMIX_LOW_MODEL,
                "minimax_h3_fl2va_pruned_int8_convrot.safetensors",
            ],
            "loras": [],
            "vae": [
                "wan_2.1_vae.safetensors",
                "minimax_h3_video_vae_fp16.safetensors",
                "minimax_h3_audio_vae_fp32.safetensors",
            ],
            "text_encoders": [
                "umt5_xxl_fp8_e4m3fn_scaled.safetensors",
                WAN22_REMIX_TEXT_ENCODER,
                "gemma-3-12b-it/model-00001-of-00005.safetensors",
                "qwen3vl_32b_minimax_h3_nvfp4_awq.safetensors",
            ],
            "clip_vision": ["clip_vision_h.safetensors"],
            "upscalers": [],
            "workflows": [],
            "capabilities": {
                "WanImageToVideo": True,
                "SaveWEBM": True,
                "ModelSamplingSD3": True,
                "KSamplerAdvanced": True,
                "CreateVideo": True,
                "LTXVGemmaCLIPModelLoader": True,
                "LTXVAudioVAELoader": True,
                "LTXVConditioning": True,
                "SaveVideo": True,
                "MiniMaxH3ImageToVideo": True,
                "VAEDecodeAudio": True,
            },
        }

    async def fake_get_json(_path, _timeout):
        return {
            "system": {
                "ram_total": 128 * 1024**3,
                "ram_free": LTX_MIN_FREE_MEMORY - 1,
            },
            "devices": [],
        }

    monkeypatch.setattr(service, "inventory", fake_inventory)
    monkeypatch.setattr(service, "_get_json", fake_get_json)

    catalog = await service.model_catalog()

    assert [item["name"] for item in catalog["image_models"]] == [
        "RealVisXL_V5.0_Lightning_fp16.safetensors",
        "qwen_image_2512_fp8_e4m3fn.safetensors",
    ]
    assert catalog["image_models"][0]["recommended"] is True
    wan, ltx, h3 = catalog["video_pipelines"]
    assert wan["installed"] is True
    assert wan["available"] is True
    assert wan["requires_input_image"] is False
    assert wan["auto_generates_input_image"] is True
    assert [item["name"] for item in wan["models"]] == [
        "wan2.1_i2v_480p_14B_fp8_scaled.safetensors",
        WAN22_REMIX_HIGH_MODEL,
    ]
    assert wan["models"][1]["companion_models"] == [WAN22_REMIX_LOW_MODEL]
    assert ltx["installed"] is True
    assert ltx["available"] is False
    assert ltx["minimum_free_memory_bytes"] == LTX_MIN_FREE_MEMORY
    assert "停止本地大语言模型" in ltx["warning"]
    assert h3["id"] == "minimax-h3"
    assert h3["installed"] is True
    assert h3["available"] is True
    assert h3["models"][0]["name"].startswith("minimax_h3")


@pytest.mark.asyncio
async def test_comfyui_generate_video_rejects_unsafe_ltx_before_submitting(monkeypatch):
    from app.services.comfyui_service import ComfyUIError, ComfyUIService

    service = ComfyUIService()
    submitted = False

    async def fake_ready():
        return {"ready": True}

    async def fake_catalog(**_kwargs):
        return {
            "video_pipelines": [{
                "id": "ltx-t2v",
                "installed": True,
                "available": False,
                "missing": [],
                "warning": "请先停止本地大语言模型再运行 LTX。",
            }],
        }

    async def fake_submit(*_args, **_kwargs):
        nonlocal submitted
        submitted = True
        return {}

    monkeypatch.setattr(service, "ensure_ready", fake_ready)
    monkeypatch.setattr(service, "model_catalog", fake_catalog)
    monkeypatch.setattr(service, "_submit_and_wait", fake_submit)

    with pytest.raises(ComfyUIError, match="停止本地大语言模型"):
        await service.generate_video("cinematic clouds", pipeline="ltx-t2v")
    assert submitted is False


@pytest.mark.asyncio
async def test_comfyui_wan_video_auto_generates_starter_image(monkeypatch):
    from app.services.comfyui_service import WAN_I2V_MODEL, ComfyUIService

    service = ComfyUIService()
    captured = {}

    async def fake_ready():
        return {"ready": True}

    async def fake_catalog(**_kwargs):
        return {
            "video_pipelines": [{
                "id": "wan-i2v",
                "installed": True,
                "available": True,
                "missing": [],
                "models": [{
                    "name": WAN_I2V_MODEL,
                    "installed": True,
                    "missing": [],
                    "recommended": True,
                }],
            }],
        }

    async def fake_generate_image(prompt, **kwargs):
        captured["starter"] = {"prompt": prompt, **kwargs}
        return {"prompt_id": "starter-prompt", "files": [{"filename": "start.png"}]}

    async def fake_upload_generated(result):
        assert result["prompt_id"] == "starter-prompt"
        return "uploaded-start.png"

    async def fake_submit(graph, timeout, metadata):
        captured["graph"] = graph
        captured["timeout"] = timeout
        captured["metadata"] = metadata
        return {"completed": True, "files": []}

    monkeypatch.setattr(service, "ensure_ready", fake_ready)
    monkeypatch.setattr(service, "model_catalog", fake_catalog)
    monkeypatch.setattr(service, "generate_image", fake_generate_image)
    monkeypatch.setattr(service, "_upload_generated_image", fake_upload_generated)
    monkeypatch.setattr(service, "_submit_and_wait", fake_submit)

    result = await service.generate_video(
        "cinematic clouds moving over a city",
        pipeline="wan-i2v",
        width=480,
        height=832,
        seed=42,
    )

    assert result["completed"] is True
    assert captured["starter"]["width"] == 480
    assert captured["starter"]["height"] == 832
    assert captured["graph"]["5"]["inputs"]["image"] == "uploaded-start.png"
    parameters = captured["metadata"]["parameters"]
    assert parameters["input_image"] is None
    assert parameters["input_image_source"] == "auto-generated"
    assert parameters["starter_image_prompt_id"] == "starter-prompt"
    assert parameters["model"] == WAN_I2V_MODEL


@pytest.mark.asyncio
async def test_comfyui_wan22_requested_model_uses_two_stage_graph(monkeypatch):
    from app.services.comfyui_service import (
        WAN22_REMIX_HIGH_MODEL,
        WAN22_REMIX_LOW_MODEL,
        ComfyUIService,
    )

    service = ComfyUIService()
    captured = {}

    async def fake_ready():
        return {"ready": True}

    async def fake_catalog(**_kwargs):
        return {
            "video_pipelines": [{
                "id": "wan-i2v",
                "installed": True,
                "available": True,
                "missing": [],
                "models": [{
                    "name": WAN22_REMIX_HIGH_MODEL,
                    "installed": True,
                    "missing": [],
                    "recommended": True,
                    "companion_models": [WAN22_REMIX_LOW_MODEL],
                }],
            }],
        }

    async def fake_upload(path):
        assert path == "/safe/session/source.png"
        return "source.png"

    async def fake_submit(graph, timeout, metadata):
        captured["graph"] = graph
        captured["metadata"] = metadata
        return {"completed": True, "files": []}

    monkeypatch.setattr(service, "ensure_ready", fake_ready)
    monkeypatch.setattr(service, "model_catalog", fake_catalog)
    monkeypatch.setattr(service, "_upload_image", fake_upload)
    monkeypatch.setattr(service, "_submit_and_wait", fake_submit)

    await service.generate_video(
        "slow cinematic movement",
        pipeline="wan-i2v",
        model=WAN22_REMIX_HIGH_MODEL,
        input_image="/safe/session/source.png",
        steps=8,
        cfg=1.0,
    )

    assert captured["graph"]["1"]["inputs"]["unet_name"] == WAN22_REMIX_HIGH_MODEL
    assert captured["graph"]["2"]["inputs"]["unet_name"] == WAN22_REMIX_LOW_MODEL
    assert captured["graph"]["15"]["class_type"] == "SaveVideo"
    assert captured["metadata"]["parameters"]["model"] == WAN22_REMIX_HIGH_MODEL
    assert captured["metadata"]["parameters"]["input_image_source"] == "local-file"


@pytest.mark.asyncio
async def test_comfyui_generate_video_routes_minimax_h3_from_model_name(monkeypatch):
    from app.services.comfyui_service import MINIMAX_H3_MODEL, ComfyUIService

    service = ComfyUIService()
    captured = {}

    async def fake_ready():
        return {"ready": True}

    async def fake_catalog(**_kwargs):
        return {
            "video_pipelines": [
                {
                    "id": "wan-i2v",
                    "installed": True,
                    "available": True,
                    "missing": [],
                    "models": [{"name": "wan2.1_i2v_480p_14B_fp8_scaled.safetensors", "installed": True}],
                },
                {
                    "id": "minimax-h3",
                    "installed": True,
                    "available": True,
                    "missing": [],
                    "models": [{"name": MINIMAX_H3_MODEL, "label": "MiniMax H3", "installed": True}],
                },
            ],
        }

    async def fake_submit(graph, timeout, metadata):
        captured["graph"] = graph
        captured["metadata"] = metadata
        return {"completed": True, "files": []}

    monkeypatch.setattr(service, "ensure_ready", fake_ready)
    monkeypatch.setattr(service, "model_catalog", fake_catalog)
    monkeypatch.setattr(service, "_submit_and_wait", fake_submit)

    await service.generate_video(
        "golden hour rooftop",
        pipeline="wan-i2v",
        model="minimax_h3",
    )

    assert captured["graph"]["5"]["class_type"] == "MiniMaxH3ImageToVideo"
    assert captured["graph"]["1"]["inputs"]["unet_name"] == MINIMAX_H3_MODEL
    assert captured["metadata"]["parameters"]["pipeline"] == "minimax-h3"
    assert captured["metadata"]["parameters"]["model"] == MINIMAX_H3_MODEL
    assert captured["metadata"]["parameters"]["input_image_source"] is None


@pytest.mark.asyncio
async def test_comfyui_generate_video_resolves_truncated_wan22_name(monkeypatch):
    from app.services.comfyui_service import (
        WAN22_REMIX_HIGH_MODEL,
        WAN22_REMIX_LOW_MODEL,
        WAN_I2V_MODEL,
        ComfyUIError,
        ComfyUIService,
    )

    service = ComfyUIService()
    captured = {}

    async def fake_ready():
        return {"ready": True}

    async def fake_catalog(**_kwargs):
        return {
            "video_pipelines": [{
                "id": "wan-i2v",
                "installed": True,
                "available": True,
                "missing": [],
                "models": [
                    {
                        "name": WAN_I2V_MODEL,
                        "installed": True,
                        "missing": [],
                        "recommended": False,
                    },
                    {
                        "name": WAN22_REMIX_HIGH_MODEL,
                        "installed": True,
                        "missing": [],
                        "recommended": True,
                        "companion_models": [WAN22_REMIX_LOW_MODEL],
                    },
                ],
            }],
        }

    async def fake_upload(path):
        return "source.png"

    async def fake_submit(graph, timeout, metadata):
        captured["graph"] = graph
        captured["metadata"] = metadata
        return {"completed": True, "files": []}

    monkeypatch.setattr(service, "ensure_ready", fake_ready)
    monkeypatch.setattr(service, "model_catalog", fake_catalog)
    monkeypatch.setattr(service, "_upload_image", fake_upload)
    monkeypatch.setattr(service, "_submit_and_wait", fake_submit)

    await service.generate_video(
        "slow cinematic movement",
        pipeline="wan-i2v",
        model="Wan2.2_Remix_NSFW_i2v_14b_high_lighting_fp8_e4m3fn_v",
        input_image="/safe/session/source.png",
    )

    assert captured["graph"]["1"]["inputs"]["unet_name"] == WAN22_REMIX_HIGH_MODEL
    assert captured["graph"]["2"]["inputs"]["unet_name"] == WAN22_REMIX_LOW_MODEL
    assert captured["metadata"]["parameters"]["model"] == WAN22_REMIX_HIGH_MODEL

    with pytest.raises(ComfyUIError, match="Unsupported wan-i2v model: missing-model"):
        await service.generate_video(
            "slow cinematic movement",
            pipeline="wan-i2v",
            model="missing-model",
            input_image="/safe/session/source.png",
        )


@pytest.mark.asyncio
async def test_comfyui_saved_selection_overrides_automatic_tool_arguments(monkeypatch):
    from app.services.comfyui_service import comfyui_service

    captured = {}

    async def fake_generate_image(prompt, **kwargs):
        captured["image"] = {"prompt": prompt, **kwargs}
        return {"completed": True}

    async def fake_generate_video(prompt, **kwargs):
        captured["video"] = {"prompt": prompt, **kwargs}
        return {"completed": True}

    monkeypatch.setattr(
        plugin_registry,
        "get",
        lambda _plugin_id: {
            "config": {
                "image_model": "selected-image.safetensors",
                "video_pipeline": "ltx-t2v",
            },
        },
    )
    monkeypatch.setattr(comfyui_service, "generate_image", fake_generate_image)
    monkeypatch.setattr(comfyui_service, "generate_video", fake_generate_video)

    await plugin_registry._dispatch(
        {"handler": "comfyui_generate_image"},
        "session",
        {"prompt": "a mug", "model": "automatic.safetensors"},
    )
    await plugin_registry._dispatch(
        {"handler": "comfyui_generate_video"},
        "session",
        {"prompt": "moving clouds", "pipeline": "wan-i2v"},
    )

    assert captured["image"]["model"] == "selected-image.safetensors"
    assert captured["video"]["pipeline"] == "ltx-t2v"

    await plugin_registry._dispatch(
        {"handler": "comfyui_generate_video"},
        "session",
        {
            "prompt": "slow movement",
            "pipeline": "wan-i2v",
            "_user_requested_model": "Wan2.2_Remix.safetensors",
            "input_image": "/safe/source.png",
        },
    )
    assert captured["video"]["pipeline"] == "wan-i2v"
    assert captured["video"]["model"] == "Wan2.2_Remix.safetensors"
    assert captured["video"]["input_image"] == "/safe/source.png"

    await plugin_registry._dispatch(
        {"handler": "comfyui_generate_image"},
        "session",
        {
            "prompt": "a rainy street",
            "model": "automatic.safetensors",
            "_user_requested_model": "bigaspv2.safetensors",
        },
    )
    assert captured["image"]["model"] == "bigaspv2.safetensors"


@pytest.mark.asyncio
async def test_comfyui_runtime_restarts_stale_active_systemd_unit(monkeypatch):
    from app.services.comfyui_runtime import ComfyUIRuntimeService

    runtime = ComfyUIRuntimeService()
    calls = []

    async def active(_unit):
        return True

    async def systemctl(*arguments, timeout=45):
        calls.append((arguments, timeout))
        return 0, ""

    monkeypatch.setattr(runtime, "_unit_active", active)
    monkeypatch.setattr(runtime, "_systemctl", systemctl)

    started, already_active, output = await runtime._start_unit(
        "aetherchat-comfyui.service"
    )

    assert (started, already_active, output) == (True, True, "")
    assert calls == [(('restart', 'aetherchat-comfyui.service'), 45)]
    assert runtime._managed_unit is None


@pytest.mark.asyncio
async def test_comfyui_runtime_diagnoses_missing_process_after_oom(monkeypatch):
    from app.services.comfyui_runtime import ComfyUIRuntimeService

    runtime = ComfyUIRuntimeService()

    async def probe(_base_url):
        return False

    async def active(_unit):
        return True

    async def docker_diagnostics(_base_url):
        return {
            "container": "comfy-spark",
            "status": "running",
            "oom_killed": True,
            "comfyui_process_running": False,
        }

    monkeypatch.setattr(runtime, "_probe", probe)
    monkeypatch.setattr(runtime, "_systemd_unit", lambda: "aetherchat-comfyui.service")
    monkeypatch.setattr(runtime, "_unit_active", active)
    monkeypatch.setattr(runtime, "_docker_diagnostics", docker_diagnostics)

    diagnostic = await runtime.diagnostics("http://127.0.0.1:8188")

    assert diagnostic["reason"] == "process_oom_killed"
    assert "内存不足" in diagnostic["message"]
    assert diagnostic["unit_active"] is True


@pytest.mark.asyncio
async def test_comfyui_runtime_starts_and_stops_only_its_managed_process(
    tmp_path,
    monkeypatch,
):
    from app.services import comfyui_runtime as runtime_module
    from app.services.comfyui_runtime import ComfyUIRuntimeService

    class FakeProcess:
        pid = 43210

        def __init__(self):
            self.returncode = None
            self.signals = []

        async def wait(self):
            return self.returncode

    runtime = ComfyUIRuntimeService()
    process = FakeProcess()
    captured = {}

    async def fake_create(*command, **kwargs):
        captured["command"] = command
        captured["kwargs"] = kwargs
        return process

    async def fake_probe(_base_url):
        return runtime._process is process and process.returncode is None

    def fake_signal(target, force=False):
        target.signals.append("kill" if force else "terminate")
        target.returncode = -15 if not force else -9

    monkeypatch.setenv("AETHERCHAT_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setattr(runtime, "_systemd_unit", lambda: None)
    monkeypatch.setattr(
        runtime,
        "_command",
        lambda _url, _roots: (
            ["/fake/python", "/fake/ComfyUI/main.py", "--port", "9199"],
            tmp_path,
        ),
    )
    monkeypatch.setattr(runtime, "_probe", fake_probe)
    monkeypatch.setattr(runtime, "_signal_process", fake_signal)
    monkeypatch.setattr(
        runtime_module.asyncio,
        "create_subprocess_exec",
        fake_create,
    )

    ready = await runtime.ensure_ready(
        "http://127.0.0.1:9199",
        [tmp_path],
    )
    assert ready["ready"] is True
    assert ready["managed"] is True
    assert ready["started_by_aetherchat"] is True
    assert ready["pid"] == process.pid
    assert captured["kwargs"]["start_new_session"] is True

    stopped = await runtime.shutdown()
    assert stopped == {
        "stopped": True,
        "pid": process.pid,
        "unit": None,
    }
    assert process.signals == ["terminate"]
    assert await runtime.shutdown() == {
        "stopped": False,
        "reason": "not_managed",
    }

    external = ComfyUIRuntimeService()

    async def external_probe(_base_url):
        return True

    monkeypatch.setattr(external, "_probe", external_probe)
    monkeypatch.setattr(external, "_systemd_unit", lambda: None)
    status = await external.ensure_ready(
        "http://127.0.0.1:9199",
        [tmp_path],
    )
    assert status["ready"] is True
    assert status["managed"] is False
    assert await external.shutdown() == {
        "stopped": False,
        "reason": "not_managed",
    }


@pytest.mark.asyncio
async def test_comfyui_runtime_owns_systemd_unit_only_when_it_starts_it(
    monkeypatch,
):
    from app.services.comfyui_runtime import ComfyUIRuntimeService

    runtime = ComfyUIRuntimeService()
    state = {"active": False, "ready": False, "stops": 0}

    async def probe(_base_url):
        return state["ready"]

    async def active(_unit):
        return state["active"]

    async def start(unit):
        state["active"] = True
        state["ready"] = True
        runtime._managed_unit = unit
        runtime._started_at = "2026-07-28T00:00:00+00:00"
        runtime._command_label = f"systemctl --user start {unit}"
        return True, False, ""

    async def stop(_unit):
        state["active"] = False
        state["ready"] = False
        state["stops"] += 1
        return True, ""

    monkeypatch.setattr(
        runtime,
        "_systemd_unit",
        lambda: "aetherchat-comfyui.service",
    )
    monkeypatch.setattr(runtime, "_probe", probe)
    monkeypatch.setattr(runtime, "_unit_active", active)
    monkeypatch.setattr(runtime, "_start_unit", start)
    monkeypatch.setattr(runtime, "_stop_unit", stop)

    status = await runtime.ensure_ready(
        "http://127.0.0.1:8188",
        [],
    )
    assert status["managed"] is True
    assert status["runtime"] == "systemd_user"
    assert status["unit"] == "aetherchat-comfyui.service"
    assert await runtime.shutdown() == {
        "stopped": True,
        "pid": None,
        "unit": "aetherchat-comfyui.service",
    }
    assert state["stops"] == 1


def test_comfyui_workflow_descriptor_exposes_nodes_edges_and_typed_overrides():
    from app.services.comfyui_service import ComfyUIError, ComfyUIService

    graph = {
        "1": {
            "class_type": "CheckpointLoaderSimple",
            "inputs": {"ckpt_name": "model.safetensors"},
        },
        "6": {
            "class_type": "KSampler",
            "inputs": {
                "model": ["1", 0],
                "seed": 123,
                "steps": 8,
                "cfg": 2.5,
            },
            "_meta": {"title": "Primary sampler"},
        },
    }
    info = {
        "CheckpointLoaderSimple": {
            "display_name": "Load checkpoint",
            "input": {"required": {"ckpt_name": [["model.safetensors", "other.safetensors"]]}},
        },
        "KSampler": {
            "display_name": "Sampler",
            "input": {
                "required": {
                    "model": ["MODEL"],
                    "seed": ["INT", {"min": 0, "max": 999}],
                    "steps": ["INT", {"min": 1, "max": 100}],
                    "cfg": ["FLOAT", {"min": 0.1, "max": 30.0}],
                },
                "optional": {
                    "denoise": ["FLOAT", {"min": 0.0, "max": 1.0}],
                    "mask": ["MASK"],
                },
            },
        },
    }
    descriptor = ComfyUIService._describe_graph(graph, info)
    assert descriptor["node_map"] == {
        "1": "Load checkpoint",
        "6": "Primary sampler",
    }
    assert descriptor["dependencies"] == [{"from": "1", "to": "6", "input": "model"}]
    assert "1 (Load checkpoint) -> 6 (Primary sampler) via model" in descriptor["dependency_text"]
    inputs = {
        (item["node_id"], item["input"]): item
        for item in descriptor["overridable_inputs"]
    }
    assert inputs[("6", "seed")]["type"] == "INT"
    assert inputs[("6", "steps")]["options"]["max"] == 100
    assert inputs[("6", "denoise")]["present"] is False
    assert ("6", "mask") not in inputs
    ComfyUIService._validate_override_value(inputs[("6", "steps")], 24)
    with pytest.raises(ComfyUIError, match="expects INT"):
        ComfyUIService._validate_override_value(inputs[("6", "steps")], "24")
    with pytest.raises(ComfyUIError, match="must be <="):
        ComfyUIService._validate_override_value(inputs[("6", "steps")], 101)
    formatted = ComfyUIService._format_node_error(
        {"node_id": "6", "exception_message": "bad input"},
        graph,
    )
    assert '"node_label": "Primary sampler"' in formatted
    assert "6=Primary sampler" in formatted
    prompt_graph = {
        **graph,
        "3": {
            "class_type": "CLIPTextEncode",
            "inputs": {"text": "雨夜中的重庆"},
            "_meta": {"title": "Positive prompt"},
        },
    }
    metadata = ComfyUIService._generation_metadata(
        prompt_id="prompt-1",
        graph=prompt_graph,
        metadata={"operation": "run_workflow"},
        files=[],
        submitted_at="2026-07-28T00:00:00+00:00",
    )
    assert metadata["provenance"]["prompts"] == [{
        "node_id": "3",
        "node_name": "Positive prompt",
        "input": "text",
        "value": "雨夜中的重庆",
    }]


@pytest.mark.asyncio
async def test_comfyui_metadata_is_written_as_agent_artifact_and_shared_entity(
    tmp_path,
):
    prompt_id = str(uuid4())
    session_id = str(uuid4())
    generated_image = tmp_path / "cover.png"
    generated_image.write_bytes(b"\x89PNG\r\n\x1a\nmock-image")
    result = await plugin_registry._persist_comfyui_metadata(
        session_id,
        "generate_image",
        "model:qwen",
        {
            "prompt_id": prompt_id,
            "files": [{
                "filename": "cover.png",
                "url": "http://127.0.0.1:8188/view?cover",
                "_local_path": str(generated_image),
            }],
            "generation_metadata": {
                "schema": "aetherchat.agent-artifact/v1",
                "schema_version": "1.0",
                "artifact_type": "comfyui_generation",
                "engine": "ComfyUI",
                "operation": "generate_image",
                "prompt_id": prompt_id,
                "submitted_at": "2026-07-28T00:00:00+00:00",
                "completed_at": "2026-07-28T00:00:01+00:00",
                "parameters": {
                    "prompt": "雨夜中的重庆",
                    "model": "qwen-image.safetensors",
                    "seed": 42,
                    "width": 1024,
                    "height": 768,
                },
                "workflow": {"node_map": {"1": "Sampler"}, "dependencies": []},
                "provenance": {
                    "models": ["qwen-image.safetensors"],
                    "seeds": [42],
                    "dimensions": [{"width": 1024, "height": 768}],
                },
                "outputs": [{"filename": "cover.png", "url": "http://127.0.0.1:8188/view?cover"}],
            },
        },
    )
    metadata_path = result["metadata_artifact"]["path"]
    assert metadata_path == f"generated/comfyui/{prompt_id}.metadata.json"
    stored = json.loads(await plugin_registry.artifacts.read(session_id, metadata_path))
    assert stored["parameters"]["prompt"] == "雨夜中的重庆"
    assert stored["producer"] == {
        "agent": "model:qwen",
        "plugin": "comfyui",
        "action": "generate_image",
    }
    media = next(
        item for item in result["generated_files"]
        if item["kind"] == "generated_media"
    )
    assert media["path"] == f"generated/comfyui/{prompt_id}/cover.png"
    native = await plugin_registry.artifacts.native_location(
        session_id,
        media["path"],
    )
    assert Path(native["path"]).read_bytes() == generated_image.read_bytes()
    assert "_local_path" not in result["files"][0]
    assert result["files"][0]["artifact_path"] == media["path"]
    shared = chat_database.get_memory(session_id)["entities"]
    entity = next(item for item in shared if item["key"] == f"asset:comfyui:{prompt_id}")
    assert entity["metadata"]["seed"] == 42
    assert entity["metadata"]["artifact_path"] == metadata_path


@pytest.mark.asyncio
async def test_comfyui_result_does_not_embed_video_as_an_image(monkeypatch):
    from app.services.comfyui_service import ComfyUIService

    class FakeResponse:
        content = b"video bytes"
        headers = {"content-type": "video/webm"}

        def raise_for_status(self):
            return None

    class FakeClient:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            return None

        async def get(self, _url):
            return FakeResponse()

    service = ComfyUIService()
    monkeypatch.setattr(service, "_client", lambda _timeout=30: FakeClient())
    result = await service._result("prompt", [{
        "filename": "clip.webm", "subfolder": "AetherChat", "type": "output", "kind": "videos",
    }])
    assert result["images"] == []
    assert result["files"][0]["kind"] == "videos"
    assert "image_data_url" not in result["files"][0]


def test_protected_desktop_document_roundtrips_and_platform_capabilities(client, tmp_path, monkeypatch):
    from datetime import date

    from docx import Document
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell.cell import MergedCell
    from pptx import Presentation
    from pptx.util import Inches
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas

    token = "offline-desktop-test-token"
    monkeypatch.setenv("AETHERCHAT_DESKTOP_TOKEN", token)
    headers = {"X-Aether-Desktop-Token": token}

    assert client.get("/desktop/status").status_code == 403
    status = client.get("/desktop/status", headers=headers)
    assert status.status_code == 200
    from app.services.desktop_service import mouse_service

    automation_available = mouse_service.available()
    assert status.json()["mouse"] is automation_available
    assert status.json()["model_computer_use"] is automation_available
    assert status.json()["computer_protocol"] == "computer_20251124"
    assert {"screenshot", "left_click", "zoom", "hold_key"} <= set(status.json()["computer_actions"])
    assert {"docx", "xlsx", "pptx", "pdf"} <= set(status.json()["document_formats"])

    markdown = tmp_path / "notes.md"
    markdown.write_text("before", encoding="utf-8")
    opened_text = client.post("/desktop/document/read", headers=headers, json={"path": str(markdown)})
    assert opened_text.json()["data"]["text"] == "before"
    saved_text = client.post("/desktop/document/write", headers=headers, json={
        "path": str(markdown),
        "source_path": str(markdown),
        "data": {"text": "after"},
        "create_backup": True,
    })
    assert saved_text.status_code == 200
    assert markdown.read_text(encoding="utf-8") == "after"
    assert Path(saved_text.json()["backup_path"]).read_text(encoding="utf-8") == "before"

    new_markdown = tmp_path / "created.md"
    created_text = client.post("/desktop/document/write", headers=headers, json={
        "path": str(new_markdown),
        "data": {"text": "created without a source file"},
        "create_backup": True,
    })
    assert created_text.status_code == 200
    assert created_text.json()["backup_path"] is None
    assert new_markdown.read_text(encoding="utf-8") == "created without a source file"

    missing_source_docx = client.post("/desktop/document/write", headers=headers, json={
        "path": str(tmp_path / "new.docx"),
        "data": {"paragraphs": ["new"]},
        "create_backup": False,
    })
    assert missing_source_docx.status_code == 400
    assert "source_path" in missing_source_docx.json()["detail"]

    gb_text = tmp_path / "gb18030.txt"
    gb_text.write_bytes("中文编码".encode("gb18030"))
    opened_gb = client.post("/desktop/document/read", headers=headers, json={"path": str(gb_text)}).json()
    assert opened_gb["data"] == {"text": "中文编码", "encoding": "gb18030"}
    assert client.post("/desktop/document/write", headers=headers, json={
        "path": str(gb_text), "source_path": str(gb_text), "data": opened_gb["data"], "create_backup": False,
    }).status_code == 200
    assert gb_text.read_bytes().decode("gb18030") == "中文编码"

    source_docx = tmp_path / "source.docx"
    output_docx = tmp_path / "edited.docx"
    document = Document()
    document.add_heading("Original title", level=1)
    document.add_paragraph("Original body")
    styled_paragraph = document.add_paragraph()
    bold_run = styled_paragraph.add_run("Bold")
    bold_run.bold = True
    italic_run = styled_paragraph.add_run("Italic")
    italic_run.italic = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.save(source_docx)
    docx_data = client.post("/desktop/document/read", headers=headers, json={"path": str(source_docx)}).json()["data"]
    docx_data["paragraphs"][0] = "Edited title"
    docx_data["tables"][0][0][1] = "Edited cell"
    assert client.post("/desktop/document/write", headers=headers, json={
        "path": str(output_docx), "source_path": str(source_docx), "data": docx_data, "create_backup": False,
    }).status_code == 200
    edited_docx = Document(output_docx)
    assert edited_docx.paragraphs[0].text == "Edited title"
    assert edited_docx.tables[0].cell(0, 1).text == "Edited cell"
    assert [run.text for run in edited_docx.paragraphs[2].runs] == ["Bold", "Italic"]
    assert edited_docx.paragraphs[2].runs[0].bold is True
    assert edited_docx.paragraphs[2].runs[1].italic is True

    source_xlsx = tmp_path / "source.xlsx"
    output_xlsx = tmp_path / "edited.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet["A1"] = "Original"
    worksheet["B1"] = 2
    worksheet["C1"] = "=B1*2"
    worksheet["D1"] = date(2026, 7, 12)
    worksheet.merge_cells("E1:F1")
    worksheet["E1"] = "Merged"
    workbook.save(source_xlsx)
    xlsx_data = client.post("/desktop/document/read", headers=headers, json={"path": str(source_xlsx)}).json()["data"]
    xlsx_data["sheets"][0]["values"][0][0] = "Edited"
    xlsx_data["sheets"][0]["values"][0][1] = 3
    assert client.post("/desktop/document/write", headers=headers, json={
        "path": str(output_xlsx), "source_path": str(source_xlsx), "data": xlsx_data, "create_backup": False,
    }).status_code == 200
    edited_xlsx = load_workbook(output_xlsx, data_only=False)
    assert edited_xlsx["Data"]["A1"].value == "Edited"
    assert edited_xlsx["Data"]["B1"].value == 3
    assert edited_xlsx["Data"]["C1"].value == "=B1*2"
    stored_date = edited_xlsx["Data"]["D1"].value
    assert (stored_date.date() if hasattr(stored_date, "date") else stored_date) == date(2026, 7, 12)
    assert "E1:F1" in {str(item) for item in edited_xlsx["Data"].merged_cells.ranges}
    assert isinstance(edited_xlsx["Data"]["F1"], MergedCell)

    source_pptx = tmp_path / "source.pptx"
    output_pptx = tmp_path / "edited.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Original slide"
    slide.placeholders[1].text = "Original body"
    styled_box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1))
    styled_box.text_frame.clear()
    styled_run_one = styled_box.text_frame.paragraphs[0].add_run()
    styled_run_one.text = "Bold"
    styled_run_one.font.bold = True
    styled_run_two = styled_box.text_frame.paragraphs[0].add_run()
    styled_run_two.text = "Italic"
    styled_run_two.font.italic = True
    presentation.save(source_pptx)
    pptx_data = client.post("/desktop/document/read", headers=headers, json={"path": str(source_pptx)}).json()["data"]
    title_entry = next(item for item in pptx_data["slides"][0]["texts"] if item["text"] == "Original slide")
    title_entry["text"] = "Edited slide"
    assert client.post("/desktop/document/write", headers=headers, json={
        "path": str(output_pptx), "source_path": str(source_pptx), "data": pptx_data, "create_backup": False,
    }).status_code == 200
    edited_pptx = Presentation(output_pptx)
    assert edited_pptx.slides[0].shapes.title.text == "Edited slide"
    edited_styled_box = next(shape for shape in edited_pptx.slides[0].shapes if shape.name == styled_box.name)
    edited_runs = edited_styled_box.text_frame.paragraphs[0].runs
    assert [run.text for run in edited_runs] == ["Bold", "Italic"]
    assert edited_runs[0].font.bold is True
    assert edited_runs[1].font.italic is True

    source_pdf = tmp_path / "source.pdf"
    output_pdf = tmp_path / "edited.pdf"
    pdf_canvas = canvas.Canvas(str(source_pdf))
    pdf_canvas.drawString(72, 720, "Original PDF")
    pdf_canvas.save()
    source_pdf_metadata = PdfReader(source_pdf).metadata
    pdf_data = client.post("/desktop/document/read", headers=headers, json={"path": str(source_pdf)}).json()["data"]
    pdf_data["metadata"]["title"] = "Edited PDF"
    pdf_data["append_note"] = "AetherChat test note"
    assert client.post("/desktop/document/write", headers=headers, json={
        "path": str(output_pdf), "source_path": str(source_pdf), "data": pdf_data, "create_backup": False,
    }).status_code == 200
    edited_pdf = PdfReader(output_pdf)
    assert len(edited_pdf.pages) == 2
    assert edited_pdf.metadata.title == "Edited PDF"
    assert edited_pdf.metadata.get("/Producer") == source_pdf_metadata.get("/Producer")

    paged_pdf = tmp_path / "paged-note.pdf"
    pdf_data["append_note"] = "\n".join(f"Long note line {index}" for index in range(180))
    assert client.post("/desktop/document/write", headers=headers, json={
        "path": str(paged_pdf), "source_path": str(source_pdf), "data": pdf_data, "create_backup": False,
    }).status_code == 200
    assert len(PdfReader(paged_pdf).pages) > 2

    current_mouse = client.post("/desktop/mouse", headers=headers, json={"action": "get"})
    if automation_available:
        assert current_mouse.status_code == 200
        assert {"x", "y"} <= set(current_mouse.json()["position"])
        invalid_move = client.post("/desktop/mouse", headers=headers, json={
            "action": "move", "x": 2_000_000_000, "y": 2_000_000_000,
        })
        assert invalid_move.status_code == 400
    else:
        assert current_mouse.status_code == 400
        assert "unavailable" in current_mouse.json()["detail"].lower() or "X11" in current_mouse.json()["detail"]

    rejected_unc = client.post("/desktop/document/read", headers=headers, json={"path": r"\\server\share\file.docx"})
    assert rejected_unc.status_code == 400


@pytest.mark.asyncio
async def test_computer_use_can_observe_windows_without_performing_actions():
    from app.services.computer_use_service import computer_use_service

    if not computer_use_service.platform_available():
        pytest.skip("No supported Windows or X11 desktop runtime")

    windows = await computer_use_service.list_windows()
    assert isinstance(windows["windows"], list)
    desktop = await computer_use_service.get_state(
        include_screenshot=False,
        include_accessibility=True,
    )
    assert desktop["snapshot_id"]
    assert isinstance(desktop["accessibility_tree"], str)
    if windows["windows"]:
        window = windows["windows"][0]
        state = await computer_use_service.get_state(
            window_id=window["window_id"],
            include_screenshot=False,
            include_accessibility=True,
        )
        assert state["window"]["window_id"] == window["window_id"]
        assert isinstance(state["elements"], list)


@pytest.mark.asyncio
async def test_computer_use_stable_and_fuzzy_selectors_can_refresh_stale_snapshot(monkeypatch):
    from app.services.computer_use_service import (
        ComputerUseError,
        ComputerUseService,
        _Snapshot,
    )

    service = ComputerUseService()
    element_id = service._stable_element_id("Button", "save-button", "保存", "0/2")
    assert element_id == service._stable_element_id("Button", "save-button", "Save", "9/8")
    service._snapshots["snapshot"] = _Snapshot(
        window_id=42,
        created_at=time.monotonic(),
        elements=[
            {
                "index": 0,
                "element_id": element_id,
                "name": "保存更改",
                "control_type": "Button",
                "enabled": True,
                "bounds": {"left": 10, "top": 20, "width": 80, "height": 24},
            },
        ],
    )
    assert service._snapshot_element("snapshot", 42, None, element_id)["name"] == "保存更改"
    assert service._snapshot_element(
        "snapshot",
        42,
        None,
        element_text="保存",
        control_type="Button",
    )["element_id"] == element_id

    def fake_resolve(
        _hwnd, snapshot_id, _index, _x, _y, _space,
        _element_id, _element_text, _control_type,
    ):
        if snapshot_id == "expired":
            raise ComputerUseError("snapshot expired")
        return (50, 32)

    async def fake_get_state(*_args, **_kwargs):
        return {"snapshot_id": "refreshed"}

    monkeypatch.setattr(service, "_resolve_point_sync", fake_resolve)
    monkeypatch.setattr(service, "get_state", fake_get_state)
    point, snapshot_id, retried = await service._resolve_point_with_retry(
        42,
        "expired",
        None,
        None,
        None,
        "window",
        element_id=element_id,
        element_text=None,
        control_type="Button",
        auto_retry=True,
    )
    assert point == (50, 32)
    assert snapshot_id == "refreshed"
    assert retried is True
