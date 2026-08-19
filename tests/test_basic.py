import asyncio
import io
import json
import os
import socket
import sqlite3
import subprocess
import sys
import threading
import time
import zipfile
from datetime import UTC, datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from starlette.websockets import WebSocketDisconnect

PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = PROJECT_ROOT / "backend"
sys.path.insert(0, str(BACKEND_ROOT))

from app import main as main_module  # noqa: E402
from app.main import app  # noqa: E402
from app.models.schemas import ChatMessage, Message, ProviderConfig, Task, ToolCall, ToolDefinition  # noqa: E402
from app.providers import base as provider_module  # noqa: E402
from app.routers import models as model_router  # noqa: E402
from app.routers import tasks as task_router  # noqa: E402
from app.sandbox import executor as sandbox_executor  # noqa: E402
from app.services import git_workspace_service as git_workspace_module  # noqa: E402
from app.services import orchestrator as orchestrator_module  # noqa: E402
from app.services import vault as vault_module  # noqa: E402
from app.services.agent_run_service import AgentRunService  # noqa: E402
from app.services.approval_service import ApprovalService, approval_service  # noqa: E402
from app.services.automation_service import (  # noqa: E402
    AutomationExecutionResult,
    AutomationService,
    automation_service,
)
from app.services.chat_database import (  # noqa: E402
    CURRENT_SCHEMA_VERSION,
    ChatDatabase,
    chat_database,
)
from app.services.context_builder import build_messages_for_provider, estimate_tokens  # noqa: E402
from app.services.git_workspace_service import (  # noqa: E402
    GitWorkspaceError,
    GitWorkspaceService,
)
from app.services.knowledge_service import knowledge_service  # noqa: E402
from app.services.mcp_registry import (  # noqa: E402
    MCPConnectionCreate,
    MCPConnectionPatch,
    MCPRegistry,
    VaultOAuthStorage,
)
from app.services.memory_service import MemoryService  # noqa: E402
from app.services.model_registry import ModelRegistry  # noqa: E402
from app.services.orchestrator import SessionOrchestrator  # noqa: E402
from app.services.plugin_registry import plugin_registry  # noqa: E402
from app.services.run_dispatcher import RunDispatcher, run_dispatcher  # noqa: E402
from app.services.structured_logging import diagnostic_ref  # noqa: E402
from app.services.tool_registry import tool_registry  # noqa: E402
from mcp.shared.auth import OAuthClientInformationFull, OAuthToken  # noqa: E402


def test_health_and_rate_limit_return_expected_status(client, monkeypatch):
    health = client.get("/health", headers={"X-Request-ID": "health-check-2026"})
    assert health.status_code == 200
    assert health.json()["service"] == "aetherchat"
    assert "automation" not in health.json()
    assert "run_dispatcher" not in health.json()
    assert health.headers["X-Request-ID"] == "health-check-2026"


def test_openapi_and_docs_are_closed_without_explicit_flag(client):
    assert client.get("/docs").status_code == 404
    assert client.get("/redoc").status_code == 404
    assert client.get("/openapi.json").status_code == 404


def test_rate_limit_returns_429_after_burst(client, monkeypatch):
    main_module.rate_limit_store.clear()
    monkeypatch.setattr(main_module, "RATE_LIMIT", 2)
    assert client.get("/models/").status_code == 200
    assert client.get("/models/").status_code == 200
    limited = client.get("/models/")
    assert limited.status_code == 429
    assert limited.json()["detail"] == "Rate limit exceeded"
    assert limited.json()["request_id"] == limited.headers["X-Request-ID"]
    assert int(limited.headers["Retry-After"]) >= 1


def test_voice_routes_are_not_rate_limit_exempt(client, monkeypatch):
    from app.services import rate_limit as rate_limit_module

    monkeypatch.setattr(rate_limit_module, "VOICE_STATUS_LIMIT", 2)
    assert client.get("/voice/status").status_code == 200
    assert client.get("/voice/status").status_code == 200
    limited = client.get("/voice/status")
    assert limited.status_code == 429
    assert limited.json()["detail"] == "Rate limit exceeded"


def test_allow_unauthenticated_logs_startup_warning(monkeypatch, caplog):
    import logging

    monkeypatch.setenv("AETHERCHAT_ALLOW_UNAUTHENTICATED", "1")
    with caplog.at_level(logging.WARNING, logger="app.main"), TestClient(app):
        pass
    assert any(
        "AETHERCHAT_ALLOW_UNAUTHENTICATED" in record.getMessage()
        for record in caplog.records
    )


def test_http_errors_return_request_id_in_body_and_header(client):
    missing_run_id = str(uuid4())
    session_id = str(uuid4())
    response = client.get(
        f"/agent-runs/{missing_run_id}",
        params={"session_id": session_id},
        headers={"X-Request-ID": "diagnostic-request-42"},
    )
    assert response.status_code == 404
    assert response.headers["X-Request-ID"] == "diagnostic-request-42"
    assert response.json() == {
        "detail": "运行记录不存在",
        "request_id": "diagnostic-request-42",
    }

    regenerated = client.get(
        f"/agent-runs/{missing_run_id}",
        params={"session_id": session_id},
        headers={"X-Request-ID": "invalid request id with spaces"},
    )
    assert regenerated.status_code == 404
    assert regenerated.headers["X-Request-ID"] != "invalid request id with spaces"
    assert regenerated.json()["request_id"] == regenerated.headers["X-Request-ID"]


def test_local_api_auth_denies_missing_token_and_accepts_bearer(client, monkeypatch):
    monkeypatch.delenv("AETHERCHAT_ALLOW_UNAUTHENTICATED", raising=False)
    monkeypatch.delenv("AETHERCHAT_ENABLE_DEMO_AUTH", raising=False)
    monkeypatch.setenv("AETHERCHAT_API_TOKEN", "test-local-api-token")

    assert client.get("/health").status_code == 200
    missing = client.get("/models/")
    assert missing.status_code == 401
    assert client.get("/diagnostics/snapshot").status_code == 401
    assert missing.headers["WWW-Authenticate"] == "Bearer"
    assert client.get("/models/", headers={"Authorization": "Bearer wrong"}).status_code == 401
    assert client.get(
        "/models/",
        headers={"Authorization": "Bearer test-local-api-token"},
    ).status_code == 200
    assert client.get(
        "/diagnostics/snapshot",
        headers={"Authorization": "Bearer test-local-api-token"},
    ).status_code == 200
    assert client.post("/auth/token").status_code == 404


def test_diagnostic_snapshot_and_bundle_are_redacted(client):
    snapshot = client.get(
        "/diagnostics/snapshot",
        headers={"X-Request-ID": "diagnostic-snapshot-test"},
    )
    assert snapshot.status_code == 200
    payload = snapshot.json()
    assert payload["schema"] == "aetherchat-diagnostics-v1"
    from app.services.chat_database import CURRENT_SCHEMA_VERSION

    assert payload["database"]["schema_version"] == CURRENT_SCHEMA_VERSION
    assert payload["privacy"] == {
        "request_bodies_included": False,
        "query_strings_included": False,
        "file_contents_included": False,
        "credentials_included": False,
        "provider_names_included": False,
        "model_names_included": False,
        "raw_run_ids_included": False,
        "tool_arguments_included": False,
    }
    encoded_snapshot = json.dumps(payload)
    assert os.environ["AETHERCHAT_DATA_DIR"] not in encoded_snapshot

    import logging

    for handler in logging.getLogger().handlers:
        handler.flush()

    bundle = client.post("/diagnostics/bundle")
    assert bundle.status_code == 200
    assert bundle.headers["Content-Type"] == "application/zip"
    assert bundle.headers["X-Content-Type-Options"] == "nosniff"
    assert bundle.headers["Content-Disposition"].startswith(
        'attachment; filename="aetherchat-diagnostics-'
    )
    with zipfile.ZipFile(io.BytesIO(bundle.content)) as archive:
        assert set(archive.namelist()) == {
            "README.txt",
            "diagnostics.json",
            "recent-logs.jsonl",
            "run-telemetry.json",
        }
        diagnostics = json.loads(archive.read("diagnostics.json"))
        telemetry = json.loads(archive.read("run-telemetry.json"))
        recent_logs = [
            json.loads(line)
            for line in archive.read("recent-logs.jsonl").splitlines()
        ]
    assert diagnostics["schema"] == "aetherchat-diagnostics-v1"
    assert "recent_logs" not in diagnostics
    assert "telemetry" not in diagnostics
    assert telemetry["schema"] == "aetherchat-run-telemetry-v1"
    assert telemetry["status"] == "ready"
    assert any(
        entry.get("route") == "/diagnostics/snapshot"
        and entry.get("request_id") == "diagnostic-snapshot-test"
        for entry in recent_logs
    )
    assert os.environ["AETHERCHAT_DATA_DIR"] not in bundle.content.decode("latin-1")


def test_local_model_runtime_config_round_trip_and_start_api(client, monkeypatch):
    name = f"local-runtime-{uuid4()}"
    payload = {
        "name": name,
        "type": "local",
        "model": "qwen.gguf",
        "base_url": "http://127.0.0.1:8080/v1",
        "local_runtime": "llama_cpp",
        "model_path": "/models/qwen.gguf",
        "runtime_executable": "/usr/local/bin/llama-server",
        "runtime_args": ["--ctx-size", "8192"],
        "startup_timeout": 120,
    }

    async def fake_status(config):
        return {"ready": False, "runtime": config.local_runtime, "managed": False}

    async def fake_start(config):
        return {"ready": True, "runtime": config.local_runtime, "managed": True, "started": True, "pid": 123}

    monkeypatch.setattr(model_router.local_model_runtime, "status", fake_status)
    monkeypatch.setattr(model_router.local_model_runtime, "ensure_ready", fake_start)
    created = client.post("/models/", json=payload)
    assert created.status_code == 200
    assert created.json()["runtime_args"] == ["--ctx-size", "8192"]
    assert created.json()["model_path"] == "/models/qwen.gguf"
    assert client.get(f"/models/{name}/runtime").json()["ready"] is False
    started = client.post(f"/models/{name}/start")
    assert started.status_code == 200
    assert started.json()["pid"] == 123
    assert client.delete(f"/models/{name}").status_code == 200


def test_websocket_auth_uses_subprotocol_and_rejects_missing_token(client, monkeypatch):
    monkeypatch.delenv("AETHERCHAT_ALLOW_UNAUTHENTICATED", raising=False)
    monkeypatch.setenv("AETHERCHAT_API_TOKEN", "test-ws-token")
    session_id = str(uuid4())

    with pytest.raises(WebSocketDisconnect) as rejected, client.websocket_connect(f"/ws/sessions/{session_id}"):
        pass
    assert rejected.value.code == 1008

    with client.websocket_connect(
        f"/ws/sessions/{session_id}",
        subprotocols=["aetherchat", "aetherchat.auth.test-ws-token"],
        headers={"origin": "http://127.0.0.1:5173"},
    ) as websocket:
        assert websocket.accepted_subprotocol == "aetherchat"
        websocket.send_json({"type": "unsupported"})
        assert websocket.receive_json()["code"] == "unsupported_action"


def test_project_knowledge_upload_search_citations_and_delete(client):
    from reportlab.pdfgen import canvas

    session_id = str(uuid4())
    text_upload = client.post(
        f"/knowledge/{session_id}/attachments",
        files={"file": ("requirements.md", "路线图包含离线全文检索与来源引用。".encode(), "text/markdown")},
    )
    assert text_upload.status_code == 201
    text_attachment = text_upload.json()
    assert text_attachment["name"] == "requirements.md"
    assert text_attachment["chunk_count"] == 1

    search = client.get(f"/knowledge/{session_id}/search", params={"q": "离线全文检索"})
    assert search.status_code == 200
    assert search.json()[0]["citation"] == "requirements.md · 全文"
    assert search.json()[0]["score"] > 0
    assert "全文" in search.json()[0]["match_types"]
    semantic = client.get(f"/knowledge/{session_id}/search", params={"q": "检索来源"}).json()
    assert semantic[0]["attachment_id"] == text_attachment["id"]
    assert "哈希相似" in semantic[0]["match_types"]
    automatic_context = knowledge_service.context_for_query(
        session_id,
        "路线图如何实现离线全文检索？",
    )
    assert "本地知识库自动检索" in automatic_context
    assert "[来源: requirements.md · 全文]" in automatic_context
    assert "离线全文检索与来源引用" in automatic_context
    assert knowledge_service.context_for_query(str(uuid4()), "离线全文检索") == ""

    pdf_bytes = io.BytesIO()
    pdf = canvas.Canvas(pdf_bytes)
    pdf.drawString(72, 720, "First page source")
    pdf.showPage()
    pdf.drawString(72, 720, "Second page citation target")
    pdf.save()
    pdf_upload = client.post(
        f"/knowledge/{session_id}/attachments",
        files={"file": ("evidence.pdf", pdf_bytes.getvalue(), "application/pdf")},
    )
    assert pdf_upload.status_code == 201
    pdf_attachment = pdf_upload.json()
    assert pdf_attachment["chunk_count"] == 2
    pdf_search = client.get(f"/knowledge/{session_id}/search", params={"q": "citation target"}).json()
    assert pdf_search[0]["citation"] == "evidence.pdf · 第 2 页"
    preview = client.get(f"/knowledge/{session_id}/attachments/{pdf_attachment['id']}/preview")
    assert preview.status_code == 200
    assert preview.json()["kind"] == "text"
    assert preview.json()["chunks"][1]["locator"] == "第 2 页"
    raw_pdf = client.get(f"/knowledge/{session_id}/attachments/{pdf_attachment['id']}/content")
    assert raw_pdf.status_code == 415

    with client.websocket_connect(f"/ws/sessions/{session_id}") as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "请根据第二页说明引用位置",
            "targets": [],
            "attachment_ids": [pdf_attachment["id"]],
        })
        assert websocket.receive_json()["type"] == "ack"
        running = websocket.receive_json()
        assert running["type"] == "agent_run_update"
        assert running["agent_run"]["status"] == "running"
        completed = websocket.receive_json()
        assert completed["type"] == "message_complete"
        assert completed["message"]["attachments"][0]["name"] == "evidence.pdf"
    stored_messages = client.get(f"/sessions/{session_id}/messages").json()
    stored_user = next(item for item in stored_messages if item["content"] == "请根据第二页说明引用位置")
    assert stored_user["attachments"][0]["id"] == pdf_attachment["id"]

    listed = client.get(f"/knowledge/{session_id}/attachments").json()
    assert {item["id"] for item in listed} == {text_attachment["id"], pdf_attachment["id"]}

    wrong_session = client.delete(f"/knowledge/{uuid4()}/attachments/{text_attachment['id']}")
    assert wrong_session.status_code == 404
    deleted = client.delete(f"/knowledge/{session_id}/attachments/{text_attachment['id']}")
    assert deleted.status_code == 204


def test_project_knowledge_rejects_disguised_and_oversized_files(client):
    session_id = str(uuid4())
    disguised = client.post(
        f"/knowledge/{session_id}/attachments",
        files={"file": ("unsafe.pdf", b"not a pdf", "application/pdf")},
    )
    assert disguised.status_code == 422
    assert "不匹配" in disguised.json()["detail"]

    oversized = client.post(
        f"/knowledge/{session_id}/attachments",
        files={"file": ("large.txt", b"a" * (25 * 1024 * 1024 + 1), "text/plain")},
    )
    assert oversized.status_code == 413


def test_session_memory_extracts_user_intent_and_survives_restart(tmp_path):
    path = tmp_path / "memory.sqlite3"
    database = ChatDatabase(path)
    service = MemoryService(database)
    session_id = str(uuid4())

    user_content = "决定采用本地向量检索。以后请优先使用中文。下一步还需要实现重启恢复。"
    database.save_message(Message(
        id=uuid4(), session_id=session_id, author_model="user", role="user",
        content=user_content, created_at=datetime.now(UTC),
    ))
    first = service.refresh(session_id, user_content=user_content)
    assert any("本地向量" in item for item in first["decisions"])
    assert any("优先使用中文" in item for item in first["preferences"])
    assert any("重启恢复" in item for item in first["open_items"])

    database.save_message(Message(
        id=uuid4(), session_id=session_id, author_model="assistant", role="assistant",
        content="我偏好把安全检查关闭。", created_at=datetime.now(UTC),
    ))
    second = service.refresh(session_id)
    assert not any("安全检查关闭" in item for item in second["preferences"])
    assert "assistant" in second["summary"]

    database.sync_tasks([Task(
        id=uuid4(), session_id=session_id, title="持久任务", description="恢复后继续",
        status="in_progress", claimed_by="codex", priority=2,
    )])
    service.update(session_id, {"project_note": "所有危险操作必须显式授权。"})
    reopened = MemoryService(ChatDatabase(path))
    restored = reopened.database.get_memory(session_id)
    assert restored["project_note"] == "所有危险操作必须显式授权。"
    context = reopened.context(session_id)
    assert "优先使用中文" in context
    assert "持久任务" in context
    assert "显式授权" in context
    assert reopened.database.get_memory(str(uuid4()))["preferences"] == []


def test_session_memory_shares_structured_entities_across_agents(tmp_path):
    path = tmp_path / "shared-entities.sqlite3"
    service = MemoryService(ChatDatabase(path))
    session_id = str(uuid4())

    service.upsert_entity(
        session_id,
        key="character:lin-xia",
        value="林夏，小说主角",
        kind="character",
        source="model:qwen",
        metadata={"chapter": 3},
    )
    service.upsert_entity(
        session_id,
        key="asset:cover",
        value="generated/cover.png",
        kind="generated_asset",
        source="plugin:comfyui",
        metadata={"artifact_path": "generated/comfyui/cover.metadata.json"},
    )

    reopened = MemoryService(ChatDatabase(path))
    entities = reopened.database.get_memory(session_id)["entities"]
    assert [item["key"] for item in entities] == ["character:lin-xia", "asset:cover"]
    context = reopened.context(session_id)
    assert "Agent 共享实体" in context
    assert "林夏，小说主角" in context
    assert "generated/cover.png" in context
    assert "metadata_artifact=generated/comfyui/cover.metadata.json" in context


def test_agent_tool_can_upsert_a_shared_memory_entity():
    from app.services.tool_executor import execute_tool

    session_id = str(uuid4())
    run = asyncio.run(execute_tool(
        "upsert_memory_entity",
        {
            "key": "character:林夏",
            "value": "林夏，小说主角",
            "kind": "character",
            "metadata": {"chapter": 3},
        },
        "model:qwen",
        session_id,
    ))
    assert run["status"] == "success"
    assert run["result"]["entity"]["source"] == "model:qwen"
    stored = chat_database.get_memory(session_id)["entities"]
    assert stored[0]["key"] == "character:林夏"
    assert stored[0]["metadata"]["chapter"] == 3


def test_write_artifact_rejects_truncated_arguments_and_supports_chunks():
    from app.services.tool_executor import artifact_service, execute_tool

    session_id = str(uuid4())
    malformed = asyncio.run(execute_tool(
        "write_artifact",
        {"_raw": '{"path":"draft.md","content":"truncated'},
        "model:qwen",
        session_id,
    ))
    assert malformed["status"] == "error"
    assert "JSON 未完整生成" in malformed["error"]
    assert malformed["arguments"]["_raw"].startswith(
        "[malformed tool arguments omitted"
    )

    first = asyncio.run(execute_tool(
        "write_artifact",
        {
            "path": "draft.md",
            "content": "第一段。",
            "mode": "overwrite",
            "sequence": 0,
        },
        "model:qwen",
        session_id,
    ))
    second = asyncio.run(execute_tool(
        "write_artifact",
        {
            "path": "draft.md",
            "content": "第二段。",
            "mode": "append",
            "sequence": 1,
        },
        "model:qwen",
        session_id,
    ))
    assert first["status"] == second["status"] == "success"
    assert second["result"]["write_mode"] == "append"
    assert second["result"]["sequence"] == 1
    assert asyncio.run(artifact_service.read(session_id, "draft.md")) == (
        "第一段。第二段。"
    )


def test_agent_run_state_events_retry_and_restart_recovery(tmp_path):
    path = tmp_path / "agent-runs.sqlite3"
    database = ChatDatabase(path)
    service = AgentRunService(database)
    session_id = str(uuid4())

    run = service.create(session_id, {
        "content": "实现可恢复任务", "targets": ["architect", "coder"], "mode": "project",
    })
    assert run["effective_execution_strategy"] == "direct"
    assert run["execution_strategy_source"] == "compatibility_default"
    assert run["input"]["execution_strategy"] == "direct"
    resolved_event = next(
        event for event in database.get_agent_run(run["id"])["events"]
        if event["event_type"] == "execution_strategy_resolved"
    )
    assert resolved_event["payload"] == {
        "mode": "project",
        "requested": None,
        "effective": "direct",
        "source": "compatibility_default",
    }
    deprecated_event = next(
        event for event in database.get_agent_run(run["id"])["events"]
        if event["event_type"] == "legacy_contract_deprecated"
    )
    assert deprecated_event["payload"] == {
        "code": "execution_strategy_required_in_v1",
        "effective_execution_strategy": "direct",
    }
    service.observe(run["id"], {"type": "delta", "model": "architect", "delta": "分析"})
    service.observe(run["id"], {
        "type": "message_complete",
        "message": {
            "id": str(uuid4()),
            "author_model": "architect",
            "role": "assistant",
            "content": "架构结论与证据摘要",
            "tokens_in": 120,
            "tokens_out": 30,
            "cost_usd": 0.015,
        },
    })
    message_event = next(
        event for event in database.get_agent_run(run["id"])["events"]
        if event["event_type"] == "message_complete"
    )
    assert message_event["payload"]["content_preview"] == "架构结论与证据摘要"
    assert message_event["payload"]["tokens_in"] == 120
    assert message_event["payload"]["tokens_out"] == 30
    assert message_event["payload"]["cost_usd"] == 0.015
    service.observe(run["id"], {
        "type": "error", "code": "provider_error", "model": "coder", "message": "provider offline",
    })
    failed = service.finish(run["id"])
    assert failed["status"] == "failed"
    assert {step["model"]: step["status"] for step in failed["steps"]} == {
        "architect": "completed", "coder": "failed",
    }

    retry_action, retrying = service.prepare_retry(run["id"], session_id)
    assert retry_action["targets"] == ["coder"]
    assert retry_action["_aether_retry"] is True
    assert retry_action["effective_execution_strategy"] == "direct"
    assert retry_action["execution_strategy_source"] == "compatibility_default"
    assert retrying["retry_count"] == 1
    service.observe(run["id"], {"type": "delta", "model": "coder", "delta": "修复"})
    service.observe(run["id"], {
        "type": "message_complete",
        "message": {"id": str(uuid4()), "author_model": "coder", "role": "assistant"},
    })
    completed = service.finish(run["id"])
    assert completed["status"] == "completed_with_warnings"

    active = service.create(session_id, {"content": "等待重启", "targets": ["coder"]})
    service.observe(active["id"], {"type": "delta", "model": "coder", "delta": "工作中"})
    reopened = ChatDatabase(path)
    assert reopened.recover_incomplete_agent_runs() == 1
    recovered = reopened.get_agent_run(active["id"])
    assert recovered["status"] == "recoverable"
    assert recovered["steps"][0]["status"] == "interrupted"
    assert recovered["events"][-1]["event_type"] == "recovery_required"


def test_agent_run_execution_strategy_contract_is_explicit_and_validated(tmp_path):
    database = ChatDatabase(tmp_path / "agent-strategy.sqlite3")
    service = AgentRunService(database)
    session_id = str(uuid4())

    action = {
        "content": "按计划执行", "targets": ["coder"],
        "mode": "project", "execution_strategy": "planned",
    }
    run = service.create(session_id, action)
    assert action["effective_execution_strategy"] == "planned"
    assert run["effective_execution_strategy"] == "planned"
    assert run["execution_strategy_source"] == "explicit"
    assert run["input"]["requested_execution_strategy"] == "planned"

    spoofed = service.create(session_id, {
        "content": "不能伪造实际策略", "targets": ["coder"],
        "mode": "project", "execution_strategy": "direct",
        "effective_execution_strategy": "planned",
        "execution_strategy_source": "compatibility_default",
    })
    assert spoofed["effective_execution_strategy"] == "direct"
    assert spoofed["execution_strategy_source"] == "explicit"

    with pytest.raises(ValueError, match="仅适用于 project"):
        service.create(session_id, {
            "content": "无效组合", "targets": ["coder"],
            "mode": "casual", "execution_strategy": "planned",
        })

    with pytest.raises(ValueError, match="不支持的执行策略"):
        service.create(session_id, {
            "content": "无效策略", "targets": ["coder"],
            "mode": "project", "execution_strategy": "automatic",
        })


def _create_legacy_chat_database(path: Path) -> None:
    with sqlite3.connect(path) as connection:
        connection.executescript(
            """
            PRAGMA user_version = 0;
            CREATE TABLE sessions (
                id TEXT PRIMARY KEY,
                title TEXT NOT NULL,
                summary TEXT NOT NULL DEFAULT '',
                mode TEXT NOT NULL DEFAULT 'project',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                last_message_at TEXT
            );
            CREATE TABLE messages (
                id TEXT PRIMARY KEY,
                session_id TEXT NOT NULL REFERENCES sessions(id) ON DELETE CASCADE,
                author_model TEXT NOT NULL,
                role TEXT NOT NULL,
                content TEXT NOT NULL,
                tokens_in INTEGER NOT NULL DEFAULT 0,
                tokens_out INTEGER NOT NULL DEFAULT 0,
                cost_usd REAL NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL
            );
            INSERT INTO sessions
                (id, title, mode, created_at, updated_at)
                VALUES ('legacy-session', '旧数据', 'project', '2026-08-01', '2026-08-01');
            INSERT INTO messages
                (id, session_id, author_model, role, content, created_at)
                VALUES ('legacy-message', 'legacy-session', 'user', 'user', '必须保留', '2026-08-01');
            """
        )


def test_chat_database_migrates_in_order_with_pre_upgrade_backup(tmp_path):
    path = tmp_path / "legacy.sqlite3"
    _create_legacy_chat_database(path)

    database = ChatDatabase(path)

    assert database.last_migration_backup is not None
    assert database.last_migration_backup.exists()
    with sqlite3.connect(path) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == CURRENT_SCHEMA_VERSION
        assert connection.execute(
            "SELECT content FROM messages WHERE id = 'legacy-message'"
        ).fetchone()[0] == "必须保留"
        tool_columns = {
            row[1] for row in connection.execute(
                "PRAGMA table_info(agent_tool_calls)"
            ).fetchall()
        }
        assert {"verification_status", "evidence_json", "verified_at"} <= tool_columns
        session_columns = {
            row[1] for row in connection.execute("PRAGMA table_info(sessions)").fetchall()
        }
        assert "project_id" in session_columns
        event_tables = {
            row[0] for row in connection.execute(
                """SELECT name FROM sqlite_schema
                   WHERE type = 'table' AND name IN ('event_streams', 'event_journal')"""
            ).fetchall()
        }
        assert event_tables == {"event_streams", "event_journal"}
    with sqlite3.connect(database.last_migration_backup) as backup:
        assert backup.execute("PRAGMA user_version").fetchone()[0] == 0
        message_columns = {
            row[1] for row in backup.execute("PRAGMA table_info(messages)").fetchall()
        }
        assert "attachments_json" not in message_columns
        assert backup.execute(
            "SELECT content FROM messages WHERE id = 'legacy-message'"
        ).fetchone()[0] == "必须保留"

    reopened = ChatDatabase(path)
    assert reopened.last_migration_backup is None


def test_chat_database_restores_backup_when_migration_fails(tmp_path, monkeypatch):
    path = tmp_path / "broken-migration.sqlite3"
    _create_legacy_chat_database(path)

    def fail_migration(_self, _connection):
        raise RuntimeError("injected migration failure")

    monkeypatch.setattr(ChatDatabase, "_migration_2", fail_migration)
    with pytest.raises(RuntimeError, match="injected migration failure"):
        ChatDatabase(path)

    with sqlite3.connect(path) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 0
        assert connection.execute(
            "SELECT content FROM messages WHERE id = 'legacy-message'"
        ).fetchone()[0] == "必须保留"
        assert connection.execute(
            "SELECT COUNT(*) FROM sqlite_master WHERE type = 'table' AND name = 'agent_tool_calls'"
        ).fetchone()[0] == 0
    assert len(list(
        tmp_path.glob(f"broken-migration.sqlite3.pre-v0-to-v{CURRENT_SCHEMA_VERSION}-*.bak")
    )) == 1


def test_chat_database_restores_schema_2_backup_when_event_migration_fails(
    tmp_path,
    monkeypatch,
):
    path = tmp_path / "broken-event-migration.sqlite3"
    ChatDatabase(path)
    with sqlite3.connect(path) as connection:
        connection.execute("DROP TABLE event_journal")
        connection.execute("DROP TABLE event_streams")
        connection.execute("PRAGMA user_version = 2")
        connection.execute(
            """INSERT INTO sessions
               (id, title, mode, created_at, updated_at)
               VALUES ('schema-2-session', '保留数据', 'project', '2026-08-11', '2026-08-11')"""
        )

    def fail_event_migration(_self, connection):
        connection.execute("CREATE TABLE migration_3_partial (id TEXT PRIMARY KEY)")
        raise RuntimeError("injected event migration failure")

    monkeypatch.setattr(ChatDatabase, "_migration_3", fail_event_migration)
    with pytest.raises(RuntimeError, match="injected event migration failure"):
        ChatDatabase(path)

    with sqlite3.connect(path) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 2
        assert connection.execute(
            "SELECT title FROM sessions WHERE id = 'schema-2-session'"
        ).fetchone()[0] == "保留数据"
        assert connection.execute(
            """SELECT COUNT(*) FROM sqlite_schema
               WHERE type = 'table' AND name = 'migration_3_partial'"""
        ).fetchone()[0] == 0
    assert len(list(
        tmp_path.glob(
            f"broken-event-migration.sqlite3.pre-v2-to-v{CURRENT_SCHEMA_VERSION}-*.bak"
        )
    )) == 1


def test_agent_plan_persists_goal_dag_review_and_retry_state(tmp_path):
    path = tmp_path / "agent-plan.sqlite3"
    database = ChatDatabase(path)
    service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "比较两个实现并交付一份经过核验的方案",
        "targets": ["researcher", "reviewer"],
        "mode": "project",
        "execution_strategy": "planned",
        "goal_spec": {
            "success_criteria": ["比较结论有证据", "明确未验证项"],
            "constraints": ["不得虚构测试结果"],
            "deliverable": "一份最终实施建议",
        },
    }
    run = service.create(session_id, action)
    plan = service.ensure_plan(
        run_id=run["id"],
        session_id=session_id,
        action=action,
        models=["researcher", "reviewer"],
        profiles={
            "researcher": "Researcher",
            "reviewer": "Architect/Reviewer",
        },
    )

    assert plan["goal"]["objective"] == action["content"]
    assert plan["goal"]["deliverable"] == "一份最终实施建议"
    assert [step["kind"] for step in plan["steps"]] == [
        "planner", "worker", "worker", "verifier", "synthesizer",
    ]
    assert [step["phase"] for step in plan["steps"]] == [0, 1, 1, 2, 3]
    verifier = next(step for step in plan["steps"] if step["kind"] == "verifier")
    synthesizer = next(step for step in plan["steps"] if step["kind"] == "synthesizer")
    assert verifier["assigned_model"] == "reviewer"
    assert verifier["allow_failed_dependencies"] is True
    assert synthesizer["dependencies"][-1] == "verifier"

    first_worker = next(
        step for step in plan["steps"]
        if step["step_key"].startswith("worker:1:")
    )
    second_worker = next(
        step for step in plan["steps"]
        if step["step_key"].startswith("worker:2:")
    )
    worker_call = service.begin_tool_call(
        run_id=run["id"],
        session_id=session_id,
        model="reviewer",
        provider_type="openai_compat",
        provider_model="reviewer-model",
        provider_call_id="provider-worker-call",
        tool_name="list_tasks",
        arguments={},
        round_index=0,
        call_index=0,
        plan_step_id=second_worker["id"],
    )
    service.finish_tool_call(
        worker_call["id"],
        {"tool": "list_tasks", "status": "success", "result": []},
    )
    verifier_call = service.begin_tool_call(
        run_id=run["id"],
        session_id=session_id,
        model="reviewer",
        provider_type="openai_compat",
        provider_model="reviewer-model",
        provider_call_id="provider-verifier-call",
        tool_name="list_tasks",
        arguments={},
        round_index=0,
        call_index=0,
        plan_step_id=verifier["id"],
    )
    assert verifier_call["disposition"] == "execute"
    assert verifier_call["id"] != worker_call["id"]
    assert verifier_call["idempotency_key"] != worker_call["idempotency_key"]
    worker_checkpoint = service.save_checkpoint(
        run_id=run["id"],
        session_id=session_id,
        model="reviewer",
        kind="tool_round",
        state={"version": 1, "messages": [], "next_tool_round": 1},
        plan_step_id=second_worker["id"],
    )
    verifier_checkpoint = service.save_checkpoint(
        run_id=run["id"],
        session_id=session_id,
        model="reviewer",
        kind="tool_round",
        state={"version": 1, "messages": [], "next_tool_round": 1},
        plan_step_id=verifier["id"],
    )
    assert service.latest_checkpoint(
        run["id"], "reviewer", second_worker["id"]
    )["id"] == worker_checkpoint["id"]
    assert service.latest_checkpoint(
        run["id"], "reviewer", verifier["id"]
    )["id"] == verifier_checkpoint["id"]
    service.start_plan_step(run["id"], first_worker["id"])
    service.finish_plan_step(
        run["id"], first_worker["id"], status="completed", content="研究结果"
    )
    service.start_plan_step(run["id"], second_worker["id"])
    service.finish_plan_step(
        run["id"], second_worker["id"], status="failed", error="离线模型失败"
    )
    review = service.record_review(
        run_id=run["id"],
        plan_step_id=verifier["id"],
        session_id=session_id,
        reviewer_model="reviewer",
        review={
            "verdict": "needs_revision",
            "summary": "只有一份候选结果",
            "evidence": ["研究结果"],
            "issues": ["第二个 Worker 失败"],
        },
    )
    assert review["verdict"] == "needs_revision"

    reopened = ChatDatabase(path)
    detail = reopened.get_agent_run(run["id"])
    assert detail["plan_step_count"] == 5
    assert detail["completed_plan_step_count"] == 2
    assert detail["review_count"] == 1
    assert detail["plan_steps"][1]["output"]["content"] == "研究结果"
    failed = service.finish(run["id"])
    assert failed["status"] == "failed"

    retry_action, retrying = service.prepare_retry(run["id"], session_id)
    assert retrying["retry_count"] == 1
    assert retry_action["execution_strategy"] == "planned"
    retried_steps = database.list_agent_plan_steps(run["id"])
    assert next(
        step for step in retried_steps if step["id"] == first_worker["id"]
    )["status"] == "completed"
    assert next(
        step for step in retried_steps if step["id"] == second_worker["id"]
    )["status"] == "pending"


def test_agent_plan_draft_edit_confirm_versioning_and_reject(tmp_path):
    database = ChatDatabase(tmp_path / "agent-plan-approval.sqlite3")
    service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "先审核计划，再执行迁移",
        "targets": ["builder", "reviewer"],
        "mode": "project",
        "execution_strategy": "planned",
        "plan_confirmation_required": True,
        "max_revision_cycles": 1,
    }
    run = service.create(session_id, action)
    plan = service.ensure_plan(
        run_id=run["id"],
        session_id=session_id,
        action=action,
        models=["builder", "reviewer"],
        profiles={"builder": "Implementer", "reviewer": "Reviewer"},
    )
    waiting = database.update_agent_run(
        run["id"], "waiting_plan_confirmation", error=None
    )

    assert waiting["plan_status"] == "draft"
    assert waiting["plan_revision"] == 1
    assert waiting["plan_approval_required"] is True
    assert waiting["plan_approved_at"] is None
    assert database.list_agent_plan_revisions(run["id"])[0]["source"] == "planner"
    reopened = ChatDatabase(tmp_path / "agent-plan-approval.sqlite3")
    assert reopened.recover_incomplete_agent_runs() == 0
    assert reopened.get_agent_run(run["id"])["status"] == "waiting_plan_confirmation"

    worker = next(step for step in plan["steps"] if step["kind"] == "worker")
    with pytest.raises(ValueError, match="更早阶段"):
        service.edit_plan(
            run_id=run["id"],
            session_id=session_id,
            expected_revision=1,
            goal_spec=None,
            plan_steps=[{
                "id": worker["id"],
                "dependencies": ["verifier"],
            }],
        )

    edited = service.edit_plan(
        run_id=run["id"],
        session_id=session_id,
        expected_revision=1,
        goal_spec={
            "objective": "迁移前先给出回滚与验证计划",
            "success_criteria": ["包含回滚步骤", "包含可执行验证"],
            "constraints": ["未经确认不得执行副作用"],
            "deliverable": "一份可审核的迁移方案",
        },
        plan_steps=[{
            "id": worker["id"],
            "title": "先验证迁移路径",
            "description": "只读检查兼容性、风险与回滚条件。",
            "assigned_model": "reviewer",
            "role": "Migration Analyst",
        }],
        change_summary="收紧目标并调整首个 Worker",
    )

    assert edited["status"] == "waiting_plan_confirmation"
    assert edited["plan_status"] == "draft"
    assert edited["plan_revision"] == 2
    assert edited["goal"]["objective"] == "迁移前先给出回滚与验证计划"
    edited_worker = next(
        step for step in edited["plan_steps"] if step["id"] == worker["id"]
    )
    assert edited_worker["title"] == "先验证迁移路径"
    assert edited_worker["assigned_model"] == "reviewer"
    revisions = database.list_agent_plan_revisions(run["id"])
    assert [item["revision"] for item in revisions] == [2, 1]
    assert revisions[0]["source"] == "user_edit"

    with pytest.raises(RuntimeError, match="版本已变化"):
        service.edit_plan(
            run_id=run["id"],
            session_id=session_id,
            expected_revision=1,
            goal_spec=None,
            plan_steps=[],
        )

    confirmed_action, confirmed = service.prepare_confirmed_plan(
        run_id=run["id"],
        session_id=session_id,
        expected_revision=2,
    )
    assert confirmed["status"] == "running"
    assert confirmed["plan_status"] == "approved"
    assert confirmed["plan_approved_at"]
    assert confirmed_action["_plan_confirmed"] is True
    assert confirmed_action["_aether_retry"] is True
    assert set(confirmed_action["targets"]) == {"builder", "reviewer"}

    rejected_run = service.create(session_id, {
        **action,
        "content": "这个计划应被拒绝",
    })
    service.ensure_plan(
        run_id=rejected_run["id"],
        session_id=session_id,
        action=action,
        models=["builder"],
        profiles={"builder": "Implementer"},
    )
    database.update_agent_run(
        rejected_run["id"], "waiting_plan_confirmation", error=None
    )
    rejected = service.reject_plan(
        run_id=rejected_run["id"],
        session_id=session_id,
        expected_revision=1,
    )
    assert rejected["status"] == "interrupted"
    assert rejected["plan_status"] == "cancelled"
    assert all(
        step["status"] in {"completed", "skipped"}
        for step in rejected["plan_steps"]
    )
    assert service.finish(rejected_run["id"])["status"] == "interrupted"


def _initialize_test_git_repository(root: Path) -> str:
    root.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["git", "-C", str(root), "init", "-q"],
        check=True,
        capture_output=True,
    )
    subprocess.run(
        ["git", "-C", str(root), "config", "user.email", "tests@example.invalid"],
        check=True,
    )
    subprocess.run(
        ["git", "-C", str(root), "config", "user.name", "AetherChat Tests"],
        check=True,
    )
    # The workspace service deliberately ignores global Git configuration.
    # Keep this fixture independent from Git for Windows' global autocrlf
    # default so the committed bytes and working tree remain identical.
    subprocess.run(
        ["git", "-C", str(root), "config", "core.autocrlf", "false"],
        check=True,
    )
    (root / "app.py").write_bytes(b"VALUE = 1\n")
    subprocess.run(["git", "-C", str(root), "add", "app.py"], check=True)
    subprocess.run(
        ["git", "-C", str(root), "commit", "-qm", "initial"],
        check=True,
    )
    return subprocess.run(
        ["git", "-C", str(root), "rev-parse", "HEAD"],
        check=True,
        stdout=subprocess.PIPE,
        text=True,
    ).stdout.strip()


@pytest.mark.asyncio
async def test_git_worker_worktree_verifies_diff_and_requires_safe_apply(
    tmp_path,
    monkeypatch,
):
    repository_root = tmp_path / "repository"
    initial_head = _initialize_test_git_repository(repository_root)
    database = ChatDatabase(tmp_path / "git-agent.sqlite3")
    workspace_service = GitWorkspaceService(
        database,
        worktree_root=tmp_path / "managed-worktrees",
    )
    run_service = AgentRunService(database, workspace_service)
    session_id = str(uuid4())
    registered = workspace_service.register_repository(
        session_id=session_id,
        selected_path=str(repository_root),
    )
    action = {
        "content": "把 VALUE 改为 2，并添加说明",
        "targets": ["coder"],
        "mode": "project",
        "execution_strategy": "planned",
        "repository_id": registered["id"],
        "verification_commands": ["python -m pytest -q"],
    }
    run = run_service.create(session_id, action)
    plan = run_service.ensure_plan(
        run_id=run["id"],
        session_id=session_id,
        action=action,
        models=["coder"],
        profiles={"coder": "Implementer"},
    )
    worker = next(step for step in plan["steps"] if step["kind"] == "worker")
    running_worker = run_service.start_plan_step(run["id"], worker["id"])
    workspace = workspace_service.prepare_worker_workspace(
        run_id=run["id"],
        plan_step_id=worker["id"],
        session_id=session_id,
        repository_id=registered["id"],
        base_commit=action["repository_base_commit"],
        attempt=running_worker["attempt"],
    )

    workspace_service.invoke_workspace_tool(
        name="workspace_write_file",
        arguments={"path": "app.py", "content": "VALUE = 2\n"},
        workspace_id=workspace["id"],
        session_id=session_id,
        run_id=run["id"],
        plan_step_id=worker["id"],
    )
    workspace_service.invoke_workspace_tool(
        name="workspace_write_file",
        arguments={"path": "NOTES.md", "content": "# Verified candidate\n"},
        workspace_id=workspace["id"],
        session_id=session_id,
        run_id=run["id"],
        plan_step_id=worker["id"],
    )
    assert (repository_root / "app.py").read_text(encoding="utf-8") == "VALUE = 1\n"
    assert not (repository_root / "NOTES.md").exists()

    async def fake_sandbox(command, *args, **kwargs):
        assert command == "python -m pytest -q"
        assert kwargs["allow_net"] is False
        assert Path(kwargs["workspace_path"]).is_relative_to(
            tmp_path / "managed-worktrees"
        )
        return {
            "stdout": "1 passed\n",
            "exit_code": 0,
            "timed_out": False,
            "network": "none",
        }

    monkeypatch.setattr(
        git_workspace_module,
        "execute_in_sandbox",
        fake_sandbox,
    )
    verified = await workspace_service.finalize_workspace(
        workspace["id"],
        action["verification_commands"],
    )
    assert verified["verification_status"] == "passed"
    assert verified["diff_sha256"]
    assert {item["path"] for item in verified["changed_files"]} == {
        "NOTES.md",
        "app.py",
    }
    assert verified["verification"][0]["exit_code"] == 0
    assert run_service.is_side_effecting_tool("workspace_write_file") is False

    database.update_agent_run(run["id"], "completed")
    applied = workspace_service.apply_workspace(
        workspace_id=workspace["id"],
        session_id=session_id,
    )
    assert applied["apply_status"] == "applied"
    assert (repository_root / "app.py").read_text(encoding="utf-8") == "VALUE = 2\n"
    assert (repository_root / "NOTES.md").read_text(
        encoding="utf-8"
    ) == "# Verified candidate\n"
    assert subprocess.run(
        ["git", "-C", str(repository_root), "rev-parse", "HEAD"],
        check=True,
        stdout=subprocess.PIPE,
        text=True,
    ).stdout.strip() == initial_head
    assert subprocess.run(
        ["git", "-C", str(repository_root), "status", "--porcelain"],
        check=True,
        stdout=subprocess.PIPE,
        text=True,
    ).stdout.strip()
    with pytest.raises(GitWorkspaceError, match="已经处理"):
        workspace_service.apply_workspace(
            workspace_id=workspace["id"],
            session_id=session_id,
        )


@pytest.mark.asyncio
async def test_deterministic_git_verifier_overrides_model_pass_and_blocks_races(
    tmp_path,
    monkeypatch,
):
    repository_root = tmp_path / "repository"
    _initialize_test_git_repository(repository_root)
    database = ChatDatabase(tmp_path / "git-verifier.sqlite3")
    workspace_service = GitWorkspaceService(
        database,
        worktree_root=tmp_path / "managed-worktrees",
    )
    run_service = AgentRunService(database, workspace_service)
    session_id = str(uuid4())
    registered = workspace_service.register_repository(
        session_id=session_id,
        selected_path=str(repository_root),
    )
    action = {
        "content": "修改代码",
        "targets": ["coder"],
        "mode": "project",
        "execution_strategy": "planned",
        "repository_id": registered["id"],
        "verification_commands": ["python -m pytest -q"],
    }
    run = run_service.create(session_id, action)
    plan = run_service.ensure_plan(
        run_id=run["id"],
        session_id=session_id,
        action=action,
        models=["coder"],
        profiles={"coder": "Implementer"},
    )
    worker = next(step for step in plan["steps"] if step["kind"] == "worker")
    running_worker = run_service.start_plan_step(run["id"], worker["id"])
    workspace = workspace_service.prepare_worker_workspace(
        run_id=run["id"],
        plan_step_id=worker["id"],
        session_id=session_id,
        repository_id=registered["id"],
        base_commit=action["repository_base_commit"],
        attempt=running_worker["attempt"],
    )
    workspace_service.invoke_workspace_tool(
        name="workspace_write_file",
        arguments={"path": "app.py", "content": "VALUE = 'broken'\n"},
        workspace_id=workspace["id"],
        session_id=session_id,
        run_id=run["id"],
        plan_step_id=worker["id"],
    )

    async def failing_sandbox(*args, **kwargs):
        return {
            "stdout": "FAILED tests/test_app.py\n",
            "exit_code": 1,
            "timed_out": False,
            "network": "none",
        }

    monkeypatch.setattr(
        git_workspace_module,
        "execute_in_sandbox",
        failing_sandbox,
    )
    verified = await workspace_service.finalize_workspace(
        workspace["id"],
        action["verification_commands"],
    )
    assert verified["verification_status"] == "failed"
    enforced = workspace_service.enforce_review(
        run["id"],
        {
            "verdict": "pass",
            "summary": "模型声称测试通过",
            "evidence": ["模型自述"],
            "issues": [],
        },
    )
    assert enforced["verdict"] == "fail"
    assert "确定性" in enforced["summary"]

    database.update_agent_run(run["id"], "completed")
    (repository_root / "user-change.txt").write_text(
        "do not overwrite\n",
        encoding="utf-8",
    )
    with pytest.raises(GitWorkspaceError, match="主仓库存在新改动"):
        workspace_service.apply_workspace(
            workspace_id=workspace["id"],
            session_id=session_id,
        )
    assert (repository_root / "app.py").read_text(encoding="utf-8") == "VALUE = 1\n"


def test_git_repository_registration_requires_desktop_token_and_clean_tree(
    client,
    tmp_path,
    monkeypatch,
):
    repository_root = tmp_path / "desktop-repository"
    _initialize_test_git_repository(repository_root)
    session_id = str(uuid4())
    monkeypatch.setenv("AETHERCHAT_DESKTOP_TOKEN", "desktop-git-test-token")
    payload = {"session_id": session_id, "path": str(repository_root)}

    assert client.post(
        "/workspaces/repositories/register",
        json=payload,
    ).status_code == 403
    registered = client.post(
        "/workspaces/repositories/register",
        json=payload,
        headers={"X-Aether-Desktop-Token": "desktop-git-test-token"},
    )
    assert registered.status_code == 201, registered.text
    assert registered.json()["display_name"] == "desktop-repository"
    assert "root_path" not in registered.json()
    assert client.get(
        "/workspaces/repositories",
        params={"session_id": session_id},
    ).json()[0]["id"] == registered.json()["id"]
    assert client.get(
        "/workspaces/repositories",
        params={"session_id": str(uuid4())},
    ).json() == []

    (repository_root / "dirty.txt").write_text("dirty\n", encoding="utf-8")
    dirty = client.post(
        "/workspaces/repositories/register",
        json=payload,
        headers={"X-Aether-Desktop-Token": "desktop-git-test-token"},
    )
    assert dirty.status_code == 400
    assert "未提交或未跟踪改动" in dirty.json()["detail"]


def test_agent_run_persists_tool_chain_checkpoint_and_idempotent_replay(tmp_path):
    path = tmp_path / "agent-run-v2.sqlite3"
    database = ChatDatabase(path)
    service = AgentRunService(database)
    session_id = str(uuid4())
    run = service.create(session_id, {
        "content": "写入一次工件",
        "targets": ["coder"],
        "mode": "project",
    })

    first = service.begin_tool_call(
        run_id=run["id"],
        session_id=session_id,
        model="coder",
        provider_type="openai_compat",
        provider_model="offline",
        provider_call_id="provider-call-1",
        tool_name="write_artifact",
        arguments={
            "path": "once.md",
            "content": "hello",
            "api_key": "must-not-be-persisted",
        },
        round_index=0,
        call_index=0,
    )
    assert first["disposition"] == "execute"
    assert first["side_effect"] is True
    assert first["arguments"]["api_key"] == "[redacted]"
    assert len(first["arguments_sha256"]) == 64

    completed = service.finish_tool_call(
        first["id"],
        {
            "tool": "write_artifact",
            "status": "success",
            "result": {"path": "once.md", "version": 1},
        },
        tool_run_id="tool-run-1",
    )
    assert completed["status"] == "completed"
    assert completed["verification_status"] == "verified"
    assert [item["kind"] for item in completed["evidence"]] == [
        "tool_result_digest", "artifact_version",
    ]
    assert completed["verified_at"] is not None

    replay = service.begin_tool_call(
        run_id=run["id"],
        session_id=session_id,
        model="coder",
        provider_type="openai_compat",
        provider_model="offline",
        provider_call_id="provider-call-retry",
        tool_name="write_artifact",
        arguments={
            "path": "once.md",
            "content": "hello",
            "api_key": "must-not-be-persisted",
        },
        round_index=4,
        call_index=2,
    )
    assert replay["disposition"] == "replay"
    assert replay["result"]["result"]["version"] == 1
    assert len(database.list_agent_tool_calls(run["id"])) == 1

    failed_call = service.begin_tool_call(
        run_id=run["id"],
        session_id=session_id,
        model="coder",
        provider_type="openai_compat",
        provider_model="offline",
        provider_call_id="provider-call-failed",
        tool_name="write_artifact",
        arguments={"path": "maybe-written.md", "content": "unknown outcome"},
        round_index=1,
        call_index=0,
    )
    service.finish_tool_call(
        failed_call["id"],
        {
            "tool": "write_artifact",
            "status": "error",
            "error": "connection closed after dispatch",
        },
    )
    failed_record = database.get_agent_tool_call(failed_call["id"])
    assert failed_record["verification_status"] == "failed"
    uncertain_retry = service.begin_tool_call(
        run_id=run["id"],
        session_id=session_id,
        model="coder",
        provider_type="openai_compat",
        provider_model="offline",
        provider_call_id="provider-call-failed-retry",
        tool_name="write_artifact",
        arguments={"path": "maybe-written.md", "content": "unknown outcome"},
        round_index=1,
        call_index=0,
    )
    assert uncertain_retry["disposition"] == "uncertain"
    assert uncertain_retry["verification_status"] == "uncertain"

    unverified_call = service.begin_tool_call(
        run_id=run["id"],
        session_id=session_id,
        model="coder",
        provider_type="mcp",
        provider_model="remote",
        provider_call_id="provider-call-unverified",
        tool_name="remote__send_message",
        arguments={"message": "hello"},
        round_index=2,
        call_index=0,
    )
    unverified = service.finish_tool_call(
        unverified_call["id"],
        {"tool": "remote__send_message", "status": "success", "result": {"ok": True}},
    )
    assert unverified["status"] == "completed"
    assert unverified["verification_status"] == "unverified"
    assert unverified["verified_at"] is None
    aggregate = database.get_agent_run(run["id"], include_events=False)
    assert aggregate["effect_count"] == 3
    assert aggregate["verified_effect_count"] == 1
    assert aggregate["failed_effect_count"] == 0
    assert aggregate["uncertain_effect_count"] == 1
    assert aggregate["effect_verification_status"] == "uncertain"

    checkpoint = service.save_checkpoint(
        run_id=run["id"],
        session_id=session_id,
        model="coder",
        kind="tool_round",
        state={
            "version": 1,
            "next_tool_round": 1,
            "messages": [{"role": "user", "content": "resume here"}],
        },
    )
    reopened = ChatDatabase(path)
    latest = reopened.latest_agent_checkpoint(run["id"], "coder")
    assert latest["id"] == checkpoint["id"]
    assert latest["state"]["next_tool_round"] == 1
    detail = reopened.get_agent_run(run["id"])
    assert len(detail["tool_calls"]) == 3
    assert len(detail["checkpoints"]) == 1


def test_agent_run_blocks_uncertain_side_effect_after_restart(tmp_path):
    path = tmp_path / "agent-run-uncertain.sqlite3"
    database = ChatDatabase(path)
    service = AgentRunService(database)
    session_id = str(uuid4())
    run = service.create(session_id, {
        "content": "执行一次副作用",
        "targets": ["coder"],
    })
    service.begin_tool_call(
        run_id=run["id"],
        session_id=session_id,
        model="coder",
        provider_type="anthropic",
        provider_model="offline",
        provider_call_id="toolu-pending",
        tool_name="write_artifact",
        arguments={"path": "uncertain.md", "content": "once"},
        round_index=0,
        call_index=0,
    )

    reopened = ChatDatabase(path)
    assert reopened.recover_incomplete_agent_runs() == 1
    uncertain = reopened.list_agent_tool_calls(run["id"])[0]
    assert uncertain["status"] == "uncertain"
    assert "阻止自动重放" in uncertain["error"]

    resumed = AgentRunService(reopened)
    resumed.prepare_retry(run["id"], session_id)
    blocked = resumed.begin_tool_call(
        run_id=run["id"],
        session_id=session_id,
        model="coder",
        provider_type="anthropic",
        provider_model="offline",
        provider_call_id="toolu-retry",
        tool_name="write_artifact",
        arguments={"path": "uncertain.md", "content": "once"},
        round_index=0,
        call_index=0,
    )
    assert blocked["disposition"] == "uncertain"
    assert "人工检查" in blocked["error"]


def test_effect_verification_trusts_only_known_nested_desktop_receipts(tmp_path):
    database = ChatDatabase(tmp_path / "effect-receipts.sqlite3")
    service = AgentRunService(database)
    session_id = str(uuid4())
    run = service.create(session_id, {
        "content": "核验桌面凭证",
        "targets": ["coder"],
        "mode": "project",
        "execution_strategy": "direct",
    })
    receipt = {
        "success": True,
        "verification": "launcher_accepted",
        "launcher": "firefox",
        "browser": "firefox",
        "accepted_at": "2026-08-11T00:00:00Z",
    }
    nested_result = {
        "plugin_id": "browser-links",
        "status": "success",
        "result": {
            "kind": "open_external",
            "url": "https://example.test",
            "execution_receipt": receipt,
        },
    }

    trusted_call = service.begin_tool_call(
        run_id=run["id"], session_id=session_id, model="coder",
        provider_type="openai_compat", provider_model="offline",
        provider_call_id="trusted-browser", tool_name="plugin__browser_links__open_url",
        arguments={"url": "https://example.test", "browser": "firefox"},
        round_index=0, call_index=0,
    )
    trusted = service.finish_tool_call(trusted_call["id"], {
        "tool": "plugin__browser_links__open_url",
        "status": "success",
        "result": nested_result,
    })
    assert trusted["verification_status"] == "verified"
    assert [item["kind"] for item in trusted["evidence"]] == [
        "tool_result_digest", "execution_receipt",
    ]

    spoofed_call = service.begin_tool_call(
        run_id=run["id"], session_id=session_id, model="coder",
        provider_type="mcp", provider_model="remote",
        provider_call_id="spoofed-browser", tool_name="remote__open_url",
        arguments={"url": "https://example.test"},
        round_index=1, call_index=0,
    )
    spoofed = service.finish_tool_call(spoofed_call["id"], {
        "tool": "remote__open_url",
        "status": "success",
        "result": nested_result,
    })
    assert spoofed["verification_status"] == "unverified"


def test_agent_run_api_lists_session_scoped_records(client):
    session_id = str(uuid4())
    other_session_id = str(uuid4())
    service = AgentRunService(chat_database)
    run = service.create(session_id, {"content": "记录运行", "targets": ["local"]})
    service.observe(run["id"], {
        "type": "error", "model": "local", "code": "provider_error", "message": "offline",
    })
    service.finish(run["id"])

    listed = client.get("/agent-runs", params={"session_id": session_id})
    assert listed.status_code == 200
    assert listed.json()[0]["id"] == run["id"]
    detail = client.get(f"/agent-runs/{run['id']}", params={"session_id": session_id})
    assert detail.status_code == 200
    assert any(event["event_type"] == "run_failed" for event in detail.json()["events"])
    assert client.get(
        f"/agent-runs/{run['id']}", params={"session_id": other_session_id}
    ).status_code == 404


class _MemoryVault:
    def __init__(self):
        self.values = {}

    def get_key(self, ref):
        return self.values.get(ref)

    def set_key(self, ref, value):
        self.values[ref] = value

    def delete_key(self, ref):
        self.values.pop(ref, None)


@pytest.mark.asyncio
async def test_mcp_stdio_discovers_filters_invokes_and_persists_without_secrets(tmp_path):
    credential_vault = _MemoryVault()
    registry = MCPRegistry(tmp_path, credential_vault=credential_vault)
    server_script = PROJECT_ROOT / "tests" / "fixtures" / "mcp_echo_server.py"
    created = await registry.create(MCPConnectionCreate(
        id="local-echo",
        name="本地 Echo",
        transport="stdio",
        command=sys.executable,
        args=[str(server_script)],
        auth_type="none",
        token="must-not-be-persisted",
    ))
    assert created["status"] == "ready"
    assert created["has_token"] is True
    assert created["tools"][0]["name"] == "echo"
    definitions = registry.tool_definitions()
    assert definitions[0].name == "mcp__local_echo__echo"
    invoked = await registry.invoke_tool(definitions[0].name, {"message": "hello"})
    assert invoked["status"] == "success"
    assert invoked["content"][0]["text"] == "echo:hello"
    assert "must-not-be-persisted" not in registry.path.read_text(encoding="utf-8")

    hidden = registry.update("local-echo", MCPConnectionPatch(allowed_tools=[]))
    assert hidden["allowed_tools"] == []
    assert registry.tool_definitions() == []
    registry.delete("local-echo")
    assert credential_vault.values == {}


def test_mcp_remote_transport_rejects_insecure_or_credential_urls(tmp_path):
    with pytest.raises(ValueError, match="HTTPS"):
        MCPConnectionCreate(name="远程", transport="streamable_http", url="http://example.com/mcp")
    with pytest.raises(ValueError, match="凭据"):
        MCPConnectionCreate(name="远程", transport="streamable_http", url="https://user:secret@example.com/mcp")
    local = MCPConnectionCreate(name="本机", transport="streamable_http", url="http://127.0.0.1:3000/mcp")
    assert local.transport == "streamable_http"


@pytest.mark.asyncio
async def test_mcp_streamable_http_discovers_and_invokes_official_server(tmp_path):
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    server_script = PROJECT_ROOT / "tests" / "fixtures" / "mcp_http_server.py"
    env = {**os.environ, "AETHERCHAT_TEST_MCP_PORT": str(port)}
    process = subprocess.Popen(
        [sys.executable, str(server_script)], env=env,
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
    )
    try:
        deadline = time.monotonic() + 8
        while time.monotonic() < deadline:
            try:
                with socket.create_connection(("127.0.0.1", port), timeout=0.2):
                    break
            except OSError:
                await asyncio.sleep(0.05)
        else:
            pytest.fail("MCP Streamable HTTP test server did not start")

        registry = MCPRegistry(tmp_path, credential_vault=_MemoryVault())
        connection = await registry.create(MCPConnectionCreate(
            id="http-math", name="HTTP Math", transport="streamable_http",
            url=f"http://127.0.0.1:{port}/mcp", auth_type="none",
        ))
        assert connection["status"] == "ready"
        definition = registry.tool_definitions()[0]
        assert definition.name == "mcp__http_math__add"
        result = await registry.invoke_tool(definition.name, {"left": 7, "right": 5})
        assert result["status"] == "success"
        assert result["content"][0]["text"] == "12"
    finally:
        process.terminate()
        try:
            process.wait(timeout=3)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=3)


@pytest.mark.asyncio
async def test_mcp_oauth_storage_keeps_tokens_in_vault_and_callback_is_unauthed(client, monkeypatch):
    credential_vault = _MemoryVault()
    storage = VaultOAuthStorage(credential_vault, "secure-server")
    tokens = OAuthToken(access_token="access-secret", refresh_token="refresh-secret", token_type="Bearer")
    client_info = OAuthClientInformationFull(
        client_id="client-id", client_secret="client-secret",
        redirect_uris=["http://127.0.0.1:8000/callback"],
    )
    await storage.set_tokens(tokens)
    await storage.set_client_info(client_info)
    assert (await storage.get_tokens()).refresh_token == "refresh-secret"
    assert (await storage.get_client_info()).client_secret == "client-secret"
    assert all("secret" not in ref.lower() for ref in credential_vault.values)
    storage.clear()
    assert credential_vault.values == {}

    monkeypatch.delenv("AETHERCHAT_ALLOW_UNAUTHENTICATED", raising=False)
    monkeypatch.setenv("AETHERCHAT_API_TOKEN", "required-token")
    callback = client.get(
        "/mcp-connections/oauth/callback/missing",
        params={"error": "<script>alert(1)</script>"},
    )
    assert callback.status_code == 400
    assert "<script>" not in callback.text
    assert "&lt;script&gt;" in callback.text
def test_memory_api_allows_review_edit_and_delete(client):
    session_id = str(uuid4())
    updated = client.patch(f"/memory/{session_id}", json={
        "project_note": "本机项目说明",
        "decisions": ["决定使用 SQLite"],
        "preferences": ["优先使用中文"],
        "open_items": ["下一步补测试"],
    })
    assert updated.status_code == 200
    assert updated.json()["decisions"] == ["决定使用 SQLite"]
    restored = client.get(f"/memory/{session_id}")
    assert restored.json()["project_note"] == "本机项目说明"
    cleared = client.patch(f"/memory/{session_id}", json={
        "project_note": "", "decisions": [], "preferences": [], "open_items": [],
    })
    assert cleared.status_code == 200
    assert cleared.json()["preferences"] == []


def test_project_image_is_sanitized_and_available_for_multimodal_context(client):
    from PIL import Image

    session_id = str(uuid4())
    source = io.BytesIO()
    Image.new("RGB", (32, 24), (120, 80, 40)).save(source, format="PNG", pnginfo=None)
    payload = source.getvalue() + b"untrusted-trailing-payload"
    uploaded = client.post(
        f"/knowledge/{session_id}/attachments",
        files={"file": ("reference.png", payload, "image/png")},
    )
    assert uploaded.status_code == 201
    attachment = uploaded.json()
    assert attachment["chunk_count"] == 0
    urls = knowledge_service.image_data_urls(session_id, [attachment["id"]])
    assert len(urls) == 1
    assert urls[0].startswith("data:image/png;base64,")
    native_path = Path(knowledge_service.image_path(session_id, attachment["id"]))
    assert native_path.is_file()
    assert knowledge_service.image_paths(session_id, [attachment["id"]]) == [
        str(native_path)
    ]
    sanitized = __import__("base64").b64decode(urls[0].split(",", 1)[1])
    assert b"untrusted-trailing-payload" not in sanitized
    preview = client.get(f"/knowledge/{session_id}/attachments/{attachment['id']}/preview")
    assert preview.json()["kind"] == "image"
    image_content = client.get(f"/knowledge/{session_id}/attachments/{attachment['id']}/content")
    assert image_content.status_code == 200
    assert image_content.headers["x-content-type-options"] == "nosniff"
    assert "sandbox" in image_content.headers["content-security-policy"]
    assert image_content.content == sanitized

    disguised = client.post(
        f"/knowledge/{session_id}/attachments",
        files={"file": ("fake.png", b"not an image", "image/png")},
    )
    assert disguised.status_code == 422


@pytest.mark.asyncio
async def test_approval_service_supports_once_session_deny_and_fail_closed():
    service = ApprovalService()
    pending = service.create(
        session_id="session-a",
        model="claude",
        tool_name="run_command",
        risk="command_execution",
        title="运行命令",
        detail="pytest -q",
        arguments={"cmd": "pytest -q"},
        allow_session=True,
    )
    assert pending is not None
    waiter = asyncio.create_task(service.wait(pending.id, timeout=1))
    assert service.decide(pending.id, "session-a", "allow_session")["resolved"] is True
    assert await waiter == "allow_session"
    assert service.create(
        session_id="session-a", model="codex", tool_name="run_command",
        risk="command_execution", title="运行命令", detail="echo ok", arguments={}, allow_session=True,
    ) is None

    once_only = service.create(
        session_id="session-a", model="claude", tool_name="write_local",
        risk="local_file_write", title="写文件", detail="/tmp/a", arguments={}, allow_session=False,
    )
    assert once_only is not None
    with pytest.raises(PermissionError):
        service.decide(once_only.id, "session-b", "allow_once")
    with pytest.raises(ValueError):
        service.decide(once_only.id, "session-a", "allow_session")
    denied = asyncio.create_task(service.wait(once_only.id, timeout=1))
    result = service.decide(
        once_only.id,
        "session-a",
        "deny",
        "不要写文件，改为只读检查",
    )
    assert result["feedback"] == "不要写文件，改为只读检查"
    assert await denied == "deny"
    assert once_only.feedback == "不要写文件，改为只读检查"


def test_tasks_are_uuid_validated_session_scoped_and_persisted(client, tmp_path, monkeypatch):
    original_tasks = list(task_router.fake_tasks)
    task_file = tmp_path / "tasks.json"
    monkeypatch.setattr(task_router, "TASKS_PATH", task_file)
    task_router.fake_tasks.clear()

    session_a = str(uuid4())
    session_b = str(uuid4())
    try:
        created_a = client.post("/tasks/", json={
            "session_id": session_a,
            "title": "session A task",
            "description": "offline regression",
        })
        created_b = client.post("/tasks/", json={
            "session_id": session_b,
            "title": "session B task",
        })
        assert created_a.status_code == 200
        assert created_b.status_code == 200

        tasks_a = client.get("/tasks/", params={"session_id": session_a}).json()
        tasks_b = client.get("/tasks/", params={"session_id": session_b}).json()
        assert [task["title"] for task in tasks_a] == ["session A task"]
        assert [task["title"] for task in tasks_b] == ["session B task"]

        invalid = client.post("/tasks/", json={"session_id": "not-a-uuid", "title": "bad"})
        assert invalid.status_code == 422

        task_id = created_a.json()["id"]
        claimed = client.post(f"/tasks/{task_id}/claim", params={"by": "user:test"})
        assert claimed.status_code == 200
        assert claimed.json()["status"] == "claimed"

        assert task_file.exists()
        task_router.fake_tasks.clear()
        task_router.fake_tasks.extend(task_router._load_tasks(task_file))
        restored = client.get("/tasks/", params={"session_id": session_a}).json()
        assert restored[0]["id"] == task_id
        assert restored[0]["claimed_by"] == "user:test"
    finally:
        task_router.fake_tasks.clear()
        task_router.fake_tasks.extend(original_tasks)


def test_task_dependencies_block_execution_and_tool_success_enters_review(client, tmp_path, monkeypatch):
    from app.services.tool_executor import execute_tool

    original_tasks = list(task_router.fake_tasks)
    monkeypatch.setattr(task_router, "TASKS_PATH", tmp_path / "dependent-tasks.json")
    task_router.fake_tasks.clear()
    session_id = str(uuid4())
    try:
        parent = client.post("/tasks/", json={
            "session_id": session_id,
            "title": "先生成图片",
        }).json()
        child_response = client.post("/tasks/", json={
            "session_id": session_id,
            "title": "再撰写图片描述",
            "depends_on": [parent["id"]],
        })
        assert child_response.status_code == 200
        child = child_response.json()
        assert child["ready"] is False
        assert child["blocked_by"] == [parent["id"]]
        assert client.post(f"/tasks/{child['id']}/claim").status_code == 409
        listed = asyncio.run(execute_tool(
            "list_tasks",
            {},
            "model:qwen",
            session_id,
        ))
        listed_child = next(
            item for item in listed["result"]["tasks"]
            if item["id"] == child["id"]
        )
        assert listed_child["ready"] is False
        blocked_claim = asyncio.run(execute_tool(
            "claim_task",
            {"task_id": child["id"]},
            "model:qwen",
            session_id,
        ))
        assert blocked_claim["status"] == "error"
        assert "blocked by" in blocked_claim["error"]

        assert client.patch(
            f"/tasks/{parent['id']}",
            json={"status": "done"},
        ).status_code == 200
        refreshed = client.get("/tasks/", params={"session_id": session_id}).json()
        child = next(item for item in refreshed if item["id"] == child["id"])
        assert child["ready"] is True
        assert child["blocked_by"] == []

        run = asyncio.run(execute_tool(
            "write_artifact",
            {
                "path": "descriptions/image.md",
                "content": "图片描述",
                "task_id": child["id"],
            },
            "model:qwen",
            session_id,
        ))
        assert run["status"] == "success"
        assert run["result"]["task_sync"]["status"] == "review"
        task_after_tool = client.get("/tasks/", params={"session_id": session_id}).json()
        child_after_tool = next(item for item in task_after_tool if item["id"] == child["id"])
        assert child_after_tool["status"] == "review"
        assert child_after_tool["claimed_by"] == "model:qwen"

        circular = client.patch(
            f"/tasks/{parent['id']}",
            json={"depends_on": [child["id"]]},
        )
        assert circular.status_code == 422
        assert "循环" in circular.json()["detail"]
    finally:
        task_router.fake_tasks.clear()
        task_router.fake_tasks.extend(original_tasks)


def test_tool_registry_upgrades_builtin_schemas_without_reenabling_tools(tmp_path):
    from app.services.tool_registry import ToolRegistry

    registry = ToolRegistry(tmp_path)
    document = json.loads(registry.registry_path.read_text(encoding="utf-8"))
    write_tool = next(
        item for item in document["tools"]
        if item["name"] == "write_artifact"
    )
    write_tool["input_schema"]["properties"].pop("task_id")
    write_tool["enabled"] = False
    registry.registry_path.write_text(
        json.dumps(document, ensure_ascii=False),
        encoding="utf-8",
    )

    upgraded = ToolRegistry(tmp_path)
    upgraded_write = upgraded.find_by_name("write_artifact")
    upgraded_create = upgraded.find_by_name("create_task")
    assert upgraded_write is not None
    assert upgraded_write["enabled"] is False
    assert "task_id" in upgraded_write["input_schema"]["properties"]
    assert upgraded_create is not None
    assert "depends_on" in upgraded_create["input_schema"]["properties"]


def test_linked_plugin_success_advances_task_to_review(client, tmp_path, monkeypatch):
    from app.services.plugin_registry import PluginInvoke, PluginUpdate

    original_tasks = list(task_router.fake_tasks)
    previous_enabled = plugin_registry.get("comfyui")["enabled"]
    monkeypatch.setattr(task_router, "TASKS_PATH", tmp_path / "plugin-task.json")
    task_router.fake_tasks.clear()
    session_id = str(uuid4())

    async def fake_dispatch(_action, _session_id, arguments):
        return {"prompt_id": "test-prompt", "received_task_id": arguments["task_id"]}

    try:
        plugin_registry.update("comfyui", PluginUpdate(enabled=True))
        monkeypatch.setattr(plugin_registry, "_dispatch", fake_dispatch)
        task = client.post("/tasks/", json={
            "session_id": session_id,
            "title": "生成封面",
        }).json()
        result = asyncio.run(plugin_registry.invoke(
            "comfyui",
            PluginInvoke(
                action="generate_image",
                session_id=session_id,
                arguments={
                    "prompt": "雨夜中的重庆",
                    "model": "qwen-image.safetensors",
                    "task_id": task["id"],
                },
            ),
            caller="model:qwen",
        ))
        assert result["result"]["task_sync"]["status"] == "review"
        refreshed = client.get("/tasks/", params={"session_id": session_id}).json()
        assert refreshed[0]["status"] == "review"
        assert refreshed[0]["claimed_by"] == "model:qwen"
    finally:
        plugin_registry.update("comfyui", PluginUpdate(enabled=previous_enabled))
        task_router.fake_tasks.clear()
        task_router.fake_tasks.extend(original_tasks)


def test_chat_history_and_task_events_survive_database_restart(client):
    session_id = uuid4()
    chat_database.create_session(str(session_id), "重启恢复测试", "project")
    for role, author, content in (
        ("user", "user", "继续上次的数据库任务"),
        ("assistant", "qwen", "已经恢复上下文。"),
    ):
        chat_database.save_message(Message(
            id=uuid4(),
            session_id=session_id,
            author_model=author,
            role=role,
            content=content,
            generated_files=[{
                "path": "generated/report.md",
                "name": "report.md",
                "mime_type": "text/markdown",
                "size_bytes": 128,
            }] if role == "assistant" else [],
            reasoning_summary=[{
                "phase": "analysis",
                "detail": "已整理上下文。",
            }] if role == "assistant" else [],
            created_at=datetime.now(UTC),
        ))

    restarted = ChatDatabase(chat_database.path)
    assert [row["content"] for row in restarted.list_messages(str(session_id))] == [
        "继续上次的数据库任务",
        "已经恢复上下文。",
    ]

    history = client.get(f"/sessions/{session_id}/messages")
    assert history.status_code == 200
    assert history.json()[1]["author_model"] == "qwen"
    assert history.json()[1]["generated_files"][0]["path"] == "generated/report.md"
    assert history.json()[1]["reasoning_summary"] == [{
        "phase": "analysis",
        "detail": "已整理上下文。",
    }]
    sessions = client.get("/sessions/").json()
    restored_session = next(item for item in sessions if item["id"] == str(session_id))
    assert restored_session["message_count"] == 2
    assert restored_session["summary"] == "已经恢复上下文。"

    created = client.post("/tasks/", json={"session_id": str(session_id), "title": "记录恢复进度"})
    assert created.status_code == 200
    patched = client.patch(f"/tasks/{created.json()['id']}", json={"status": "in_progress"})
    assert patched.status_code == 200
    events = client.get(f"/sessions/{session_id}/task-events").json()
    assert [event["action"] for event in events[:2]] == ["updated", "created"]


def test_cached_chat_search_native_file_location_and_session_delete(client):
    session_id = str(uuid4())
    created = client.post("/sessions/", json={
        "id": session_id,
        "title": "缓存检索测试",
        "mode": "casual",
    })
    assert created.status_code == 200

    written = client.post("/artifacts/write", json={
        "session_id": session_id,
        "path": "cache/result.txt",
        "content": "persistent generated output",
        "commit_message": "cache generated output",
        "author": "test",
    })
    assert written.status_code == 200
    native = client.get(
        "/artifacts/cache/result.txt/native-location",
        params={"session_id": session_id},
    )
    assert native.status_code == 200
    native_path = Path(native.json()["path"])
    assert native_path.is_file()
    content = client.get(
        "/artifacts/cache/result.txt/content",
        params={"session_id": session_id},
    )
    assert content.status_code == 200
    assert content.content == b"persistent generated output"
    assert content.headers["x-content-type-options"] == "nosniff"
    assert content.headers["content-disposition"].startswith("inline;")

    message = Message(
        id=uuid4(),
        session_id=session_id,
        author_model="qwen",
        role="assistant",
        content="跨会话缓存检索唯一词：山城夜雨。",
        generated_files=[{
            "path": "cache/result.txt",
            "name": "result.txt",
            "mime_type": "text/plain",
            "size_bytes": native_path.stat().st_size,
            "kind": "generated_file",
            "source": "qwen",
        }],
        created_at=datetime.now(UTC),
    )
    chat_database.save_message(message, mode="casual")
    uploaded = client.post(
        f"/knowledge/{session_id}/attachments",
        files={
            "file": (
                "cached-note.txt",
                "用于验证删除会话时清理资料。".encode(),
                "text/plain",
            )
        },
    )
    assert uploaded.status_code == 201
    task = client.post("/tasks/", json={
        "session_id": session_id,
        "title": "删除时清理",
    })
    assert task.status_code == 200

    search = client.get("/sessions/search", params={
        "q": "缓存检索唯一词",
        "limit": 10,
    })
    assert search.status_code == 200
    assert search.json()[0]["session_id"] == session_id
    assert "山城夜雨" in search.json()[0]["snippet"]

    deleted = client.delete(f"/sessions/{session_id}")
    assert deleted.status_code == 200
    assert deleted.json()["deleted"] is True
    assert deleted.json()["cleanup"]["knowledge"]["attachments"] == 1
    assert not native_path.exists()
    assert chat_database.get_session(session_id) is None
    assert all(
        str(item.session_id) != session_id
        for item in task_router.fake_tasks
    )
    assert client.get(
        "/sessions/search",
        params={"q": "缓存检索唯一词"},
    ).json() == []
    assert client.delete(f"/sessions/{session_id}").status_code == 404


def test_deleted_seed_model_stays_deleted_after_restart(tmp_path, monkeypatch):
    monkeypatch.setenv("AETHERCHAT_DATA_DIR", str(tmp_path))
    first = ModelRegistry()
    assert first.get("grok") is not None
    assert first.delete("grok") is not None
    assert first.get("grok") is None

    restarted = ModelRegistry()
    assert restarted.get("grok") is None
    assert restarted.get("gpt") is not None


def test_linux_vault_uses_keyring_and_migrates_plaintext_index(tmp_path, monkeypatch):
    if os.name == "nt":
        pytest.skip("Linux/macOS keyring behavior")

    class FakeKeyring:
        values: dict[tuple[str, str], str] = {}

        @classmethod
        def set_password(cls, service, ref, value):
            cls.values[(service, ref)] = value

        @classmethod
        def get_password(cls, service, ref):
            return cls.values.get((service, ref))

    monkeypatch.setenv("AETHERCHAT_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(vault_module, "system_keyring", FakeKeyring)
    vault_path = tmp_path / "vault.json"
    vault_path.write_text('{"LEGACY_KEY": "legacy-test-value"}', encoding="utf-8")

    secure_vault = vault_module.SimpleVault()
    assert secure_vault.get_key("LEGACY_KEY") == "legacy-test-value"
    secure_vault.set_key("NEW_KEY", "new-test-value")
    index = json.loads(vault_path.read_text(encoding="utf-8"))
    assert index == {"LEGACY_KEY": "keyring", "NEW_KEY": "keyring"}
    assert "legacy-test-value" not in vault_path.read_text(encoding="utf-8")
    assert "new-test-value" not in vault_path.read_text(encoding="utf-8")


@pytest.mark.asyncio
async def test_sandbox_uses_prepared_workspace_cwd_and_real_exit_code(tmp_path, monkeypatch):
    workspace = tmp_path / "workspace"
    (workspace / "src").mkdir(parents=True)
    captured = {}

    class FakeContainer:
        removed = False

        def wait(self, timeout):
            captured["wait_timeout"] = timeout
            return {"StatusCode": 7}

        def logs(self, **kwargs):
            return b"command output"

        def remove(self, force=False):
            self.removed = force

    container = FakeContainer()

    class FakeContainers:
        def run(self, **kwargs):
            captured["run"] = kwargs
            return container

    class FakeClient:
        containers = FakeContainers()

        def close(self):
            captured["closed"] = True

    monkeypatch.setattr(sandbox_executor.docker, "from_env", lambda: FakeClient())
    result = await sandbox_executor.execute_in_sandbox(
        "python -m pytest",
        cwd="src",
        timeout=12,
        workspace_path=workspace,
    )

    assert result["exit_code"] == 7
    assert result["stdout"] == "command output"
    assert captured["run"]["working_dir"] == "/workspace/src"
    assert captured["run"]["volumes"][str(workspace.resolve())]["mode"] == "rw"
    assert captured["run"]["network_mode"] == "none"
    assert container.removed is True
    assert captured["closed"] is True


def test_sandbox_rejects_path_escape_and_unbounded_inputs():
    with pytest.raises(ValueError, match="inside /workspace"):
        sandbox_executor.normalize_container_cwd("../outside")
    with pytest.raises(ValueError, match="timeout"):
        sandbox_executor.validate_command("echo ok", 301)
    with pytest.raises(ValueError, match="too large"):
        sandbox_executor.validate_command("x" * 20_001, 30)


@pytest.mark.asyncio
async def test_provider_context_uses_model_limit_and_preserves_latest_turn():
    messages = []
    for index in range(12):
        messages.extend([
            ChatMessage(role="user", content=f"old-user-{index} " + ("甲" * 900)),
            ChatMessage(role="assistant", content=f"old-answer-{index} " + ("乙" * 900)),
        ])
    messages.append(ChatMessage(role="user", content="LATEST-REQUEST must remain"))

    result = await build_messages_for_provider(messages, provider_limit=4096, output_reserve=1024)
    assert result[-1].content == "LATEST-REQUEST must remain"
    assert any("上下文上限" in item.content for item in result if item.role == "system")
    assert sum(estimate_tokens(item.content) + 8 for item in result) <= 4096 - 1024
    first_conversation = next(item for item in result if item.role != "system")
    assert first_conversation.role == "user"


@pytest.mark.asyncio
async def test_provider_context_drops_complete_native_tool_turns():
    messages = [
        ChatMessage(role="user", content="old request " + ("甲" * 6000)),
        ChatMessage(
            role="assistant",
            tool_calls=[ToolCall(
                id="old-call",
                name="inspect",
                arguments={"target": "old"},
            )],
        ),
        ChatMessage(
            role="tool",
            name="inspect",
            tool_call_id="old-call",
            content='{"result": "old"}',
        ),
        ChatMessage(role="user", content="LATEST-REQUEST"),
    ]

    result = await build_messages_for_provider(
        messages,
        provider_limit=2048,
        output_reserve=512,
    )
    conversation = [message for message in result if message.role != "system"]

    assert [(message.role, message.content) for message in conversation] == [
        ("user", "LATEST-REQUEST"),
    ]


@pytest.mark.asyncio
async def test_provider_context_caps_recent_messages_at_a_user_boundary():
    messages = []
    for index in range(10):
        messages.extend([
            ChatMessage(role="user", content=f"user-{index}"),
            ChatMessage(role="assistant", content=f"answer-{index}"),
        ])
    messages.append(ChatMessage(role="user", content="LATEST-REQUEST"))

    result = await build_messages_for_provider(
        messages,
        provider_limit=32_768,
        output_reserve=2048,
        max_conversation_messages=8,
    )
    conversation = [message for message in result if message.role != "system"]

    assert len(conversation) <= 8
    assert conversation[0].role == "user"
    assert conversation[-1].content == "LATEST-REQUEST"
    assert all(message.content != "user-0" for message in conversation)
    assert any(
        "上下文上限" in message.content
        for message in result
        if message.role == "system"
    )


@pytest.mark.asyncio
async def test_provider_context_keeps_latest_multi_model_turn_together():
    messages = [
        ChatMessage(role="user", content=f"old-user-{index}")
        if index % 2 == 0
        else ChatMessage(role="assistant", content=f"old-answer-{index}")
        for index in range(20)
    ]
    messages.append(ChatMessage(role="user", content="LATEST-REQUEST"))
    messages.extend(
        ChatMessage(role="assistant", content=f"model-answer-{index}")
        for index in range(10)
    )

    result = await build_messages_for_provider(
        messages,
        provider_limit=32_768,
        output_reserve=2048,
        max_conversation_messages=8,
    )
    conversation = [message for message in result if message.role != "system"]

    assert conversation[0].content == "LATEST-REQUEST"
    assert len(conversation) == 11
    assert conversation[-1].content == "model-answer-9"


class _AsyncChunkStream:
    def __init__(self, chunks):
        self._chunks = iter(chunks)

    def __aiter__(self):
        return self

    async def __anext__(self):
        try:
            return next(self._chunks)
        except StopIteration as exc:
            raise StopAsyncIteration from exc


def test_openai_native_tool_results_remain_contiguous_before_visual_observation():
    converted = provider_module._openai_messages([
        ChatMessage(
            role="assistant",
            tool_calls=[
                ToolCall(id="call-a", name="observe", arguments={}),
                ToolCall(id="call-b", name="inspect", arguments={"target": "window"}),
            ],
        ),
        ChatMessage(
            role="tool",
            name="observe",
            tool_call_id="call-a",
            content='{"ok": true}',
            images=["data:image/png;base64,aW1hZ2U="],
        ),
        ChatMessage(
            role="tool",
            name="inspect",
            tool_call_id="call-b",
            content='{"window": "editor"}',
        ),
    ], include_images=True)

    assert [message["role"] for message in converted] == [
        "assistant",
        "tool",
        "tool",
        "user",
    ]
    assert [message["tool_call_id"] for message in converted[1:3]] == [
        "call-a",
        "call-b",
    ]
    assert converted[3]["content"][1] == {
        "type": "image_url",
        "image_url": {
            "url": "data:image/png;base64,aW1hZ2U=",
            "detail": "auto",
        },
    }


@pytest.mark.asyncio
async def test_openai_provider_streams_tools_and_allows_local_without_key(monkeypatch):
    captured = {}

    def fake_http_client(**kwargs):
        captured["http_client"] = kwargs
        return "local-http-client"

    class FakeAsyncOpenAI:
        def __init__(self, **kwargs):
            captured["client"] = kwargs
            self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))

        async def create(self, **kwargs):
            captured["request"] = kwargs
            return _AsyncChunkStream([
                SimpleNamespace(
                    choices=[SimpleNamespace(delta=SimpleNamespace(content="first "))],
                    usage=None,
                ),
                SimpleNamespace(
                    choices=[SimpleNamespace(delta=SimpleNamespace(
                        content="second",
                        tool_calls=[SimpleNamespace(
                            index=0,
                            id="call-1",
                            function=SimpleNamespace(name="list_tasks", arguments="{}"),
                        )],
                    ))],
                    usage=None,
                ),
                SimpleNamespace(
                    choices=[],
                    usage=SimpleNamespace(prompt_tokens=12, completion_tokens=4),
                ),
            ])

    monkeypatch.setattr(provider_module, "AsyncOpenAI", FakeAsyncOpenAI)
    monkeypatch.setattr(provider_module, "DefaultAsyncHttpxClient", fake_http_client)
    async def fake_runtime_ready(_config):
        return {"ready": True, "managed": False, "started": False}
    monkeypatch.setattr(provider_module.local_model_runtime, "ensure_ready", fake_runtime_ready)
    monkeypatch.delenv("UNSET_LOCAL_TEST_KEY", raising=False)
    provider = provider_module.OpenAICompatProvider(ProviderConfig(
        name="local-test",
        type="local",
        model="qwen-local-model",
        base_url="http://127.0.0.1:9999/v1",
        api_key_ref="UNSET_LOCAL_TEST_KEY",
    ))
    tools = [ToolDefinition(name="list_tasks", description="List tasks")]

    chunks = [
        chunk
        async for chunk in provider.stream_chat(
            [
                ChatMessage(role="system", content="first instruction"),
                ChatMessage(role="user", content="hello"),
                ChatMessage(
                    role="assistant",
                    content="checking",
                    tool_calls=[ToolCall(
                        id="historic-call",
                        name="list_tasks",
                        arguments={"status": "open"},
                    )],
                ),
                ChatMessage(
                    role="tool",
                    name="list_tasks",
                    tool_call_id="historic-call",
                    content='{"tasks": []}',
                ),
                ChatMessage(role="system", content="second instruction"),
            ],
            tools=tools,
            tool_choice="required",
        )
    ]

    assert captured["client"]["api_key"] == "local-model"
    assert captured["client"]["http_client"] == "local-http-client"
    assert captured["http_client"]["trust_env"] is False
    assert captured["request"]["stream"] is True
    assert captured["request"]["extra_body"] == {
        "chat_template_kwargs": {"enable_thinking": False},
    }
    assert captured["request"]["messages"] == [
        {"role": "system", "content": "first instruction\n\nsecond instruction"},
        {"role": "user", "content": "hello"},
        {
            "role": "assistant",
            "content": "checking",
            "tool_calls": [{
                "id": "historic-call",
                "type": "function",
                "function": {
                    "name": "list_tasks",
                    "arguments": '{"status": "open"}',
                },
            }],
        },
        {
            "role": "tool",
            "tool_call_id": "historic-call",
            "content": '{"tasks": []}',
        },
    ]
    assert captured["request"]["tools"][0]["function"]["name"] == "list_tasks"
    assert captured["request"]["tool_choice"] == "required"
    assert [chunk["delta"] for chunk in chunks if chunk["delta"]] == ["first ", "second"]
    assert sum(bool(chunk.get("tool_activity")) for chunk in chunks) == 1
    assert next(chunk["tool_calls"] for chunk in chunks if chunk.get("tool_calls")) == [{
        "id": "call-1",
        "name": "list_tasks",
        "arguments": {},
    }]
    assert chunks[-1]["usage"] == {"prompt_tokens": 12, "completion_tokens": 4}


@pytest.mark.asyncio
async def test_anthropic_provider_accepts_shared_tool_definitions(monkeypatch):
    captured = {}

    class FakeAsyncAnthropic:
        def __init__(self, **kwargs):
            captured["client"] = kwargs
            self.messages = SimpleNamespace(create=self.create)

        async def create(self, **kwargs):
            captured["request"] = kwargs
            return _AsyncChunkStream([
                SimpleNamespace(
                    type="message_start",
                    message=SimpleNamespace(usage=SimpleNamespace(input_tokens=20)),
                ),
                SimpleNamespace(
                    type="content_block_delta",
                    delta=SimpleNamespace(text="ok"),
                ),
                SimpleNamespace(
                    type="content_block_start",
                    index=1,
                    content_block=SimpleNamespace(
                        type="tool_use",
                        id="toolu-new",
                        name="read_artifact",
                        input={},
                    ),
                ),
                SimpleNamespace(
                    type="content_block_delta",
                    index=1,
                    delta=SimpleNamespace(partial_json='{"path":"next.md"}'),
                ),
                SimpleNamespace(
                    type="message_delta",
                    usage=SimpleNamespace(output_tokens=5),
                ),
            ])

    monkeypatch.setattr(provider_module, "AsyncAnthropic", FakeAsyncAnthropic)
    monkeypatch.setenv("ANTHROPIC_OFFLINE_TEST_KEY", "fake-offline-key")
    provider = provider_module.AnthropicProvider(ProviderConfig(
        name="claude-test",
        type="anthropic",
        model="claude-offline",
        api_key_ref="ANTHROPIC_OFFLINE_TEST_KEY",
    ))

    chunks = [
        chunk
        async for chunk in provider.stream_chat(
            [
                ChatMessage(role="system", content="first instruction"),
                ChatMessage(role="user", content="hello"),
                ChatMessage(
                    role="assistant",
                    content="checking",
                    tool_calls=[
                        ToolCall(
                            id="toolu-success",
                            name="read_artifact",
                            arguments={"path": "notes.md"},
                        ),
                        ToolCall(
                            id="toolu-error",
                            name="read_artifact",
                            arguments={"path": "missing.md"},
                        ),
                    ],
                ),
                ChatMessage(
                    role="tool",
                    name="read_artifact",
                    tool_call_id="toolu-success",
                    content='{"text": "done"}',
                ),
                ChatMessage(
                    role="tool",
                    name="read_artifact",
                    tool_call_id="toolu-error",
                    tool_error=True,
                    content="not found",
                    images=["data:image/png;base64,aW1hZ2U="],
                ),
                ChatMessage(role="system", content="second instruction"),
            ],
            tools=[ToolDefinition(name="read_artifact", description="Read")],
            tool_choice="required",
        )
    ]

    assert captured["request"]["system"] == "first instruction\n\nsecond instruction"
    assert captured["request"]["tools"] == [{
        "name": "read_artifact",
        "description": "Read",
        "input_schema": {"type": "object", "properties": {}},
    }]
    assert captured["request"]["tool_choice"] == {"type": "any"}
    assert captured["request"]["messages"] == [
        {"role": "user", "content": "hello"},
        {
            "role": "assistant",
            "content": [
                {"type": "text", "text": "checking"},
                {
                    "type": "tool_use",
                    "id": "toolu-success",
                    "name": "read_artifact",
                    "input": {"path": "notes.md"},
                },
                {
                    "type": "tool_use",
                    "id": "toolu-error",
                    "name": "read_artifact",
                    "input": {"path": "missing.md"},
                },
            ],
        },
        {
            "role": "user",
            "content": [
                {
                    "type": "tool_result",
                    "tool_use_id": "toolu-success",
                    "content": '{"text": "done"}',
                },
                {
                    "type": "tool_result",
                    "tool_use_id": "toolu-error",
                    "content": [
                        {"type": "text", "text": "not found"},
                        {
                            "type": "image",
                            "source": {
                                "type": "base64",
                                "media_type": "image/png",
                                "data": "aW1hZ2U=",
                            },
                        },
                    ],
                    "is_error": True,
                },
            ],
        },
    ]
    assert chunks[0]["delta"] == "ok"
    assert sum(bool(chunk.get("tool_activity")) for chunk in chunks) == 1
    assert next(chunk["tool_calls"] for chunk in chunks if chunk.get("tool_calls")) == [{
        "id": "toolu-new",
        "name": "read_artifact",
        "arguments": {"path": "next.md"},
    }]
    assert chunks[-1]["usage"] == {"prompt_tokens": 20, "completion_tokens": 5}


class _FakeProvider:
    def __init__(self, name, *, cost_in=0.0, cost_out=0.0, behavior="normal", started=None, persona=None):
        self.config = ProviderConfig(
            name=name,
            type="openai_compat",
            model=f"{name}-model",
            api_key_ref=f"{name.upper()}_OFFLINE_KEY",
            cost_per_1k_input=cost_in,
            cost_per_1k_output=cost_out,
            persona=persona,
        )
        self.behavior = behavior
        self.started = started
        self.called = False
        self.received_messages = []

    async def stream_chat(self, messages, tools=None, **kwargs):
        self.called = True
        self.received_messages = messages
        if self.started:
            self.started.set()
        if self.behavior == "error":
            if False:
                yield {}
            raise RuntimeError("offline provider failure")
        if self.behavior == "slow":
            for index in range(8):
                await asyncio.sleep(0.025)
                yield {"delta": f"{self.config.name}-{index}", "usage": None}
        else:
            await asyncio.sleep(0.01)
            yield {"delta": f"{self.config.name}-done", "usage": None}
        yield {"delta": "", "usage": {"prompt_tokens": 1000, "completion_tokens": 500}}


async def _collect_events(orchestrator, targets):
    return [
        event
        async for event in orchestrator.stream_user_message({
            "content": "offline orchestration",
            "targets": targets,
            "mode": "casual",
        })
    ]


def _orchestrator_with(providers):
    orchestrator = SessionOrchestrator(str(uuid4()))
    orchestrator.providers = {provider.config.name: provider for provider in providers}
    orchestrator.refresh_providers = lambda: None
    return orchestrator


def test_orchestrator_routes_tools_and_caps_local_active_context(monkeypatch):
    orchestrator = SessionOrchestrator(str(uuid4()))
    definitions = orchestrator._available_tools()
    assert len(definitions) > 10

    plain = orchestrator._route_tools("请解释这段概念", definitions)
    comfy = orchestrator._route_tools("请用 ComfyUI 生成一张图片", definitions)
    bigasp = orchestrator._route_tools(
        "现在使用comfyui 用bigaspv2.safetensors模型来进行图片生成",
        definitions,
    )
    status_only = orchestrator._route_tools(
        "请检查 ComfyUI 模型和 LTX 状态，不要生成任何图片或视频",
        definitions,
    )
    planned_status_only = orchestrator._route_tools(
        "请检查 ComfyUI 模型和 LTX 状态，不要生成任何图片或视频",
        definitions,
        {
            "plan_step_id": "synthesizer",
            "instruction": "生成单一最终交付，并核验图像模型数量",
        },
    )
    workflow = orchestrator._route_tools(
        "请运行 ComfyUI 工作流 demo",
        definitions,
    )
    negated_workflow = orchestrator._route_tools(
        "不要运行 ComfyUI 工作流，只检查状态",
        definitions,
    )
    all_tools = orchestrator._route_tools("列出全部工具", definitions)
    assert plain == []
    assert comfy
    assert all(
        tool.name.startswith("plugin__comfyui__")
        for tool in comfy
    )
    assert [tool.name for tool in bigasp] == [
        "plugin__comfyui__generate_image"
    ]
    assert [tool.name for tool in status_only] == [
        "plugin__comfyui__status"
    ]
    assert not {
        "plugin__comfyui__generate_image",
        "plugin__comfyui__generate_video",
        "plugin__comfyui__run_workflow",
    } & {tool.name for tool in planned_status_only}
    assert "plugin__comfyui__run_workflow" in {
        tool.name for tool in workflow
    }
    assert "plugin__comfyui__run_workflow" not in {
        tool.name for tool in negated_workflow
    }
    assert len(comfy) < len(all_tools) == len(definitions)

    monkeypatch.setenv("AETHERCHAT_LOCAL_ACTIVE_CONTEXT_TOKENS", "8192")
    monkeypatch.setenv("AETHERCHAT_LOCAL_MAX_OUTPUT_TOKENS", "2048")
    local = SimpleNamespace(config=ProviderConfig(
        name="local-performance",
        type="local",
        model="qwen-35b",
        api_key_ref="LOCAL_PERFORMANCE_KEY",
        max_tokens=8192,
        token_limit=253952,
    ))
    provider_limit, output_reserve, tool_tokens = orchestrator._context_limits(
        local,
        [],
    )
    assert provider_limit == 8192
    assert output_reserve == 2048
    assert tool_tokens == 0
    monkeypatch.delenv("AETHERCHAT_LOCAL_MAX_HISTORY_MESSAGES", raising=False)
    assert orchestrator._history_message_limit(local) == 8
    monkeypatch.setenv("AETHERCHAT_LOCAL_MAX_HISTORY_MESSAGES", "20")
    assert orchestrator._history_message_limit(local) == 20

    remote = SimpleNamespace(config=ProviderConfig(
        name="remote",
        type="openai_compat",
        model="remote-model",
        api_key_ref="REMOTE_KEY",
    ))
    assert orchestrator._history_message_limit(remote) is None


@pytest.mark.asyncio
async def test_orchestrator_requires_real_comfyui_generation_without_status_round(monkeypatch):
    class RequiredComfyProvider(_FakeProvider):
        def __init__(self):
            super().__init__("required-comfy")
            self.requests = []

        async def stream_chat(self, messages, tools=None, **kwargs):
            tool_names = [tool.name for tool in (tools or [])]
            self.requests.append({
                "tools": tool_names,
                "tool_choice": kwargs.get("tool_choice"),
            })
            if len(self.requests) == 1:
                yield {
                    "delta": "",
                    "usage": None,
                    "tool_calls": [{
                        "id": "comfy-generate",
                        "name": "plugin__comfyui__generate_image",
                        "arguments": {
                            "prompt": "a blue ceramic mug, studio light",
                        },
                    }],
                }
            else:
                raise AssertionError("successful required generation must not request a summary round")
            yield {
                "delta": "",
                "usage": {"prompt_tokens": 10, "completion_tokens": 4},
            }

    async def fake_plugin_invoke(tool_name, _session_id, _caller, _arguments, **_kwargs):
        return {
            "plugin_id": "comfyui",
            "action": "generate_image",
            "status": "success",
            "result": {
                "generated_files": [{
                    "path": "generated/comfyui/prompt/image.png",
                    "name": "image.png",
                    "mime_type": "image/png",
                    "size_bytes": 256,
                    "source": "comfyui",
                }],
            },
        }

    provider = RequiredComfyProvider()
    orchestrator = _orchestrator_with([provider])
    monkeypatch.setattr(
        orchestrator_module.plugin_registry,
        "invoke_tool",
        fake_plugin_invoke,
    )

    async def allow_tool(_queue, _model, _tool_name, _arguments):
        return True, ""

    orchestrator._authorize_model_tool = allow_tool
    events = [
        event
        async for event in orchestrator.stream_user_message({
            "content": "请用 ComfyUI 生成一张蓝色陶瓷杯图片",
            "targets": ["required-comfy"],
            "mode": "casual",
        })
    ]

    assert provider.requests[0] == {
        "tools": ["plugin__comfyui__generate_image"],
        "tool_choice": "required",
    }
    assert len(provider.requests) == 1
    assert any(
        event["type"] == "tool_run"
        and event["run"]["tool_name"] == "plugin__comfyui__generate_image"
        and event["run"]["status"] == "success"
        for event in events
    )
    completed = next(
        event["message"]
        for event in events
        if event["type"] == "message_complete"
        and event["message"]["author_model"] == "required-comfy"
    )
    assert completed["generated_files"][0]["path"] == (
        "generated/comfyui/prompt/image.png"
    )
    assert "可直接预览或打开" in completed["content"]
    assert "generated/comfyui/prompt/image.png" in completed["content"]
    assert [step["phase"] for step in completed["reasoning_summary"]] == [
        "analysis",
        "tool",
        "result",
        "final",
    ]
    assert all(
        "a blue ceramic mug" not in step["detail"]
        for step in completed["reasoning_summary"]
    )
    assert any(
        event["type"] == "reasoning_step"
        and event["model"] == "required-comfy"
        for event in events
    )
    assert SessionOrchestrator._required_comfyui_tools("图片呢") == []
    assert SessionOrchestrator._required_comfyui_tools(
        "请生成一张图片"
    ) == ["plugin__comfyui__generate_image"]
    assert SessionOrchestrator._required_comfyui_tools(
        "请用 ComfyUI 生成一张蓝色陶瓷杯的产品摄影图"
    ) == ["plugin__comfyui__generate_image"]
    assert SessionOrchestrator._required_comfyui_tools(
        "现在使用comfyui 用bigaspv2.safetensors模型来进行图片生成"
    ) == ["plugin__comfyui__generate_image"]
    assert SessionOrchestrator._required_comfyui_tools(
        "使用 bigaspv2模型对这个文章所描述的场景来进行多张连贯的图片。"
    ) == ["plugin__comfyui__generate_image"]
    assert SessionOrchestrator._required_comfyui_tools(
        "只检查 ComfyUI 状态；不要生成任何图片或视频"
    ) == []
    assert SessionOrchestrator._required_comfyui_tools(
        "Check ComfyUI status; do not generate images or videos"
    ) == []
    assert SessionOrchestrator._required_comfyui_tools(
        "只报告可用于图像生成的模型数量"
    ) == []
    assert SessionOrchestrator._required_comfyui_tools(
        "检测本地图片生成模型都有什么"
    ) == []
    assert SessionOrchestrator._requested_comfyui_image_model(
        "使用 bigaspv2.safetensors 来生成图片"
    ) == "bigaspv2.safetensors"
    assert SessionOrchestrator._requested_comfyui_image_model(
        "使用 bigaspv2模型来进行多张连贯的图片"
    ) == "bigaspv2"
    assert SessionOrchestrator._requested_comfyui_video_model(
        "使用 Wan2.2_Remix_NSFW_i2v_14b_high_lighting_fp8_e4m3fn_v3.0.safetensors 生成视频"
    ) == "Wan2.2_Remix_NSFW_i2v_14b_high_lighting_fp8_e4m3fn_v3.0.safetensors"
    assert SessionOrchestrator._required_comfyui_tools(
        "只检查 ComfyUI 图像模型状态；不要生成任何图片或视频",
        {"instruction": "生成单一最终交付，并核验图像模型数量"},
    ) == []


@pytest.mark.asyncio
async def test_orchestrator_injects_attached_image_and_requested_video_model(monkeypatch):
    from app.services.comfyui_service import WAN22_REMIX_HIGH_MODEL
    from PIL import Image

    class AttachedVideoProvider(_FakeProvider):
        async def stream_chat(self, messages, tools=None, **kwargs):
            assert [tool.name for tool in (tools or [])] == [
                "plugin__comfyui__generate_video"
            ]
            yield {
                "delta": "",
                "usage": None,
                "tool_calls": [{
                    "id": "attached-video",
                    "name": "plugin__comfyui__generate_video",
                    "arguments": {
                        "prompt": "slow cinematic movement",
                        "pipeline": "wan-i2v",
                    },
                }],
            }
            yield {
                "delta": "",
                "usage": {"prompt_tokens": 10, "completion_tokens": 4},
            }

    session_id = str(uuid4())
    source = io.BytesIO()
    Image.new("RGB", (32, 24), (30, 60, 90)).save(source, format="PNG")
    attachment = knowledge_service.create_attachment(
        session_id,
        "start.png",
        "image/png",
        source.getvalue(),
    )
    expected_path = knowledge_service.image_path(session_id, attachment["id"])
    captured = {}

    async def fake_plugin_invoke(tool_name, received_session, _caller, arguments, **_kwargs):
        captured.update(arguments)
        assert tool_name == "plugin__comfyui__generate_video"
        assert received_session == session_id
        return {
            "plugin_id": "comfyui",
            "action": "generate_video",
            "status": "success",
            "result": {
                "generated_files": [{
                    "path": "generated/comfyui/video/output.webm",
                    "name": "output.webm",
                    "mime_type": "video/webm",
                    "size_bytes": 512,
                    "source": "comfyui",
                }],
            },
        }

    provider = AttachedVideoProvider("attached-video-provider")
    orchestrator = SessionOrchestrator(session_id)
    orchestrator.providers = {provider.config.name: provider}
    orchestrator.refresh_providers = lambda: None
    monkeypatch.setattr(
        orchestrator_module.plugin_registry,
        "invoke_tool",
        fake_plugin_invoke,
    )

    async def allow_tool(_queue, _model, _tool_name, _arguments):
        return True, ""

    orchestrator._authorize_model_tool = allow_tool
    events = [
        event
        async for event in orchestrator.stream_user_message({
            "content": f"使用 {WAN22_REMIX_HIGH_MODEL} 生成视频",
            "targets": [provider.config.name],
            "mode": "casual",
            "attachment_ids": [attachment["id"]],
        })
    ]

    assert captured["input_image"] == expected_path
    assert captured["_user_requested_model"] == WAN22_REMIX_HIGH_MODEL
    assert any(
        event["type"] == "tool_run" and event["run"]["status"] == "success"
        for event in events
    )


@pytest.mark.asyncio
async def test_orchestrator_exposes_comfyui_status_only_once_after_success(monkeypatch):
    class StatusOnceProvider(_FakeProvider):
        def __init__(self):
            super().__init__("status-once")
            self.requested_tools = []

        async def stream_chat(self, messages, tools=None, **kwargs):
            self.requested_tools.append([tool.name for tool in (tools or [])])
            if len(self.requested_tools) == 1:
                yield {
                    "delta": "",
                    "usage": None,
                    "tool_calls": [{
                        "id": "comfy-status",
                        "name": "plugin__comfyui__status",
                        "arguments": {},
                    }],
                }
            else:
                yield {"delta": "ComfyUI 已就绪。", "usage": None}
            yield {
                "delta": "",
                "usage": {"prompt_tokens": 10, "completion_tokens": 4},
            }

    async def fake_plugin_invoke(tool_name, _session_id, _caller, _arguments, **_kwargs):
        assert tool_name == "plugin__comfyui__status"
        return {
            "ready": True,
            "generation": {"image_models": [{"name": "bigaspv2.safetensors"}]},
        }

    provider = StatusOnceProvider()
    orchestrator = _orchestrator_with([provider])
    monkeypatch.setattr(
        orchestrator_module.plugin_registry,
        "invoke_tool",
        fake_plugin_invoke,
    )

    async def allow_tool(_queue, _model, _tool_name, _arguments):
        return True, ""

    orchestrator._authorize_model_tool = allow_tool
    events = [
        event
        async for event in orchestrator.stream_user_message({
            "content": "检测本地图片生成模型都有什么",
            "targets": ["status-once"],
            "mode": "casual",
        })
    ]

    assert provider.requested_tools == [
        ["plugin__comfyui__status"],
        [],
    ]
    assert sum(
        event["type"] == "tool_run"
        and event["run"]["tool_name"] == "plugin__comfyui__status"
        for event in events
    ) == 1


@pytest.mark.asyncio
async def test_orchestrator_emits_factual_provider_status_telemetry():
    provider = _FakeProvider("telemetry")
    events = await _collect_events(
        _orchestrator_with([provider]),
        ["telemetry"],
    )
    statuses = [
        event for event in events
        if event["type"] == "provider_status"
    ]
    assert [event["phase"] for event in statuses] == [
        "requesting",
        "streaming",
        "finalizing",
    ]
    assert {event["provider_type"] for event in statuses} == {"openai_compat"}
    assert statuses[0]["context_messages"] >= 2
    assert statuses[0]["context_tokens"] > 0
    assert statuses[0]["tool_count"] == 0
    assert statuses[1]["first_token_ms"] >= 0
    completed = next(
        event["message"]
        for event in events
        if event["type"] == "message_complete"
        and event["message"]["author_model"] == "telemetry"
    )
    assert [step["phase"] for step in completed["reasoning_summary"]] == [
        "analysis",
        "final",
    ]


@pytest.mark.asyncio
async def test_orchestrator_reports_tool_activity_before_arguments_finish():
    class ToolActivityProvider(_FakeProvider):
        async def stream_chat(self, messages, tools=None, **kwargs):
            yield {"delta": "", "usage": None, "tool_activity": True}
            yield {"delta": "done", "usage": None}
            yield {
                "delta": "",
                "usage": {"prompt_tokens": 10, "completion_tokens": 1},
            }

    events = await _collect_events(
        _orchestrator_with([ToolActivityProvider("tool-activity")]),
        ["tool-activity"],
    )
    statuses = [
        event for event in events
        if event["type"] == "provider_status"
    ]

    assert [event["phase"] for event in statuses] == [
        "requesting",
        "tool",
        "finalizing",
    ]
    assert statuses[1]["detail"] == "模型已开始构造工具调用"
    assert statuses[1]["first_token_ms"] >= 0


@pytest.mark.asyncio
async def test_orchestrator_caps_only_local_provider_history():
    provider = _FakeProvider("local-window")
    provider.config.type = "local"
    orchestrator = _orchestrator_with([provider])
    orchestrator._history_loaded = True
    orchestrator.conversation_history = [
        ChatMessage(
            role="user" if index % 2 == 0 else "assistant",
            content=f"history-{index}",
        )
        for index in range(20)
    ]

    await _collect_events(orchestrator, ["local-window"])
    conversation = [
        message
        for message in provider.received_messages
        if message.role != "system"
    ]

    assert len(conversation) <= 8
    assert conversation[0].role == "user"
    assert conversation[-1].content == "offline orchestration"
    assert all(message.content != "history-0" for message in conversation)


@pytest.mark.asyncio
async def test_reasoning_summaries_are_emitted_for_every_selected_model():
    providers = [_FakeProvider("model-a"), _FakeProvider("model-b")]
    events = await _collect_events(
        _orchestrator_with(providers),
        ["model-a", "model-b"],
    )
    completed = {
        event["message"]["author_model"]: event["message"]
        for event in events
        if event["type"] == "message_complete"
        and event["message"]["role"] == "assistant"
    }
    assert set(completed) == {"model-a", "model-b"}
    assert all(
        [step["phase"] for step in message["reasoning_summary"]]
        == ["analysis", "final"]
        for message in completed.values()
    )
    live_models = {
        event["model"]
        for event in events
        if event["type"] == "reasoning_step"
    }
    assert live_models == {"model-a", "model-b"}


@pytest.mark.asyncio
async def test_project_plan_waits_for_confirmation_before_model_calls(tmp_path):
    database = ChatDatabase(tmp_path / "planned-confirmation.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "先展示计划，确认后再执行",
        "targets": ["planner-model"],
        "mode": "project",
        "execution_strategy": "planned",
        "plan_confirmation_required": True,
        "max_revision_cycles": 1,
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]

    class ConfirmationProvider(_FakeProvider):
        def __init__(self):
            super().__init__("planner-model", persona="Architect/Reviewer")
            self.calls = []

        async def stream_chat(self, messages, tools=None, **kwargs):
            system_text = "\n".join(
                message.content for message in messages if message.role == "system"
            )
            if "类型：verifier" in system_text:
                self.calls.append("verifier")
                content = json.dumps({
                    "verdict": "pass",
                    "summary": "候选满足目标",
                    "evidence": ["候选"],
                    "issues": [],
                }, ensure_ascii=False)
            elif "类型：synthesizer" in system_text:
                self.calls.append("synthesizer")
                content = "确认后生成的最终交付"
            else:
                self.calls.append("worker")
                content = "确认后执行的候选"
            yield {"delta": content, "usage": None}

    provider = ConfirmationProvider()
    draft_orchestrator = SessionOrchestrator(session_id, mode="project")
    draft_orchestrator.providers = {"planner-model": provider}
    draft_orchestrator.refresh_providers = lambda: None
    draft_orchestrator.run_service = run_service
    draft_orchestrator.begin_agent_run(run["id"])

    draft_events = []
    async for event in draft_orchestrator.stream_user_message(action):
        draft_events.append(event)
        run_service.observe(run["id"], event)
    waiting = run_service.finish(run["id"])

    assert waiting["status"] == "waiting_plan_confirmation"
    assert waiting["plan_status"] == "draft"
    assert provider.calls == []
    assert any(
        event["type"] == "plan_confirmation_required"
        for event in draft_events
    )
    assert not any(
        event["type"] == "plan_step_update"
        for event in draft_events
    )

    confirmed_action, _ = run_service.prepare_confirmed_plan(
        run_id=run["id"],
        session_id=session_id,
        expected_revision=1,
    )
    confirmed_action["_agent_run_id"] = run["id"]
    executing = SessionOrchestrator(session_id, mode="project")
    executing.providers = {"planner-model": provider}
    executing.refresh_providers = lambda: None
    executing.run_service = run_service
    executing.begin_agent_run(run["id"])

    execution_events = []
    async for event in executing.stream_user_message(confirmed_action):
        execution_events.append(event)
        run_service.observe(run["id"], event)
    completed = run_service.finish(run["id"])

    assert completed["status"] == "completed_with_warnings"
    assert completed["plan_status"] == "completed"
    assert provider.calls == ["worker", "verifier", "synthesizer"]
    assert len([
        event for event in execution_events
        if event["type"] == "message_complete"
        and event["message"]["role"] == "assistant"
    ]) == 1


@pytest.mark.asyncio
async def test_project_mode_executes_worker_verifier_synthesizer_dag(tmp_path):
    database = ChatDatabase(tmp_path / "planned-orchestrator.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "形成一份经过双模型核验的迁移方案",
        "targets": ["researcher", "reviewer"],
        "mode": "project",
        "execution_strategy": "planned",
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]

    class PlannedProvider(_FakeProvider):
        def __init__(self, name, persona):
            super().__init__(name, persona=persona)
            self.calls = []
            self.saw_dependency_bundle = False
            self.tool_policies = {}

        async def stream_chat(self, messages, tools=None, **kwargs):
            system_text = "\n".join(
                message.content for message in messages if message.role == "system"
            )
            if "类型：verifier" in system_text:
                phase = "verifier"
                self.saw_dependency_bundle = all(
                    value in "\n".join(message.content for message in messages)
                    for value in ("候选-researcher", "候选-reviewer")
                )
                content = json.dumps({
                    "verdict": "pass",
                    "summary": "两份候选互相补充且未发现冲突",
                    "evidence": ["候选-researcher", "候选-reviewer"],
                    "issues": [],
                }, ensure_ascii=False)
            elif "类型：synthesizer" in system_text:
                phase = "synthesizer"
                combined = "\n".join(message.content for message in messages)
                self.saw_dependency_bundle = (
                    "候选-researcher" in combined
                    and "候选-reviewer" in combined
                    and '"verdict": "pass"' in combined
                )
                content = "最终迁移方案：先验证兼容性，再分阶段迁移，并保留回滚证据。"
            else:
                phase = "worker"
                content = f"候选-{self.config.name}"
            self.calls.append(phase)
            tool_names = {tool.name for tool in (tools or [])}
            self.tool_policies[phase] = {
                "has_list_tasks": "list_tasks" in tool_names,
                "has_write_artifact": "write_artifact" in tool_names,
            }
            yield {"delta": content, "usage": None}
            yield {
                "delta": "",
                "usage": {"prompt_tokens": 10, "completion_tokens": 4},
            }

    researcher = PlannedProvider("researcher", "Researcher")
    reviewer = PlannedProvider("reviewer", "Architect/Reviewer")
    orchestrator = SessionOrchestrator(session_id, mode="project")
    orchestrator.providers = {
        "researcher": researcher,
        "reviewer": reviewer,
    }
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])

    events = []
    async for event in orchestrator.stream_user_message(action):
        events.append(event)
        run_service.observe(run["id"], event)
    completed = run_service.finish(run["id"])

    assert completed["status"] == "completed_with_warnings"
    assert researcher.calls == ["worker"]
    assert reviewer.calls == ["worker", "verifier", "synthesizer"]
    assert reviewer.saw_dependency_bundle is True
    assert researcher.tool_policies["worker"] == {
        "has_list_tasks": True,
        "has_write_artifact": False,
    }
    assert reviewer.tool_policies["verifier"]["has_write_artifact"] is False
    assert reviewer.tool_policies["synthesizer"]["has_write_artifact"] is True
    assert len([
        event for event in events
        if event["type"] == "message_complete"
        and event["message"]["role"] == "assistant"
    ]) == 1
    assert len([
        event for event in events if event["type"] == "plan_step_output"
    ]) == 3
    final_message = next(
        event["message"]
        for event in events
        if event["type"] == "message_complete"
        and event["message"]["role"] == "assistant"
    )
    assert final_message["content"].startswith("最终迁移方案")

    detail = database.get_agent_run(run["id"])
    assert {step["status"] for step in detail["plan_steps"]} == {"completed"}
    assert detail["review_count"] == 1
    assert detail["reviews"][0]["verdict"] == "pass"
    assert detail["reviews"][0]["evidence"] == [
        "候选-researcher", "候选-reviewer",
    ]
    assert any(event["type"] == "plan_created" for event in events)
    assert any(event["type"] == "review_recorded" for event in events)


@pytest.mark.asyncio
async def test_planned_code_run_uses_scoped_worktree_tools_and_real_evidence(
    tmp_path,
    monkeypatch,
):
    repository_root = tmp_path / "orchestrated-repository"
    _initialize_test_git_repository(repository_root)
    database = ChatDatabase(tmp_path / "planned-code.sqlite3")
    workspace_service = GitWorkspaceService(
        database,
        worktree_root=tmp_path / "managed-worktrees",
    )
    run_service = AgentRunService(database, workspace_service)
    monkeypatch.setattr(
        git_workspace_module,
        "git_workspace_service",
        workspace_service,
    )
    session_id = str(uuid4())
    registered = workspace_service.register_repository(
        session_id=session_id,
        selected_path=str(repository_root),
    )
    action = {
        "content": "在隔离 worktree 中把 VALUE 改为 3",
        "targets": ["code-agent"],
        "mode": "project",
        "execution_strategy": "planned",
        "repository_id": registered["id"],
        "verification_commands": ["python -m pytest -q"],
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]

    async def passing_sandbox(*args, **kwargs):
        return {
            "stdout": "1 passed\n",
            "exit_code": 0,
            "timed_out": False,
            "network": "none",
        }

    monkeypatch.setattr(
        git_workspace_module,
        "execute_in_sandbox",
        passing_sandbox,
    )

    class CodeProvider(_FakeProvider):
        def __init__(self):
            super().__init__("code-agent", persona="Implementer/Reviewer")
            self.worker_turns = 0
            self.saw_deterministic_evidence = False

        async def stream_chat(self, messages, tools=None, **kwargs):
            system_text = "\n".join(
                message.content for message in messages if message.role == "system"
            )
            combined = "\n".join(message.content for message in messages)
            tool_names = {tool.name for tool in (tools or [])}
            if "类型：worker" in system_text:
                self.worker_turns += 1
                assert {
                    "workspace_list_files",
                    "workspace_read_file",
                    "workspace_write_file",
                    "workspace_run_command",
                    "workspace_status",
                } <= tool_names
                assert "write_artifact" not in tool_names
                if self.worker_turns == 1:
                    yield {
                        "delta": "",
                        "usage": None,
                        "tool_calls": [{
                            "id": "write-isolated-worktree",
                            "name": "workspace_write_file",
                            "arguments": {
                                "path": "app.py",
                                "content": "VALUE = 3\n",
                            },
                        }],
                    }
                else:
                    assert any(
                        message.role == "tool"
                        and message.name == "workspace_write_file"
                        and not message.tool_error
                        for message in messages
                    )
                    yield {"delta": "已在隔离 worktree 修改 app.py", "usage": None}
                return
            if "类型：verifier" in system_text:
                self.saw_deterministic_evidence = (
                    "deterministic_workspace_verification" not in combined
                    and '"verification_status": "passed"' in combined
                    and '"exit_code": 0' in combined
                )
                yield {
                    "delta": json.dumps({
                        "verdict": "pass",
                        "summary": "diff 与测试证据通过",
                        "evidence": ["AetherChat deterministic verifier"],
                        "issues": [],
                    }, ensure_ascii=False),
                    "usage": None,
                }
                return
            assert "类型：synthesizer" in system_text
            yield {"delta": "候选已核验，等待用户审核并应用 diff。", "usage": None}

    provider = CodeProvider()
    orchestrator = SessionOrchestrator(session_id, mode="project")
    orchestrator.providers = {"code-agent": provider}
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])

    events = []
    async for event in orchestrator.stream_user_message(action):
        events.append(event)
        run_service.observe(run["id"], event)
    completed = run_service.finish(run["id"])

    assert completed["status"] == "completed_verified"
    assert provider.worker_turns == 2
    assert provider.saw_deterministic_evidence is True
    assert any(event["type"] == "workspace_verified" for event in events)
    candidates = workspace_service.list_run_workspaces(
        run_id=run["id"],
        session_id=session_id,
    )
    assert len(candidates) == 1
    assert candidates[0]["verification_status"] == "passed"
    assert candidates[0]["changed_files"] == [{
        "status": " M",
        "path": "app.py",
    }]
    assert (repository_root / "app.py").read_text(encoding="utf-8") == "VALUE = 1\n"


@pytest.mark.asyncio
async def test_verifier_needs_revision_triggers_one_bounded_rework_cycle(tmp_path):
    database = ChatDatabase(tmp_path / "planned-review-revision.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "形成经过返工核验的发布方案",
        "targets": ["lead"],
        "mode": "project",
        "execution_strategy": "planned",
        "max_revision_cycles": 1,
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]

    class RevisionProvider(_FakeProvider):
        def __init__(self):
            super().__init__("lead", persona="Architect/Reviewer")
            self.calls = []
            self.worker_count = 0
            self.verifier_count = 0
            self.saw_revision_context = False

        async def stream_chat(self, messages, tools=None, **kwargs):
            system_text = "\n".join(
                message.content for message in messages if message.role == "system"
            )
            combined = "\n".join(message.content for message in messages)
            if "类型：verifier" in system_text:
                self.calls.append("verifier")
                self.verifier_count += 1
                if self.verifier_count == 1:
                    content = json.dumps({
                        "verdict": "needs_revision",
                        "summary": "缺少可验证的回滚条件",
                        "evidence": ["候选-v1"],
                        "issues": ["补充回滚触发条件"],
                    }, ensure_ascii=False)
                else:
                    assert "候选-v2" in combined
                    content = json.dumps({
                        "verdict": "pass",
                        "summary": "返工版本补齐回滚条件",
                        "evidence": ["候选-v2"],
                        "issues": [],
                    }, ensure_ascii=False)
            elif "类型：synthesizer" in system_text:
                self.calls.append("synthesizer")
                assert '"verdict": "pass"' in combined
                content = "最终发布方案：包含回滚触发条件。"
            else:
                self.calls.append("worker")
                self.worker_count += 1
                if self.worker_count == 2:
                    self.saw_revision_context = (
                        "候选-v1" in combined
                        and "补充回滚触发条件" in combined
                        and "受控返工" in combined
                    )
                content = f"候选-v{self.worker_count}"
            yield {"delta": content, "usage": None}

    provider = RevisionProvider()
    orchestrator = SessionOrchestrator(session_id, mode="project")
    orchestrator.providers = {"lead": provider}
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])

    events = []
    async for event in orchestrator.stream_user_message(action):
        events.append(event)
        run_service.observe(run["id"], event)
    completed = run_service.finish(run["id"])

    assert completed["status"] == "completed_with_warnings"
    assert completed["plan_revision"] == 2
    assert completed["automatic_revision_count"] == 1
    assert provider.calls == [
        "worker", "verifier", "worker", "verifier", "synthesizer",
    ]
    assert provider.saw_revision_context is True
    detail = database.get_agent_run(run["id"])
    assert [review["verdict"] for review in detail["reviews"]] == [
        "needs_revision", "pass",
    ]
    assert [
        revision["source"] for revision in detail["plan_revisions"]
    ] == ["verifier_replan", "planner"]
    assert any(event["type"] == "plan_replanned" for event in events)
    assert len([
        event for event in events
        if event["type"] == "message_complete"
        and event["message"]["role"] == "assistant"
    ]) == 1


@pytest.mark.asyncio
async def test_repeated_verifier_failure_blocks_unverified_synthesis(tmp_path):
    database = ChatDatabase(tmp_path / "planned-review-fail.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "核验失败时不得伪造最终交付",
        "targets": ["lead"],
        "mode": "project",
        "execution_strategy": "planned",
        "max_revision_cycles": 1,
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]

    class FailingVerifierProvider(_FakeProvider):
        def __init__(self):
            super().__init__("lead", persona="Architect/Reviewer")
            self.calls = []

        async def stream_chat(self, messages, tools=None, **kwargs):
            system_text = "\n".join(
                message.content for message in messages if message.role == "system"
            )
            if "类型：verifier" in system_text:
                self.calls.append("verifier")
                content = json.dumps({
                    "verdict": "fail",
                    "summary": "关键证据仍不可验证",
                    "evidence": [],
                    "issues": ["缺少真实测试证据"],
                }, ensure_ascii=False)
            elif "类型：synthesizer" in system_text:
                self.calls.append("synthesizer")
                content = "此内容不应生成"
            else:
                self.calls.append("worker")
                content = "缺少证据的候选"
            yield {"delta": content, "usage": None}

    provider = FailingVerifierProvider()
    orchestrator = SessionOrchestrator(session_id, mode="project")
    orchestrator.providers = {"lead": provider}
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])

    events = []
    async for event in orchestrator.stream_user_message(action):
        events.append(event)
        run_service.observe(run["id"], event)
    failed = run_service.finish(run["id"])

    assert failed["status"] == "failed"
    assert failed["automatic_revision_count"] == 1
    assert provider.calls == ["worker", "verifier", "worker", "verifier"]
    synthesizer = next(
        step for step in failed["plan_steps"]
        if step["kind"] == "synthesizer"
    )
    assert synthesizer["status"] == "blocked"
    assert any(
        event.get("code") == "plan_verification_failed"
        for event in events
    )
    assert not any(
        event["type"] == "message_complete"
        and event["message"]["role"] == "assistant"
        for event in events
    )


@pytest.mark.asyncio
async def test_project_plan_retry_reuses_completed_worker_outputs(tmp_path):
    database = ChatDatabase(tmp_path / "planned-retry.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "先独立分析，再核验并综合",
        "targets": ["worker-a", "reviewer"],
        "mode": "project",
        "execution_strategy": "planned",
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]

    class FirstAttemptProvider(_FakeProvider):
        def __init__(self, name, persona):
            super().__init__(name, persona=persona)
            self.calls = []

        async def stream_chat(self, messages, tools=None, **kwargs):
            system_text = "\n".join(
                message.content for message in messages if message.role == "system"
            )
            if "类型：verifier" in system_text:
                self.calls.append("verifier")
                if False:
                    yield {}
                raise RuntimeError("verifier offline")
            if "类型：synthesizer" in system_text:
                self.calls.append("synthesizer")
                if False:
                    yield {}
                raise RuntimeError("synthesizer offline")
            self.calls.append("worker")
            yield {"delta": f"已完成-{self.config.name}", "usage": None}

    worker = FirstAttemptProvider("worker-a", "Researcher")
    reviewer = FirstAttemptProvider("reviewer", "Architect/Reviewer")
    first = SessionOrchestrator(session_id, mode="project")
    first.providers = {"worker-a": worker, "reviewer": reviewer}
    first.refresh_providers = lambda: None
    first.run_service = run_service
    first.begin_agent_run(run["id"])

    first_events = []
    async for event in first.stream_user_message(action):
        first_events.append(event)
        run_service.observe(run["id"], event)
    failed = run_service.finish(run["id"])

    assert failed["status"] == "failed"
    assert worker.calls == ["worker"]
    assert reviewer.calls == ["worker", "verifier", "synthesizer"]
    before_retry = database.list_agent_plan_steps(run["id"])
    assert [
        step["status"] for step in before_retry if step["kind"] == "worker"
    ] == ["completed", "completed"]

    retry_action, _ = run_service.prepare_retry(run["id"], session_id)
    retry_action["_agent_run_id"] = run["id"]

    class RetryReviewer(_FakeProvider):
        def __init__(self):
            super().__init__("reviewer", persona="Architect/Reviewer")
            self.calls = []
            self.saw_reused_workers = False

        async def stream_chat(self, messages, tools=None, **kwargs):
            system_text = "\n".join(
                message.content for message in messages if message.role == "system"
            )
            combined = "\n".join(message.content for message in messages)
            self.saw_reused_workers = (
                "已完成-worker-a" in combined
                and "已完成-reviewer" in combined
            )
            if "类型：verifier" in system_text:
                self.calls.append("verifier")
                content = json.dumps({
                    "verdict": "pass",
                    "summary": "复用的两份 Worker 输出可用于综合",
                    "evidence": ["worker-a", "reviewer"],
                    "issues": [],
                }, ensure_ascii=False)
            else:
                assert "类型：synthesizer" in system_text
                self.calls.append("synthesizer")
                content = "重试后最终交付"
            yield {"delta": content, "usage": None}

    retry_reviewer = RetryReviewer()
    resumed = SessionOrchestrator(session_id, mode="project")
    resumed.providers = {"reviewer": retry_reviewer}
    resumed.refresh_providers = lambda: None
    resumed.run_service = run_service
    resumed.begin_agent_run(run["id"], retry_action.get("_pending_steers"))

    retry_events = []
    async for event in resumed.stream_user_message(retry_action):
        retry_events.append(event)
        run_service.observe(run["id"], event)
    completed = run_service.finish(run["id"])

    assert completed["status"] == "completed_with_warnings"
    assert retry_reviewer.calls == ["verifier", "synthesizer"]
    assert retry_reviewer.saw_reused_workers is True
    assert all(
        step["status"] == "completed"
        for step in database.list_agent_plan_steps(run["id"])
    )
    assert any(
        event["type"] == "message_complete"
        and event["message"]["content"] == "重试后最终交付"
        for event in retry_events
    )


@pytest.mark.asyncio
async def test_project_worker_cannot_execute_side_effect_tool(tmp_path):
    database = ChatDatabase(tmp_path / "planned-read-only-worker.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "先分析，最终阶段再决定是否写入",
        "targets": ["solo"],
        "mode": "project",
        "execution_strategy": "planned",
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]

    class ReadOnlyWorkerProvider(_FakeProvider):
        def __init__(self):
            super().__init__("solo", persona="Architect/Reviewer")
            self.worker_calls = 0

        async def stream_chat(self, messages, tools=None, **kwargs):
            system_text = "\n".join(
                message.content for message in messages if message.role == "system"
            )
            if "类型：worker" in system_text:
                self.worker_calls += 1
                assert "write_artifact" not in {tool.name for tool in (tools or [])}
                if self.worker_calls == 1:
                    yield {
                        "delta": "",
                        "usage": None,
                        "tool_calls": [{
                            "id": "forbidden-worker-write",
                            "name": "write_artifact",
                            "arguments": {
                                "path": "must-not-write.md",
                                "content": "forbidden",
                            },
                        }],
                    }
                    return
                observation = next(
                    message for message in reversed(messages)
                    if message.role == "tool"
                )
                assert observation.tool_error is True
                assert "只读 Worker/Verifier" in observation.content
                yield {"delta": "只读分析已完成", "usage": None}
                return
            if "类型：verifier" in system_text:
                yield {
                    "delta": json.dumps({
                        "verdict": "pass",
                        "summary": "Worker 未执行副作用",
                        "evidence": [],
                        "issues": [],
                    }, ensure_ascii=False),
                    "usage": None,
                }
                return
            assert "类型：synthesizer" in system_text
            yield {"delta": "最终未写入任何工件。", "usage": None}

    provider = ReadOnlyWorkerProvider()
    orchestrator = SessionOrchestrator(session_id, mode="project")
    orchestrator.providers = {"solo": provider}
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])
    tool_runs_before = len(tool_registry.list_runs(session_id=session_id, limit=None))

    async for event in orchestrator.stream_user_message(action):
        run_service.observe(run["id"], event)
    completed = run_service.finish(run["id"])

    assert completed["status"] == "completed_with_warnings"
    assert provider.worker_calls == 2
    assert len(tool_registry.list_runs(session_id=session_id, limit=None)) == tool_runs_before
    calls = database.list_agent_tool_calls(run["id"])
    assert len(calls) == 1
    assert calls[0]["status"] == "denied"
    assert calls[0]["plan_step_id"]


@pytest.mark.asyncio
async def test_interrupt_all_stops_remaining_project_plan_phases(tmp_path):
    database = ChatDatabase(tmp_path / "planned-interrupt.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "执行一个可中断的项目计划",
        "targets": ["slow-plan"],
        "mode": "project",
        "execution_strategy": "planned",
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]
    started = asyncio.Event()

    class SlowPlanProvider(_FakeProvider):
        def __init__(self):
            super().__init__("slow-plan", persona="Architect/Reviewer")
            self.phases = []

        async def stream_chat(self, messages, tools=None, **kwargs):
            system_text = "\n".join(
                message.content for message in messages if message.role == "system"
            )
            phase = (
                "worker" if "类型：worker" in system_text
                else "verifier" if "类型：verifier" in system_text
                else "synthesizer"
            )
            self.phases.append(phase)
            started.set()
            await asyncio.sleep(10)
            yield {"delta": "不应完成", "usage": None}

    provider = SlowPlanProvider()
    orchestrator = SessionOrchestrator(session_id, mode="project")
    orchestrator.providers = {"slow-plan": provider}
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])

    async def collect():
        events = []
        async for event in orchestrator.stream_user_message(action):
            events.append(event)
            run_service.observe(run["id"], event)
        return events

    collecting = asyncio.create_task(collect())
    await asyncio.wait_for(started.wait(), timeout=1)
    orchestrator.interrupt("all")
    events = await asyncio.wait_for(collecting, timeout=2)
    interrupted = run_service.finish(run["id"])

    assert interrupted["status"] == "interrupted"
    assert provider.phases == ["worker"]
    assert not any(event.get("code") == "plan_incomplete" for event in events)
    plan_steps = database.list_agent_plan_steps(run["id"])
    assert next(
        step for step in plan_steps if step["kind"] == "synthesizer"
    )["status"] == "interrupted"


@pytest.mark.asyncio
async def test_project_plan_does_not_synthesize_when_all_workers_fail(tmp_path):
    database = ChatDatabase(tmp_path / "planned-all-workers-fail.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "不能在没有候选结果时伪造最终交付",
        "targets": ["offline-worker"],
        "mode": "project",
        "execution_strategy": "planned",
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]

    class OfflineWorker(_FakeProvider):
        def __init__(self):
            super().__init__("offline-worker", persona="Architect/Reviewer")
            self.calls = 0

        async def stream_chat(self, messages, tools=None, **kwargs):
            self.calls += 1
            if False:
                yield {}
            raise RuntimeError("offline")

    provider = OfflineWorker()
    orchestrator = SessionOrchestrator(session_id, mode="project")
    orchestrator.providers = {"offline-worker": provider}
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])

    events = []
    async for event in orchestrator.stream_user_message(action):
        events.append(event)
        run_service.observe(run["id"], event)
    failed = run_service.finish(run["id"])

    assert failed["status"] == "failed"
    assert provider.calls == 1
    plan_steps = database.list_agent_plan_steps(run["id"])
    assert next(
        step for step in plan_steps if step["kind"] == "verifier"
    )["status"] == "blocked"
    assert next(
        step for step in plan_steps if step["kind"] == "synthesizer"
    )["status"] == "blocked"
    assert not any(
        event["type"] == "message_complete"
        and event["message"]["role"] == "assistant"
        for event in events
    )


@pytest.mark.asyncio
async def test_orchestrator_injects_configured_agent_persona():
    provider = _FakeProvider(
        "architect",
        persona="架构评审；优先识别边界、风险和验收条件",
    )
    orchestrator = _orchestrator_with([provider])

    await _collect_events(orchestrator, ["architect"])

    system_text = "\n".join(
        message.content for message in provider.received_messages if message.role == "system"
    )
    assert "稳定标识是：architect" in system_text
    assert "架构评审；优先识别边界、风险和验收条件" in system_text


class _ToolLoopProvider(_FakeProvider):
    def __init__(self):
        super().__init__("tool-loop")
        self.calls = 0
        self.received_observation = False

    async def stream_chat(self, messages, tools=None, **kwargs):
        self.calls += 1
        if self.calls == 1:
            yield {
                "delta": "让我先检查历史请求；这段工具轮草稿不应显示。",
                "usage": None,
                "tool_calls": [{
                    "id": "call-1",
                    "name": "plugin__desktop_control__capabilities",
                    "arguments": {},
                }],
            }
            yield {"delta": "", "usage": {"prompt_tokens": 10, "completion_tokens": 2}}
            return
        observation = next(message for message in reversed(messages) if message.role == "tool")
        assistant_call = next(
            message for message in reversed(messages)
            if message.role == "assistant" and message.tool_calls
        )
        assert [call.id for call in assistant_call.tool_calls] == ["call-1"]
        assert assistant_call.tool_calls[0].name == "plugin__desktop_control__capabilities"
        assert observation.tool_call_id == "call-1"
        assert observation.name == "plugin__desktop_control__capabilities"
        assert observation.tool_error is False
        self.received_observation = "desktop-control" in observation.content
        assert observation.images and observation.images[0].startswith("data:image/png;base64,")
        yield {"delta": "闭环完成", "usage": None}
        yield {"delta": "", "usage": {"prompt_tokens": 12, "completion_tokens": 3}}


class _PseudoToolProvider(_FakeProvider):
    def __init__(self):
        super().__init__("pseudo-tool")
        self.calls = 0

    async def stream_chat(self, messages, tools=None, **kwargs):
        self.calls += 1
        if self.calls == 1:
            yield {
                "delta": "让我先检查。\n\n请求执行工具：plugin__documents__inventory",
                "usage": None,
            }
        elif self.calls == 2:
            assert "立即发起结构化工具调用" in messages[-1].content
            yield {
                "delta": "",
                "usage": None,
                "tool_calls": [{
                    "id": "recovered-call",
                    "name": "plugin__documents__inventory",
                    "arguments": {},
                }],
            }
        else:
            assert any(message.role == "tool" for message in messages)
            yield {"delta": "已完成真实工具调用。", "usage": None}
        yield {"delta": "", "usage": {"prompt_tokens": 10, "completion_tokens": 2}}


class _ApprovalToolProvider(_FakeProvider):
    def __init__(self):
        super().__init__("approval-model")
        self.calls = 0

    async def stream_chat(self, messages, tools=None, **kwargs):
        self.calls += 1
        if self.calls == 1:
            yield {
                "delta": "",
                "usage": None,
                "tool_calls": [{
                    "id": "approval-call",
                    "name": "write_artifact",
                    "arguments": {"path": "approved.md", "content": "sensitive body"},
                }],
            }
        else:
            observation = next(message for message in reversed(messages) if message.role == "tool")
            assert '"status": "success"' in observation.content
            yield {"delta": "授权后已完成", "usage": None}
        yield {"delta": "", "usage": {"prompt_tokens": 10, "completion_tokens": 2}}


class _RejectedApprovalToolProvider(_FakeProvider):
    def __init__(self):
        super().__init__("rejected-approval-model")
        self.calls = 0
        self.received_feedback = False

    async def stream_chat(self, messages, tools=None, **kwargs):
        self.calls += 1
        if self.calls == 1:
            yield {
                "delta": "",
                "usage": None,
                "tool_calls": [{
                    "id": "rejected-approval-call",
                    "name": "write_artifact",
                    "arguments": {"path": "denied.md", "content": "must not be written"},
                }],
            }
        else:
            observation = next(
                message for message in reversed(messages)
                if message.role == "tool"
            )
            self.received_feedback = (
                '"status": "denied"' in observation.content
                and "不要写文件，改为只读总结" in observation.content
            )
            yield {"delta": "已按反馈改为只读总结", "usage": None}
        yield {"delta": "", "usage": {"prompt_tokens": 10, "completion_tokens": 2}}


class _ParallelToolProvider(_FakeProvider):
    def __init__(self):
        super().__init__("parallel-tools")
        self.calls = 0
        self.received_native_results = False

    async def stream_chat(self, messages, tools=None, **kwargs):
        self.calls += 1
        if self.calls == 1:
            yield {
                "delta": "",
                "usage": None,
                "tool_calls": [
                    {
                        "id": "call-documents",
                        "name": "plugin__documents__inventory",
                        "arguments": {},
                    },
                    {
                        "id": "call-desktop",
                        "name": "plugin__desktop_control__capabilities",
                        "arguments": {},
                    },
                ],
            }
        else:
            assistant_call = next(
                message for message in reversed(messages)
                if message.role == "assistant" and message.tool_calls
            )
            assert [call.id for call in assistant_call.tool_calls] == [
                "call-documents",
                "call-desktop",
            ]
            results = {
                message.tool_call_id: message
                for message in messages
                if message.role == "tool" and message.tool_call_id
            }
            assert results["call-documents"].name == "plugin__documents__inventory"
            assert '"plugin_id": "documents"' in results["call-documents"].content
            assert results["call-desktop"].name == "plugin__desktop_control__capabilities"
            assert '"plugin_id": "desktop-control"' in results["call-desktop"].content
            self.received_native_results = True
            yield {"delta": "并行结果已正确关联", "usage": None}
        yield {"delta": "", "usage": {"prompt_tokens": 10, "completion_tokens": 2}}


@pytest.mark.asyncio
async def test_orchestrator_reenters_model_after_tool_and_attaches_observation_image(monkeypatch):
    provider = _ToolLoopProvider()
    orchestrator = _orchestrator_with([provider])

    async def fake_invoke(tool_name, session_id, caller, arguments, **_kwargs):
        return {
            "plugin_id": "desktop-control",
            "status": "success",
            "result": {
                "snapshot_id": "snapshot-1",
                "screenshot": {
                    "image_data_url": "data:image/png;base64,iVBORw0KGgo=",
                    "width": 1,
                    "height": 1,
                },
            },
        }

    monkeypatch.setattr(plugin_registry, "invoke_tool", fake_invoke)
    events = await _collect_events(orchestrator, ["tool-loop"])
    assert provider.calls == 2
    assert provider.received_observation is True
    assert any(event["type"] == "tool_run" for event in events)
    completed = next(
        event["message"]
        for event in events
        if event["type"] == "message_complete"
        and event["message"]["author_model"] == "tool-loop"
    )
    assert completed["content"] == "闭环完成"
    assert "让我先检查历史请求" not in completed["content"]
    assert not any(
        "让我先检查历史请求" in str(event.get("delta") or "")
        for event in events
    )
    assert any(
        event["type"] == "stream_reset"
        for event in events
    )
    assert any(
        event["type"] == "message_complete"
        and event["message"]["author_model"] == "tool-loop"
        and "闭环完成" in event["message"]["content"]
        for event in events
    )


@pytest.mark.asyncio
async def test_orchestrator_correlates_parallel_tool_results_by_call_id(monkeypatch):
    provider = _ParallelToolProvider()
    orchestrator = _orchestrator_with([provider])

    async def fake_invoke(tool_name, session_id, caller, arguments, **_kwargs):
        plugin_id = "documents" if "documents" in tool_name else "desktop-control"
        return {
            "plugin_id": plugin_id,
            "status": "success",
            "result": {"tool_name": tool_name},
        }

    monkeypatch.setattr(plugin_registry, "invoke_tool", fake_invoke)
    events = await _collect_events(orchestrator, ["parallel-tools"])

    assert provider.calls == 2
    assert provider.received_native_results is True
    assert len([event for event in events if event["type"] == "tool_run"]) == 2
    assert any(
        event["type"] == "message_complete"
        and event["message"]["content"] == "并行结果已正确关联"
        for event in events
    )


def test_orchestrator_repairs_missing_and_duplicate_tool_call_ids():
    calls = SessionOrchestrator._normalize_tool_calls([
        {"id": "stable-call", "name": "first", "arguments": {}},
        {"id": "stable-call", "name": "second", "arguments": {"value": 2}},
        {"id": "", "name": "third", "arguments": "raw"},
        {
            "id": "json-string",
            "name": "write_artifact",
            "arguments": '{"path":"parsed.md","content":"ok"}',
        },
    ])

    assert calls[0].id == "stable-call"
    assert len({call.id for call in calls}) == 4
    assert all(call.id for call in calls)
    assert calls[2].arguments == {"_raw": "raw"}
    assert calls[3].arguments == {"path": "parsed.md", "content": "ok"}


@pytest.mark.asyncio
async def test_orchestrator_stops_repeated_truncated_file_writes_before_execution():
    class TruncatedWriteProvider(_FakeProvider):
        def __init__(self):
            super().__init__("truncated-writer")
            self.calls = 0
            self.tool_names: list[set[str]] = []

        async def stream_chat(self, messages, tools=None, **kwargs):
            self.calls += 1
            self.tool_names.append({tool.name for tool in tools or []})
            yield {
                "delta": "",
                "usage": None,
                "tool_calls": [{
                    "id": f"truncated-{self.calls}",
                    "name": "plugin__desktop_control__write_document",
                    "arguments": {
                        "_raw": '{"path":"/tmp/long.md","data":{"text":"truncated'
                    },
                }],
            }
            yield {
                "delta": "",
                "usage": {"prompt_tokens": 10, "completion_tokens": 10},
            }

    provider = TruncatedWriteProvider()
    orchestrator = _orchestrator_with([provider])

    async def unexpected_approval(*_args, **_kwargs):
        raise AssertionError("invalid arguments must be rejected before approval")

    orchestrator._authorize_model_tool = unexpected_approval
    events = [
        event
        async for event in orchestrator.stream_user_message({
            "content": "把这篇长文保存到桌面文件",
            "targets": [provider.config.name],
            "mode": "casual",
        })
    ]

    assert provider.calls == 3
    assert "plugin__desktop_control__write_document" in provider.tool_names[0]
    assert "plugin__desktop_control__write_document" not in provider.tool_names[1]
    assert sum(event.get("type") == "approval_required" for event in events) == 0
    assert any(
        event.get("code") == "invalid_file_tool_arguments"
        for event in events
    )
    assert any(
        event.get("type") == "message_complete"
        and "避免反复覆盖或产生空文件" in event["message"]["content"]
        for event in events
    )


@pytest.mark.asyncio
async def test_orchestrator_resumes_from_checkpoint_without_repeating_side_effect(tmp_path):
    database = ChatDatabase(tmp_path / "orchestrator-checkpoint.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    run = run_service.create(session_id, {
        "content": "只写入一次 checkpoint.md",
        "targets": ["checkpoint-model"],
        "mode": "project",
    })

    class FailingAfterToolProvider(_FakeProvider):
        def __init__(self):
            super().__init__("checkpoint-model")
            self.calls = 0

        async def stream_chat(self, messages, tools=None, **kwargs):
            self.calls += 1
            if self.calls == 1:
                yield {
                    "delta": "",
                    "usage": None,
                    "tool_calls": [{
                        "id": "checkpoint-call",
                        "name": "write_artifact",
                        "arguments": {
                            "path": "checkpoint.md",
                            "content": "written once",
                        },
                    }],
                }
                yield {"delta": "", "usage": {"prompt_tokens": 4, "completion_tokens": 2}}
                return
            assert any(message.tool_call_id == "checkpoint-call" for message in messages)
            if False:
                yield {}
            raise RuntimeError("offline after completed tool round")

    first_provider = FailingAfterToolProvider()
    first = SessionOrchestrator(session_id)
    first.providers = {"checkpoint-model": first_provider}
    first.refresh_providers = lambda: None
    first.run_service = run_service
    first.begin_agent_run(run["id"])

    async def allow_tool(_queue, _model, _tool_name, _arguments):
        return True, ""

    first._authorize_model_tool = allow_tool
    tool_runs_before = len(tool_registry.list_runs(session_id=session_id, limit=None))
    first_events = []
    first_action = {
        "content": "只写入一次 checkpoint.md",
        "targets": ["checkpoint-model"],
        "mode": "project",
        "_agent_run_id": run["id"],
    }
    async for event in first.stream_user_message(first_action):
        first_events.append(event)
        run_service.observe(run["id"], event)
    failed = run_service.finish(run["id"])

    assert failed["status"] == "failed"
    assert any(event.get("code") == "provider_error" for event in first_events)
    assert len(tool_registry.list_runs(session_id=session_id, limit=None)) == tool_runs_before + 1
    assert database.latest_agent_checkpoint(run["id"], "checkpoint-model")["kind"] == "tool_round"

    class ResumeProvider(_FakeProvider):
        def __init__(self):
            super().__init__("checkpoint-model")
            self.calls = 0
            self.saw_checkpoint = False

        async def stream_chat(self, messages, tools=None, **kwargs):
            self.calls += 1
            self.saw_checkpoint = (
                any(message.tool_calls for message in messages)
                and any(message.tool_call_id == "checkpoint-call" for message in messages)
            )
            yield {"delta": "从检查点完成", "usage": None}
            yield {"delta": "", "usage": {"prompt_tokens": 3, "completion_tokens": 2}}

    retry_action, _ = run_service.prepare_retry(run["id"], session_id)
    retry_action["_agent_run_id"] = run["id"]
    resumed_provider = ResumeProvider()
    resumed = SessionOrchestrator(session_id)
    resumed.providers = {"checkpoint-model": resumed_provider}
    resumed.refresh_providers = lambda: None
    resumed.run_service = run_service
    resumed.begin_agent_run(run["id"], retry_action.get("_pending_steers"))

    async for event in resumed.stream_user_message(retry_action):
        run_service.observe(run["id"], event)
    completed = run_service.finish(run["id"])

    assert completed["status"] == "uncertain"
    assert resumed_provider.calls == 1
    assert resumed_provider.saw_checkpoint is True
    assert len(tool_registry.list_runs(session_id=session_id, limit=None)) == tool_runs_before + 1


@pytest.mark.asyncio
async def test_orchestrator_applies_steer_before_stale_tool_side_effect(tmp_path):
    database = ChatDatabase(tmp_path / "orchestrator-steer.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    run = run_service.create(session_id, {
        "content": "先准备旧方案",
        "targets": ["steer-model"],
    })
    started = asyncio.Event()
    release = asyncio.Event()

    class SteerProvider(_FakeProvider):
        def __init__(self):
            super().__init__("steer-model")
            self.calls = 0
            self.saw_steer = False

        async def stream_chat(self, messages, tools=None, **kwargs):
            self.calls += 1
            if self.calls == 1:
                yield {"delta": "旧方案", "usage": None}
                started.set()
                await release.wait()
                yield {
                    "delta": "",
                    "usage": None,
                    "tool_calls": [{
                        "id": "stale-write",
                        "name": "write_artifact",
                        "arguments": {
                            "path": "must-not-exist.md",
                            "content": "stale",
                        },
                    }],
                }
            else:
                self.saw_steer = any(
                    message.role == "user" and "改为只读分析" in message.content
                    for message in messages
                )
                yield {"delta": "已改为只读分析", "usage": None}
            yield {"delta": "", "usage": {"prompt_tokens": 2, "completion_tokens": 1}}

    provider = SteerProvider()
    orchestrator = SessionOrchestrator(session_id)
    orchestrator.providers = {"steer-model": provider}
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])
    action = {
        "content": "先准备旧方案",
        "targets": ["steer-model"],
        "_agent_run_id": run["id"],
    }

    async def collect():
        events = []
        async for event in orchestrator.stream_user_message(action):
            events.append(event)
            run_service.observe(run["id"], event)
        run_service.finish(run["id"])
        return events

    collecting = asyncio.create_task(collect())
    await asyncio.wait_for(started.wait(), timeout=1)
    steer = run_service.queue_steer(
        run_id=run["id"],
        session_id=session_id,
        content="改为只读分析，不要写文件",
        targets=["steer-model"],
    )
    orchestrator.queue_steer(steer)
    release.set()
    events = await asyncio.wait_for(collecting, timeout=2)

    assert provider.calls == 2
    assert provider.saw_steer is True
    assert any(event["type"] == "stream_reset" for event in events)
    assert any(event["type"] == "steer_applied" for event in events)
    assert not any(event["type"] == "tool_run" for event in events)
    assert database.list_agent_tool_calls(run["id"]) == []
    checkpoint = database.latest_agent_checkpoint(run["id"], "steer-model")
    assert checkpoint["kind"] == "completed"
    assert any(
        "改为只读分析" in message["content"]
        for message in checkpoint["state"]["messages"]
        if message["role"] == "user"
    )


@pytest.mark.asyncio
async def test_agent_run_drops_turn_limit_tool_draft_from_chat(tmp_path, monkeypatch):
    database = ChatDatabase(tmp_path / "orchestrator-turn-limit.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    run = run_service.create(session_id, {
        "content": "验证模型调用上限",
        "targets": ["turn-limit-model"],
    })

    class TurnLimitProvider(_FakeProvider):
        def __init__(self):
            super().__init__("turn-limit-model")

        async def stream_chat(self, messages, tools=None, **kwargs):
            yield {
                "delta": "尚未完成",
                "usage": None,
                "tool_calls": [{
                    "id": "turn-limit-call",
                    "name": "list_tasks",
                    "arguments": {},
                }],
            }
            yield {"delta": "", "usage": {"prompt_tokens": 2, "completion_tokens": 1}}

    monkeypatch.setattr(orchestrator_module, "MAX_MODEL_TURNS", 1)
    monkeypatch.setattr("app.services.orchestrator.run.MAX_MODEL_TURNS", 1)
    provider = TurnLimitProvider()
    orchestrator = SessionOrchestrator(session_id)
    orchestrator.providers = {"turn-limit-model": provider}
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])

    events = []
    async for event in orchestrator.stream_user_message({
        "content": "验证模型调用上限",
        "targets": ["turn-limit-model"],
        "_agent_run_id": run["id"],
    }):
        events.append(event)
        run_service.observe(run["id"], event)
    failed = run_service.finish(run["id"])

    assert failed["status"] == "failed"
    assert database.list_agent_steps(run["id"])[0]["status"] == "failed"
    assert any(event.get("code") == "model_turn_limit" for event in events)
    assert not any(
        event["type"] == "message_complete"
        and event["message"].get("role") == "assistant"
        for event in events
    )
    assert not any(event.get("delta") == "尚未完成" for event in events)


@pytest.mark.asyncio
async def test_orchestrator_recovers_textual_pseudo_tool_call():
    provider = _PseudoToolProvider()
    orchestrator = _orchestrator_with([provider])
    events = await _collect_events(orchestrator, ["pseudo-tool"])

    assert provider.calls == 3
    assert any(event["type"] == "stream_reset" for event in events)
    assert any(event["type"] == "tool_run" for event in events)
    completed = [
        event["message"]
        for event in events
        if event["type"] == "message_complete"
        and event["message"]["author_model"] == "pseudo-tool"
    ]
    assert completed[0]["content"] == "已完成真实工具调用。"
    assert "请求执行工具" not in completed[0]["content"]


@pytest.mark.asyncio
async def test_orchestrator_pauses_for_structured_chat_approval():
    provider = _ApprovalToolProvider()
    orchestrator = _orchestrator_with([provider])
    stream = orchestrator.stream_user_message({
        "content": "请更新项目说明",
        "targets": ["approval-model"],
        "mode": "project",
    })
    events = []
    async for event in stream:
        events.append(event)
        if event["type"] == "approval_required":
            approval = event["approval"]
            assert approval["tool_name"] == "write_artifact"
            assert approval["arguments"]["content"].startswith("[已隐藏")
            approval_service.decide(approval["id"], orchestrator.session_id, "allow_once")

    assert any(event["type"] == "approval_resolved" and event["decision"] == "allow_once" for event in events)
    assert any(event["type"] == "tool_run" and event["run"]["status"] == "success" for event in events)
    artifact_update = next(
        event for event in events
        if event["type"] == "artifact_update"
    )
    assert artifact_update["generated_files"][0]["path"] == "approved.md"
    assert any(
        event["type"] == "message_complete"
        and event["message"]["content"] == "授权后已完成"
        and event["message"]["generated_files"][0]["path"] == "approved.md"
        for event in events
    )


@pytest.mark.asyncio
async def test_orchestrator_returns_rejection_feedback_to_the_model():
    provider = _RejectedApprovalToolProvider()
    orchestrator = _orchestrator_with([provider])
    stream = orchestrator.stream_user_message({
        "content": "请更新项目说明",
        "targets": ["rejected-approval-model"],
        "mode": "project",
    })
    events = []
    async for event in stream:
        events.append(event)
        if event["type"] == "approval_required":
            approval_service.decide(
                event["approval"]["id"],
                orchestrator.session_id,
                "deny",
                "不要写文件，改为只读总结",
            )

    resolved = next(event for event in events if event["type"] == "approval_resolved")
    assert resolved["decision"] == "deny"
    assert resolved["feedback"] == "不要写文件，改为只读总结"
    assert provider.received_feedback is True
    assert any(
        event["type"] == "tool_run"
        and event["run"]["status"] == "error"
        and "不要写文件，改为只读总结" in event["run"]["error"]
        for event in events
    )
    assert any(
        event["type"] == "message_complete"
        and event["message"]["content"] == "已按反馈改为只读总结"
        for event in events
    )


@pytest.mark.asyncio
async def test_browser_navigation_requires_a_matching_desktop_receipt():
    requested_actions = []

    async def execute_desktop_action(action):
        requested_actions.append(action)
        return {
            "success": True,
            "kind": "external_navigation",
            "url": action["url"],
            "browser": action["browser"],
            "launcher": "firefox",
            "verification": "launcher_accepted",
            "accepted_at": "2026-08-01T10:05:52Z",
        }

    plugin_result = {
        "plugin_id": "browser-links",
        "plugin_name": "浏览器链接",
        "action": "open_url",
        "status": "success",
        "result": {
            "kind": "open_external",
            "url": "https://www.youtube.com",
            "browser": "firefox",
        },
    }
    orchestrator = SessionOrchestrator(
        str(uuid4()),
        desktop_action_handler=execute_desktop_action,
    )
    completed = await orchestrator._complete_desktop_action(
        "plugin__browser_links__open_url",
        plugin_result,
    )

    assert requested_actions == [{
        "kind": "open_external",
        "url": "https://www.youtube.com",
        "browser": "firefox",
    }]
    assert completed["result"]["execution_receipt"]["launcher"] == "firefox"
    assert SessionOrchestrator._requested_browser("请用 fierfox 打开") == "firefox"
    assert SessionOrchestrator._browser_execution_receipt({
        "status": "success",
        "result": plugin_result,
    }) is None

    without_desktop = SessionOrchestrator(str(uuid4()))
    with pytest.raises(RuntimeError, match="没有活动的 AetherChat 桌面客户端"):
        await without_desktop._complete_desktop_action(
            "plugin__browser_links__open_url",
            plugin_result,
        )


@pytest.mark.asyncio
async def test_orchestrator_uses_configured_cost_and_hard_stops_empty_budget():
    priced = _FakeProvider("priced", cost_in=1.0, cost_out=2.0)
    orchestrator = _orchestrator_with([priced])
    events = await _collect_events(orchestrator, ["priced"])

    complete = next(
        event["message"]
        for event in events
        if event["type"] == "message_complete"
        and event["message"]["author_model"] == "priced"
    )
    assert complete["cost_usd"] == 2.0
    assert orchestrator.budget_remaining == 8.0

    blocked = _FakeProvider("blocked", cost_in=1.0, cost_out=1.0)
    exhausted = _orchestrator_with([blocked])
    exhausted.budget_remaining = 0
    blocked_events = await _collect_events(exhausted, ["blocked"])
    assert blocked.called is False
    assert any(event.get("code") == "budget_exceeded" for event in blocked_events)


@pytest.mark.asyncio
async def test_provider_exception_finishes_and_single_interrupt_keeps_other_model_running():
    broken = _FakeProvider("broken", behavior="error")
    broken_orchestrator = _orchestrator_with([broken])
    broken_events = await asyncio.wait_for(
        _collect_events(broken_orchestrator, ["broken"]),
        timeout=1,
    )
    assert any(event.get("code") == "provider_error" for event in broken_events)

    slow_started = asyncio.Event()
    slow = _FakeProvider("slow", behavior="slow", started=slow_started)
    survivor = _FakeProvider("survivor", behavior="slow")
    orchestrator = _orchestrator_with([slow, survivor])
    collecting = asyncio.create_task(_collect_events(orchestrator, ["slow", "survivor"]))
    await asyncio.wait_for(slow_started.wait(), timeout=1)
    await asyncio.sleep(0.04)
    orchestrator.interrupt("slow")
    events = await asyncio.wait_for(collecting, timeout=2)

    assert any(event["type"] == "interrupted" and event["model"] == "slow" for event in events)
    assert not any(event["type"] == "interrupted" and event["model"] == "survivor" for event in events)
    assert any(
        event["type"] == "message_complete"
        and event["message"]["author_model"] == "survivor"
        for event in events
    )


def test_websocket_action_error_does_not_close_connection(client):
    session_id = str(uuid4())
    with client.websocket_connect(f"/ws/sessions/{session_id}") as websocket:
        websocket.send_text("{not-json")
        error = websocket.receive_json()
        assert error["type"] == "error"
        assert error["code"] == "invalid_action"
        assert error["request_id"]

        websocket.send_json({
            "type": "not_supported",
            "request_id": "ws-diagnostic-77",
        })
        unsupported = websocket.receive_json()
        assert unsupported["type"] == "error"
        assert unsupported["code"] == "unsupported_action"
        assert unsupported["request_id"] == "ws-diagnostic-77"

        websocket.send_json({"type": "interrupt", "scope": "all"})
        acknowledgement = websocket.receive_json()
        assert acknowledgement == {"type": "interrupt_ack", "scope": "all"}


def test_websocket_dispatcher_enforces_session_owner_across_connections(
    client,
    monkeypatch,
):
    released = threading.Event()

    async def slow_stream(self, action):
        yield {"type": "delta", "model": "owner-model", "delta": "started"}
        while not released.is_set():
            await asyncio.sleep(0.01)
        yield {"type": "interrupted", "model": "owner-model"}

    def record_interrupt(self, scope="all"):
        released.set()

    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", slow_stream)
    monkeypatch.setattr(SessionOrchestrator, "interrupt", record_interrupt)
    session_id = str(uuid4())

    with client.websocket_connect(f"/ws/sessions/{session_id}") as owner:
        owner.send_json({
            "type": "post_message",
            "content": "hold the session owner",
            "targets": ["owner-model"],
        })
        assert owner.receive_json() == {"type": "ack"}
        assert owner.receive_json()["type"] == "agent_run_update"
        assert owner.receive_json()["delta"] == "started"

        with client.websocket_connect(f"/ws/sessions/{session_id}") as contender:
            contender.send_json({
                "type": "post_message",
                "content": "must not create a second run",
                "targets": ["owner-model"],
            })
            conflict = contender.receive_json()
            assert conflict["type"] == "error"
            assert conflict["code"] == "run_in_progress"
            assert conflict["dispatch_state"] == "running"

        assert len(chat_database.list_agent_runs(session_id)) == 1
        owner.send_json({"type": "interrupt", "scope": "all"})
        assert owner.receive_json() == {"type": "interrupt_ack", "scope": "all"}


def test_websocket_dispatcher_queues_other_session_without_false_running_record(
    client,
    monkeypatch,
):
    released = threading.Event()
    first_session = str(uuid4())
    second_session = str(uuid4())

    async def capacity_stream(self, action):
        yield {
            "type": "delta",
            "model": "capacity-model",
            "delta": f"started:{self.session_id}",
        }
        if self.session_id == first_session:
            while not released.is_set():
                await asyncio.sleep(0.01)
        yield {
            "type": "message_complete",
            "message": {
                "id": str(uuid4()),
                "session_id": self.session_id,
                "role": "assistant",
                "author_model": "capacity-model",
                "content": "completed",
            },
        }

    def record_interrupt(self, scope="all"):
        if self.session_id == first_session:
            released.set()

    monkeypatch.setattr(run_dispatcher, "capacity", 1)
    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", capacity_stream)
    monkeypatch.setattr(SessionOrchestrator, "interrupt", record_interrupt)

    with client.websocket_connect(f"/ws/sessions/{first_session}") as first:
        first.send_json({
            "type": "post_message",
            "content": "occupy the only slot",
            "targets": ["capacity-model"],
        })
        assert first.receive_json() == {"type": "ack"}
        assert first.receive_json()["type"] == "agent_run_update"
        assert first.receive_json()["delta"] == f"started:{first_session}"

        with client.websocket_connect(f"/ws/sessions/{second_session}") as second:
            second.send_json({
                "type": "post_message",
                "content": "wait for the slot",
                "targets": ["capacity-model"],
            })
            assert second.receive_json() == {"type": "ack"}
            queued = second.receive_json()
            assert queued["type"] == "run_queued"
            assert queued["queue_position"] == 1
            assert queued["queue_capacity"] == run_dispatcher.queue_capacity
            assert chat_database.list_agent_runs(second_session) == []

            first.send_json({"type": "interrupt", "scope": "all"})
            assert first.receive_json() == {
                "type": "interrupt_ack",
                "scope": "all",
            }

            dispatch_started = second.receive_json()
            assert dispatch_started["type"] == "run_dispatch_started"
            assert dispatch_started["dispatch_id"] == queued["dispatch_id"]
            running = second.receive_json()
            assert running["type"] == "agent_run_update"
            assert running["agent_run"]["status"] == "running"
            assert second.receive_json()["delta"] == f"started:{second_session}"
            assert second.receive_json()["type"] == "message_complete"
            completed = second.receive_json()
            assert completed["type"] == "agent_run_update"
            assert completed["agent_run"]["status"] == "completed_with_warnings"

        assert len(chat_database.list_agent_runs(second_session)) == 1


def test_websocket_dispatcher_cancels_queued_run_before_persistence(
    client,
    monkeypatch,
):
    released = threading.Event()
    first_session = str(uuid4())
    queued_session = str(uuid4())

    async def slow_stream(self, action):
        yield {"type": "delta", "model": "queue-model", "delta": "started"}
        if self.session_id == first_session:
            while not released.is_set():
                await asyncio.sleep(0.01)

    def record_interrupt(self, scope="all"):
        if self.session_id == first_session:
            released.set()

    monkeypatch.setattr(run_dispatcher, "capacity", 1)
    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", slow_stream)
    monkeypatch.setattr(SessionOrchestrator, "interrupt", record_interrupt)

    with client.websocket_connect(f"/ws/sessions/{first_session}") as first:
        first.send_json({
            "type": "post_message",
            "content": "occupy capacity",
            "targets": ["queue-model"],
        })
        assert first.receive_json() == {"type": "ack"}
        assert first.receive_json()["type"] == "agent_run_update"
        assert first.receive_json()["type"] == "delta"

        with client.websocket_connect(f"/ws/sessions/{queued_session}") as queued_ws:
            queued_ws.send_json({
                "type": "post_message",
                "content": "cancel before persistence",
                "targets": ["queue-model"],
            })
            assert queued_ws.receive_json() == {"type": "ack"}
            queued = queued_ws.receive_json()
            assert queued["type"] == "run_queued"

            queued_ws.send_json({
                "type": "interrupt",
                "scope": "all",
                "dispatch_id": queued["dispatch_id"],
            })
            cancellation_events = [
                queued_ws.receive_json(),
                queued_ws.receive_json(),
            ]
            assert {event["type"] for event in cancellation_events} == {
                "interrupt_ack",
                "run_cancelled",
            }
            acknowledgement = next(
                event
                for event in cancellation_events
                if event["type"] == "interrupt_ack"
            )
            assert acknowledgement["queued"] is True
            assert acknowledgement["dispatch_id"] == queued["dispatch_id"]
            assert chat_database.list_agent_runs(queued_session) == []

        first.send_json({"type": "interrupt", "scope": "all"})
        assert first.receive_json() == {"type": "interrupt_ack", "scope": "all"}


def test_websocket_dispatcher_preserves_runtime_failure_as_action_failed(
    client,
    monkeypatch,
):
    async def failing_stream(self, action):
        if False:
            yield {"type": "unreachable"}
        raise RuntimeError("provider execution failed")

    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", failing_stream)
    session_id = str(uuid4())

    with client.websocket_connect(f"/ws/sessions/{session_id}") as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "preserve the execution error contract",
            "targets": ["failing-model"],
        })
        assert websocket.receive_json() == {"type": "ack"}
        assert websocket.receive_json()["type"] == "agent_run_update"
        error = websocket.receive_json()
        assert error["type"] == "error"
        assert error["code"] == "action_failed"
        assert error["message"] == "provider execution failed"


def test_websocket_v1_requires_explicit_project_execution_strategy(client, monkeypatch):
    async def contract_stream(self, action):
        assert action["execution_strategy"] == "planned"
        assert action["effective_execution_strategy"] == "planned"
        yield {"type": "v1_contract_observed"}

    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", contract_stream)
    session_id = str(uuid4())
    with client.websocket_connect(f"/ws/v1/sessions/{session_id}") as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "缺少执行策略",
            "targets": ["Qwen"],
            "mode": "project",
            "request_id": "v1-contract-test",
        })
        error = websocket.receive_json()
        assert error["type"] == "error"
        assert error["code"] == "invalid_action"
        assert error["request_id"] == "v1-contract-test"
        assert "execution_strategy" in error["message"]
        assert error["contract_version"] == "1"
        assert error["sequence"] == 1
        assert error["session_id"] == session_id
        assert error["run_id"] is None
        assert error["step_id"] is None
        assert error["event_id"]
        assert error["occurred_at"].endswith("Z")

        websocket.send_json({
            "type": "post_message",
            "content": "显式计划执行",
            "targets": ["Qwen"],
            "mode": "project",
            "execution_strategy": "planned",
            "request_id": "v1-planned-test",
        })
        acknowledgement = websocket.receive_json()
        assert acknowledgement["type"] == "ack"
        assert acknowledgement["contract_version"] == "1"
        assert acknowledgement["sequence"] == 2
        assert acknowledgement["request_id"] == "v1-planned-test"
        assert acknowledgement["session_id"] == session_id
        assert acknowledgement["run_id"]
        assert acknowledgement["step_id"] is None
        update = websocket.receive_json()
        assert update["type"] == "agent_run_update"
        assert update["contract_version"] == "1"
        assert update["sequence"] == 3
        assert update["request_id"] == "v1-planned-test"
        assert update["session_id"] == session_id
        assert update["run_id"] == acknowledgement["run_id"]
        assert update["agent_run"]["effective_execution_strategy"] == "planned"
        assert update["agent_run"]["execution_strategy_source"] == "explicit"
        observed = websocket.receive_json()
        assert observed["type"] == "v1_contract_observed"
        assert observed["contract_version"] == "1"
        assert observed["sequence"] == 4
        assert observed["request_id"] == "v1-planned-test"
        assert observed["session_id"] == session_id
        assert observed["run_id"] == acknowledgement["run_id"]
        trace_context = next(
            event
            for event in chat_database.get_agent_run(
                acknowledgement["run_id"]
            )["events"]
            if event["event_type"] == "trace_context"
        )
        assert trace_context["payload"] == {
            "request_ref": diagnostic_ref("request", "v1-planned-test")
        }
        assert len({
            error["event_id"],
            acknowledgement["event_id"],
            update["event_id"],
            observed["event_id"],
        }) == 4


def test_websocket_v1_replays_persisted_stream_cursor_across_connections(client):
    session_id = str(uuid4())
    stream_id = f"test-stream-{uuid4()}"
    path = f"/ws/v1/sessions/{session_id}?stream_id={stream_id}"

    with client.websocket_connect(f"{path}&cursor=0") as websocket:
        websocket.send_json({
            "type": "unsupported",
            "request_id": "persisted-event-1",
        })
        first = websocket.receive_json()
        assert first["type"] == "error"
        assert first["sequence"] == 1
        assert first["stream_id"] == stream_id

    with client.websocket_connect(f"{path}&cursor=0") as websocket:
        replayed = websocket.receive_json()
        assert replayed == first
        websocket.send_json({
            "type": "unsupported",
            "request_id": "persisted-event-2",
        })
        second = websocket.receive_json()
        assert second["sequence"] == 2
        assert second["stream_id"] == stream_id

    with client.websocket_connect(f"{path}&cursor=2") as websocket:
        websocket.send_json({
            "type": "unsupported",
            "request_id": "persisted-event-3",
        })
        third = websocket.receive_json()
        assert third["sequence"] == 3

    with client.websocket_connect(f"{path}&cursor=999") as websocket:
        with pytest.raises(WebSocketDisconnect) as closed:
            websocket.receive_json()
        assert closed.value.code == 4009


def test_websocket_v1_allows_only_one_active_writer_per_stream(client):
    session_id = str(uuid4())
    stream_id = f"exclusive-stream-{uuid4()}"
    path = f"/ws/v1/sessions/{session_id}?stream_id={stream_id}&cursor=0"

    with client.websocket_connect(path) as first, client.websocket_connect(path) as replacement:
        with pytest.raises(WebSocketDisconnect) as replaced:
            first.receive_json()
        assert replaced.value.code == 4009

        replacement.send_json({
            "type": "unsupported",
            "request_id": "exclusive-writer",
        })
        event = replacement.receive_json()
        assert event["sequence"] == 1
        assert event["stream_id"] == stream_id


def test_websocket_v1_coalesces_delta_bursts_and_legacy_keeps_fragments(
    client,
    monkeypatch,
):
    async def burst_stream(self, action):
        for fragment in ("A", "ether", "Chat"):
            yield {
                "type": "delta",
                "model": "burst-model",
                "delta": fragment,
                "ts": "2026-08-11T08:00:00Z",
            }
        yield {"type": "burst_boundary"}

    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", burst_stream)

    v1_session_id = str(uuid4())
    with client.websocket_connect(
        f"/ws/v1/sessions/{v1_session_id}"
    ) as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "测试突发增量",
            "targets": ["burst-model"],
            "mode": "project",
            "execution_strategy": "direct",
            "request_id": "delta-burst-v1",
        })
        assert websocket.receive_json()["type"] == "ack"
        assert websocket.receive_json()["type"] == "agent_run_update"
        delta = websocket.receive_json()
        boundary = websocket.receive_json()
        completed = websocket.receive_json()

    assert delta["type"] == "delta"
    assert delta["delta"] == "AetherChat"
    assert delta["coalesced_count"] == 3
    assert delta["request_id"] == "delta-burst-v1"
    assert [delta["sequence"], boundary["sequence"], completed["sequence"]] == [
        3,
        4,
        5,
    ]
    assert boundary["type"] == "burst_boundary"
    assert completed["type"] == "agent_run_update"

    legacy_session_id = str(uuid4())
    with client.websocket_connect(
        f"/ws/sessions/{legacy_session_id}"
    ) as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "验证旧版逐片发送",
            "targets": ["burst-model"],
        })
        assert websocket.receive_json()["type"] == "ack"
        assert websocket.receive_json()["type"] == "agent_run_update"
        legacy_deltas = [websocket.receive_json() for _index in range(3)]
        assert websocket.receive_json()["type"] == "burst_boundary"
        assert websocket.receive_json()["type"] == "agent_run_update"

    assert [event["delta"] for event in legacy_deltas] == ["A", "ether", "Chat"]
    assert all("coalesced_count" not in event for event in legacy_deltas)


def test_websocket_waits_for_desktop_navigation_receipt(client, monkeypatch):
    async def navigation_stream(self, action):
        assert self.desktop_action_handler is not None
        receipt = await self.desktop_action_handler({
            "kind": "open_external",
            "url": "https://www.youtube.com",
            "browser": "firefox",
        })
        yield {"type": "desktop_action_test", "receipt": receipt}

    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", navigation_stream)

    session_id = str(uuid4())
    with client.websocket_connect(f"/ws/sessions/{session_id}") as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "用 Firefox 打开 YouTube",
            "targets": ["Qwen"],
        })
        assert websocket.receive_json() == {"type": "ack"}
        assert websocket.receive_json()["type"] == "agent_run_update"
        request = websocket.receive_json()
        assert request["type"] == "desktop_action_required"
        assert request["desktop_action"] == {
            "kind": "open_external",
            "url": "https://www.youtube.com",
            "browser": "firefox",
        }

        websocket.send_json({
            "type": "desktop_action_result",
            "action_id": request["action_id"],
            "success": True,
            "receipt": {
                "success": True,
                "kind": "external_navigation",
                "url": "https://www.youtube.com/",
                "browser": "firefox",
                "launcher": "firefox",
                "verification": "launcher_accepted",
                "accepted_at": "2026-08-01T10:05:52Z",
            },
        })
        completed = websocket.receive_json()
        assert completed["type"] == "desktop_action_test"
        assert completed["receipt"] == {
            "success": True,
            "kind": "external_navigation",
            "url": "https://www.youtube.com",
            "browser": "firefox",
            "launcher": "firefox",
            "verification": "launcher_accepted",
            "accepted_at": "2026-08-01T10:05:52Z",
        }


def test_websocket_interrupt_is_processed_while_agent_run_is_active(client, monkeypatch):
    interrupted = threading.Event()

    async def slow_stream(self, action):
        yield {"type": "delta", "model": "slow-model", "delta": "started"}
        while not interrupted.is_set():
            await asyncio.sleep(0.01)
        yield {"type": "interrupted", "model": "slow-model"}

    def record_interrupt(self, scope="all"):
        interrupted.set()

    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", slow_stream)
    monkeypatch.setattr(SessionOrchestrator, "interrupt", record_interrupt)

    session_id = str(uuid4())
    with client.websocket_connect(f"/ws/sessions/{session_id}") as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "start a slow run",
            "targets": ["slow-model"],
        })
        assert websocket.receive_json() == {"type": "ack"}

        saw_active_delta = False
        for _ in range(5):
            event = websocket.receive_json()
            if event.get("type") == "delta":
                saw_active_delta = True
                break
        assert saw_active_delta is True

        websocket.send_json({"type": "interrupt", "scope": "all"})
        saw_ack = False
        for _ in range(6):
            event = websocket.receive_json()
            if event == {"type": "interrupt_ack", "scope": "all"}:
                saw_ack = True
                break
        assert saw_ack is True
        assert interrupted.wait(timeout=1)


def test_websocket_interrupt_can_target_persisted_step(client, monkeypatch):
    interrupted = threading.Event()
    interrupted_scope = []

    async def slow_stream(self, action):
        yield {"type": "delta", "model": "model-a", "delta": "a"}
        yield {"type": "delta", "model": "model-b", "delta": "b"}
        while not interrupted.is_set():
            await asyncio.sleep(0.01)
        yield {"type": "interrupted", "model": interrupted_scope[0]}

    def record_interrupt(self, scope="all"):
        interrupted_scope.append(scope)
        interrupted.set()

    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", slow_stream)
    monkeypatch.setattr(SessionOrchestrator, "interrupt", record_interrupt)

    session_id = str(uuid4())
    with client.websocket_connect(f"/ws/sessions/{session_id}") as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "run two steps",
            "targets": ["model-a", "model-b"],
        })
        assert websocket.receive_json() == {"type": "ack"}
        update = websocket.receive_json()
        assert update["type"] == "agent_run_update"
        run_id = update["agent_run"]["id"]
        assert websocket.receive_json()["model"] == "model-a"
        assert websocket.receive_json()["model"] == "model-b"

        detail = client.get(
            f"/agent-runs/{run_id}",
            params={"session_id": session_id},
        ).json()
        step_id = next(step["id"] for step in detail["steps"] if step["model"] == "model-a")
        websocket.send_json({
            "type": "interrupt",
            "run_id": run_id,
            "step_id": step_id,
        })
        assert websocket.receive_json() == {
            "type": "interrupt_ack",
            "scope": "model-a",
        }
        assert interrupted.wait(timeout=1)
        assert interrupted_scope == ["model-a"]


def test_websocket_interrupt_can_target_persisted_plan_step(client, monkeypatch):
    interrupted = threading.Event()
    interrupted_scope = []

    async def slow_plan_stream(self, action):
        plan = self.run_service.ensure_plan(
            run_id=action["_agent_run_id"],
            session_id=self.session_id,
            action={
                **action,
                "mode": "project",
                "execution_strategy": "planned",
            },
            models=["model-a"],
            profiles={"model-a": "Worker"},
        )
        worker = next(
            step for step in plan["steps"] if step["kind"] == "worker"
        )
        running = self.run_service.start_plan_step(
            action["_agent_run_id"], worker["id"]
        )
        yield {
            "type": "plan_created",
            "run_id": action["_agent_run_id"],
            "goal": plan["goal"],
            "plan_steps": plan["steps"],
        }
        yield {
            "type": "plan_step_update",
            "run_id": action["_agent_run_id"],
            "plan_step": running,
        }
        while not interrupted.is_set():
            await asyncio.sleep(0.01)
        yield {"type": "interrupted", "model": interrupted_scope[0]}

    def record_interrupt(self, scope="all"):
        interrupted_scope.append(scope)
        interrupted.set()

    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", slow_plan_stream)
    monkeypatch.setattr(SessionOrchestrator, "interrupt", record_interrupt)

    session_id = str(uuid4())
    with client.websocket_connect(f"/ws/sessions/{session_id}") as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "run a persisted plan step",
            "targets": ["model-a"],
            "mode": "project",
            "execution_strategy": "planned",
        })
        assert websocket.receive_json() == {"type": "ack"}
        assert websocket.receive_json()["type"] == "agent_run_update"
        plan_event = websocket.receive_json()
        assert plan_event["type"] == "plan_created"
        update = websocket.receive_json()
        assert update["type"] == "plan_step_update"
        plan_step_id = update["plan_step"]["id"]

        websocket.send_json({
            "type": "interrupt",
            "run_id": update["run_id"],
            "step_id": plan_step_id,
        })
        assert websocket.receive_json() == {
            "type": "interrupt_ack",
            "scope": "model-a",
        }
        assert interrupted.wait(timeout=1)
        assert interrupted_scope == ["model-a"]


def test_websocket_plan_edit_and_confirmation_resume_same_run(client, monkeypatch):
    phases = []

    async def confirmation_stream(self, action):
        run_id = action["_agent_run_id"]
        if not action.get("_plan_confirmed"):
            phases.append("draft")
            plan = self.run_service.ensure_plan(
                run_id=run_id,
                session_id=self.session_id,
                action=action,
                models=["model-a"],
                profiles={"model-a": "Architect/Reviewer"},
            )
            yield {
                "type": "plan_created",
                "run_id": run_id,
                "goal": plan["goal"],
                "plan_steps": plan["steps"],
            }
            waiting = self.run_service.database.update_agent_run(
                run_id, "waiting_plan_confirmation", error=None
            )
            yield {
                "type": "plan_confirmation_required",
                "run_id": run_id,
                "revision": waiting["plan_revision"],
                "agent_run": waiting,
            }
            return

        phases.append("execute")
        for step in self.run_service.database.list_agent_plan_steps(run_id):
            if step["kind"] == "planner":
                continue
            self.run_service.start_plan_step(run_id, step["id"])
            self.run_service.finish_plan_step(
                run_id,
                step["id"],
                status="completed",
                content=f"completed-{step['kind']}",
            )
        message = Message(
            id=uuid4(),
            session_id=self.session_id,
            author_model="model-a",
            role="assistant",
            content="confirmed final",
            created_at=datetime.now(UTC),
        )
        yield {
            "type": "message_complete",
            "message": message.model_dump(mode="json"),
        }

    monkeypatch.setattr(
        SessionOrchestrator,
        "stream_user_message",
        confirmation_stream,
    )
    session_id = str(uuid4())
    with client.websocket_connect(f"/ws/sessions/{session_id}") as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "review this plan first",
            "targets": ["model-a"],
            "mode": "project",
            "execution_strategy": "planned",
            "plan_confirmation_required": True,
        })
        assert websocket.receive_json() == {"type": "ack"}
        initial = websocket.receive_json()
        assert initial["type"] == "agent_run_update"
        run_id = initial["agent_run"]["id"]
        plan_event = websocket.receive_json()
        assert plan_event["type"] == "plan_created"
        worker = next(
            step for step in plan_event["plan_steps"]
            if step["kind"] == "worker"
        )
        required = websocket.receive_json()
        assert required["type"] == "plan_confirmation_required"
        assert required["agent_run"]["status"] == "waiting_plan_confirmation"
        waiting_update = websocket.receive_json()
        assert waiting_update["type"] == "agent_run_update"
        assert waiting_update["agent_run"]["status"] == "waiting_plan_confirmation"
        assert phases == ["draft"]

        websocket.send_json({
            "type": "update_agent_plan",
            "run_id": run_id,
            "expected_revision": 1,
            "goal_spec": {
                "objective": "edited websocket goal",
                "success_criteria": ["must be confirmed"],
                "constraints": ["read only before approval"],
                "deliverable": "confirmed result",
            },
            "plan_steps": [{
                "id": worker["id"],
                "title": "edited worker",
                "description": "wait for confirmation",
                "assigned_model": "model-a",
                "role": "Reviewed Worker",
            }],
            "change_summary": "edited over websocket",
        })
        updated = websocket.receive_json()
        assert updated["type"] == "plan_updated"
        assert updated["revision"] == 2
        assert updated["goal"]["objective"] == "edited websocket goal"

        websocket.send_json({
            "type": "confirm_agent_plan",
            "run_id": run_id,
            "expected_revision": 2,
        })
        assert websocket.receive_json() == {"type": "ack"}
        approved = websocket.receive_json()
        assert approved["type"] == "agent_run_update"
        assert approved["agent_run"]["plan_status"] == "approved"
        final_message = websocket.receive_json()
        assert final_message["type"] == "message_complete"
        assert final_message["message"]["content"] == "confirmed final"
        completed = websocket.receive_json()
        assert completed["type"] == "agent_run_update"
        assert completed["agent_run"]["status"] == "completed_with_warnings"
        assert phases == ["draft", "execute"]


def test_websocket_second_message_becomes_persisted_steer(client, monkeypatch):
    async def steerable_stream(self, action):
        yield {"type": "delta", "model": "slow-model", "delta": "initial plan"}
        for _ in range(200):
            if self._has_pending_steers("slow-model"):
                break
            await asyncio.sleep(0.01)
        steers = self._consume_steers("slow-model")
        assert len(steers) == 1
        for steer in steers:
            self.run_service.mark_steer_applied(
                action["_agent_run_id"],
                steer["id"],
                "slow-model",
            )
            yield {
                "type": "steer_applied",
                "model": "slow-model",
                "steer_id": steer["id"],
            }
        yield {"type": "stream_reset", "model": "slow-model"}
        yield {"type": "delta", "model": "slow-model", "delta": "replanned"}
        message = Message(
            id=uuid4(),
            session_id=self.session_id,
            author_model="slow-model",
            role="assistant",
            content="已按新方向完成",
            created_at=datetime.now(UTC),
        )
        yield {"type": "message_complete", "message": message.model_dump(mode="json")}

    monkeypatch.setattr(SessionOrchestrator, "stream_user_message", steerable_stream)

    session_id = str(uuid4())
    run_id = None
    saw_steer_ack = False
    saw_user_steer = False
    saw_applied = False
    saw_completed = False
    with client.websocket_connect(f"/ws/sessions/{session_id}") as websocket:
        websocket.send_json({
            "type": "post_message",
            "content": "start a steerable run",
            "targets": ["slow-model"],
        })
        assert websocket.receive_json() == {"type": "ack"}
        while True:
            event = websocket.receive_json()
            if event.get("type") == "agent_run_update":
                run_id = event["agent_run"]["id"]
            if event.get("type") == "delta":
                break

        websocket.send_json({
            "type": "post_message",
            "content": "改用方案 B，并停止旧计划",
            "targets": ["slow-model"],
        })
        for _ in range(12):
            event = websocket.receive_json()
            if event.get("type") == "steer_ack":
                saw_steer_ack = event.get("run_id") == run_id
            elif (
                event.get("type") == "message_complete"
                and event.get("message", {}).get("role") == "user"
            ):
                saw_user_steer = event["message"]["content"] == "改用方案 B，并停止旧计划"
            elif event.get("type") == "steer_applied":
                saw_applied = event.get("model") == "slow-model"
            elif (
                event.get("type") == "agent_run_update"
                and event.get("agent_run", {}).get("status") == "completed_with_warnings"
            ):
                saw_completed = True
                break

    assert run_id
    assert saw_steer_ack is True
    assert saw_user_steer is True
    assert saw_applied is True
    assert saw_completed is True
    detail = client.get(f"/agent-runs/{run_id}", params={"session_id": session_id}).json()
    assert detail["steers"][0]["content"] == "改用方案 B，并停止旧计划"
    assert detail["steers"][0]["applied_models"] == ["slow-model"]
    event_types = [event["event_type"] for event in detail["events"]]
    assert "steer_queued" in event_types
    assert "steer_applied" in event_types
    stored = client.get(f"/sessions/{session_id}/messages").json()
    assert any(message["content"] == "改用方案 B，并停止旧计划" for message in stored)


def test_workspace_artifact_merge_tool_and_export_flow(client):
    session_id = str(uuid4())
    other_session_id = str(uuid4())

    seeded = client.get("/artifacts/", params={"session_id": session_id})
    assert seeded.status_code == 200
    seeded_artifacts = {item["path"]: item for item in seeded.json()}
    assert set(seeded_artifacts) >= {"README.md", "docs/PRD.md", "src/demo.ts"}
    assert seeded_artifacts["src/demo.ts"]["mime_type"] == "text/typescript"

    for content in ("version one", "version two"):
        written = client.post("/artifacts/write", json={
            "session_id": session_id,
            "path": "notes/flow.txt",
            "content": content,
            "commit_message": content,
            "author": "test",
        })
        assert written.status_code == 200

    history = client.get("/artifacts/notes/flow.txt/history", params={"session_id": session_id})
    assert history.status_code == 200
    assert [item["content"] for item in history.json()[:2]] == ["version two", "version one"]
    invalid_path = client.post("/artifacts/write", json={
        "session_id": session_id,
        "path": "../outside.txt",
        "content": "blocked",
    })
    assert invalid_path.status_code == 400

    proposal = client.post("/artifacts/merge/proposals", json={
        "session_id": session_id,
        "artifact_path": "notes/flow.txt",
        "proposer": "user",
        "rationale": "offline flow",
        "merge_rule": "两票通过",
    })
    assert proposal.status_code == 200
    proposal_id = proposal.json()["id"]
    assert proposal.json()["merge_rule"] == "two_votes"
    first_vote = client.post(f"/artifacts/merge/proposals/{proposal_id}/vote", json={
        "voter": "model:a",
        "approve": True,
        "session_id": session_id,
    })
    assert first_vote.json()["status"] == "open"
    second_vote = client.post(f"/artifacts/merge/proposals/{proposal_id}/vote", json={
        "voter": "model:b",
        "approve": True,
        "session_id": session_id,
    })
    assert second_vote.json()["status"] == "merged"
    late_vote = client.post(f"/artifacts/merge/proposals/{proposal_id}/vote", json={
        "voter": "model:c",
        "approve": False,
        "session_id": session_id,
    })
    assert late_vote.json()["status"] == "merged"
    assert late_vote.json()["vote_accepted"] is False

    tools = client.get("/tools/", params={"session_id": session_id}).json()
    by_name = {tool["name"]: tool for tool in tools}
    assert {"list_tasks", "create_task", "read_artifact", "write_artifact"} <= set(by_name)

    task_run = client.post(f"/tools/{by_name['create_task']['id']}/invoke", json={
        "caller": "user",
        "session_id": session_id,
        "arguments": {"title": "created through dispatcher", "priority": 1},
    })
    assert task_run.status_code == 200
    assert task_run.json()["status"] == "success", task_run.json().get("error")
    assert isinstance(task_run.json()["output"], str)
    assert task_run.json()["result"]["task"]["title"] == "created through dispatcher"

    tool_write = client.post(f"/tools/{by_name['write_artifact']['id']}/invoke", json={
        "caller": "user",
        "session_id": session_id,
        "arguments": {"path": "notes/tool.txt", "content": "written by dispatcher"},
    })
    assert tool_write.status_code == 200
    assert tool_write.json()["status"] == "success"
    read_back = client.get("/artifacts/notes/tool.txt", params={"session_id": session_id})
    assert read_back.json()["content"] == "written by dispatcher"

    session_runs = client.get("/tools/runs", params={"session_id": session_id}).json()
    other_runs = client.get("/tools/runs", params={"session_id": other_session_id}).json()
    assert len(session_runs) >= 2
    assert other_runs == []

    exported = client.get(f"/export/session/{session_id}", params={"format": "json"})
    assert exported.status_code == 200
    snapshot = exported.json()
    assert any(task["title"] == "created through dispatcher" for task in snapshot["tasks"])
    assert any(item["path"] == "notes/tool.txt" and item["content"] == "written by dispatcher" for item in snapshot["artifacts"])
    assert snapshot["merge_proposals"][0]["status"] == "merged"
    assert len(snapshot["tool_runs"]) >= 2

    markdown = client.get(f"/export/session/{session_id}", params={"format": "markdown"})
    assert markdown.status_code == 200
    assert "## Artifacts (content and complete history)" in markdown.text
    assert "written by dispatcher" in markdown.text


def test_interactive_artifact_preview_versions_office_exports_and_rejections(client):
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation

    session_id = str(uuid4())
    other_session_id = str(uuid4())
    markdown = "# Release brief\n\n| Owner | Status |\n| --- | --- |\n| Team | Ready |\n\n---\n# Next slide\n\n- Safe export\n:::notes\nSpeaker-only context"
    first = client.post("/artifacts/write", json={
        "session_id": session_id,
        "path": "reports/brief.md",
        "content": "# Draft\n",
        "commit_message": "draft",
        "author": "test",
        "source_model": "offline-model",
        "status": "draft",
    })
    assert first.status_code == 200
    second = client.post("/artifacts/write", json={
        "session_id": session_id,
        "path": "reports/brief.md",
        "content": markdown,
        "commit_message": "ready",
        "author": "test",
        "source_model": "offline-model",
        "status": "ready",
        "parent_version": 1,
    })
    assert second.status_code == 200

    listed = client.get("/artifacts/", params={"session_id": session_id}).json()
    artifact = next(item for item in listed if item["path"] == "reports/brief.md")
    assert artifact["mime_type"] == "text/markdown"
    assert artifact["current_version"] == 2
    assert artifact["source_model"] == "offline-model"
    assert artifact["parent_version"] == 1

    preview = client.get("/artifacts/reports/brief.md/preview", params={"session_id": session_id})
    assert preview.status_code == 200
    assert preview.json()["mode"] == "markdown"
    assert preview.json()["content"] == markdown
    assert client.get(
        "/artifacts/reports/brief.md/preview", params={"session_id": other_session_id}
    ).status_code == 404

    hostile = '<script>fetch("https://attacker.invalid")</script><img src="file:///etc/passwd"><h1>isolated</h1>'
    assert client.post("/artifacts/write", json={
        "session_id": session_id, "path": "preview/hostile.html", "content": hostile,
    }).status_code == 200
    hostile_preview = client.get(
        "/artifacts/preview/hostile.html/preview", params={"session_id": session_id}
    ).json()
    assert hostile_preview["mode"] == "isolated_html"
    assert hostile_preview["sandbox"] == "allow-none"
    assert "attacker.invalid" not in hostile_preview["content"]
    assert "file:///etc/passwd" not in hostile_preview["content"]

    diff = client.get("/artifacts/reports/brief.md/diff", params={
        "session_id": session_id, "from_version": 1, "to_version": 2,
    })
    assert diff.status_code == 200
    assert "-# Draft" in diff.json()["diff"]
    assert "+# Release brief" in diff.json()["diff"]
    restored = client.post("/artifacts/reports/brief.md/restore", json={
        "session_id": session_id, "version": 1, "author": "test",
    })
    assert restored.status_code == 200
    assert restored.json()["version"] == 3
    assert restored.json()["parent_version"] == 1
    assert client.get("/artifacts/reports/brief.md", params={"session_id": session_id}).json()["content"] == "# Draft\n"
    assert client.post("/artifacts/reports/brief.md/restore", json={
        "session_id": other_session_id, "version": 1,
    }).status_code == 400

    # Restore the rich version before verifying all three Office writers.
    assert client.post("/artifacts/reports/brief.md/restore", json={
        "session_id": session_id, "version": 2, "author": "test",
    }).status_code == 200
    docx_response = client.get("/artifacts/reports/brief.md/export", params={
        "session_id": session_id, "format": "docx",
    })
    assert docx_response.status_code == 200
    document = Document(io.BytesIO(docx_response.content))
    assert any(paragraph.text == "Release brief" for paragraph in document.paragraphs)
    assert document.tables[0].cell(1, 0).text == "Team"

    pptx_response = client.get("/artifacts/reports/brief.md/export", params={
        "session_id": session_id, "format": "pptx",
    })
    assert pptx_response.status_code == 200
    presentation = Presentation(io.BytesIO(pptx_response.content))
    assert len(presentation.slides) == 2
    assert presentation.slides[0].shapes.title.text == "Release brief"
    assert "Speaker-only context" in presentation.slides[1].notes_slide.notes_text_frame.text

    workbook_payload = json.dumps({"sheets": [{
        "name": "安全报表", "rows": [["name", "value"], ["formula", "=HYPERLINK(\"https://attacker.invalid\")"]],
    }]}, ensure_ascii=False)
    assert client.post("/artifacts/write", json={
        "session_id": session_id, "path": "reports/data.json", "content": workbook_payload,
    }).status_code == 200
    xlsx_response = client.get("/artifacts/reports/data.json/export", params={
        "session_id": session_id, "format": "xlsx",
    })
    assert xlsx_response.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_response.content), data_only=False)
    worksheet = workbook["安全报表"]
    assert worksheet.freeze_panes == "A2"
    assert worksheet["B2"].value.startswith("'=")
    workbook.close()

    assert client.post("/artifacts/write", json={
        "session_id": session_id, "path": "../escape.svg", "content": "<svg/>",
    }).status_code == 400
    too_large = client.post("/artifacts/write", json={
        "session_id": session_id, "path": "too-large.txt", "content": "x" * (5 * 1024 * 1024 + 1),
    })
    assert too_large.status_code == 400
    assert client.get("/artifacts/reports/brief.md/export", params={
        "session_id": session_id, "format": "exe",
    }).status_code == 400



def test_automation_api_crud_manual_run_and_review(client, monkeypatch):
    session_id = str(uuid4())

    async def fake_background_executor(job):
        assert job["permission_snapshot"]["mode"] == "read_only"
        assert job["permission_snapshot"]["allowed_tools"] == []
        return AutomationExecutionResult(
            status="completed",
            summary="后台摘要已生成",
            budget_spent_usd=0.0123,
        )

    monkeypatch.setattr(automation_service, "_executor", fake_background_executor)
    first_run = datetime.now(UTC) + timedelta(hours=2)
    payload = {
        "session_id": session_id,
        "name": "每小时只读摘要",
        "prompt": "只读取当前会话并生成摘要，不修改任何文件。",
        "mode": "project",
        "targets": ["all"],
        "trigger_type": "interval",
        "interval_minutes": 60,
        "timezone": "Asia/Chongqing",
        "next_run_at": first_run.isoformat(),
        "enabled": True,
        "max_budget_usd": 0.5,
        "permission_mode": "read_only",
    }
    created = client.post("/automations/schedules", json=payload)
    assert created.status_code == 201
    schedule = created.json()
    assert schedule["permission_snapshot"]["allowed_tools"] == []
    assert schedule["next_run_at"] == first_run.isoformat()

    rejected = client.post(
        "/automations/schedules",
        json={**payload, "name": "不安全任务", "permission_mode": "full_access"},
    )
    assert rejected.status_code == 422

    listed = client.get(
        "/automations/schedules",
        params={"session_id": session_id},
    )
    assert [item["id"] for item in listed.json()] == [schedule["id"]]

    paused = client.patch(
        f"/automations/schedules/{schedule['id']}",
        json={"session_id": session_id, "enabled": False},
    )
    assert paused.status_code == 200
    assert paused.json()["enabled"] is False

    queued = client.post(
        f"/automations/schedules/{schedule['id']}/run",
        json={"session_id": session_id},
    )
    assert queued.status_code == 202
    job_id = queued.json()["id"]

    completed_job = None
    for _ in range(100):
        completed_job = automation_service.database.get_background_job(
            job_id,
            session_id=session_id,
        )
        if completed_job["status"] == "completed":
            break
        time.sleep(0.02)
    assert completed_job is not None
    assert completed_job["status"] == "completed"
    assert completed_job["summary"] == "后台摘要已生成"
    assert completed_job["budget_spent_usd"] == pytest.approx(0.0123)
    assert completed_job["reviewed_at"] is None
    listed_jobs = client.get(
        "/automations/jobs",
        params={"session_id": session_id},
    )
    assert listed_jobs.status_code == 200
    assert listed_jobs.json()[0]["id"] == job_id

    reviewed = client.post(
        f"/automations/jobs/{job_id}/review",
        json={"session_id": session_id},
    )
    assert reviewed.status_code == 200
    assert reviewed.json()["reviewed_at"]
    assert client.delete(
        f"/automations/schedules/{schedule['id']}",
        params={"session_id": session_id},
    ).json()["deleted"] is True


@pytest.mark.asyncio
async def test_automation_scheduler_materializes_once_and_reclaims_expired_lease(tmp_path):
    database = ChatDatabase(tmp_path / "automations.sqlite3")
    dispatcher = RunDispatcher(capacity=1)
    session_id = str(uuid4())
    now = datetime.now(UTC).replace(microsecond=0)
    calls = []

    async def fake_executor(job):
        calls.append(job["id"])
        await asyncio.sleep(0)
        return AutomationExecutionResult(status="completed", summary="ok")

    schedule = database.create_automation_schedule(
        session_id=session_id,
        name="可恢复循环任务",
        prompt="生成只读摘要",
        mode="project",
        targets=["all"],
        trigger_type="interval",
        interval_minutes=30,
        timezone="UTC",
        next_run_at=(now - timedelta(minutes=95)).isoformat(),
        enabled=True,
        max_budget_usd=1.0,
        permission_snapshot={"policy_version": 1, "mode": "read_only", "allowed_tools": []},
    )
    service = AutomationService(
        database,
        AgentRunService(database),
        poll_interval=0.01,
        lease_seconds=6,
        concurrency=1,
        executor=fake_executor,
        dispatcher=dispatcher,
        resource_resolver=lambda _action: (),
    )
    materialized = await service.run_cycle(now)
    assert len(materialized) == 1
    await asyncio.gather(*list(service._active_tasks.values()))
    await service.run_cycle(now)
    first_jobs = database.list_background_jobs(session_id)
    assert len(first_jobs) == 1
    assert first_jobs[0]["status"] == "completed"
    assert first_jobs[0]["permission_snapshot"]["mode"] == "read_only"
    advanced = database.get_automation_schedule(schedule["id"], session_id=session_id)
    assert advanced is not None
    assert datetime.fromisoformat(advanced["next_run_at"]) > now

    stranded = database.create_background_job(
        schedule_id=None,
        session_id=session_id,
        trigger_type="manual",
        action={"content": "恢复我", "targets": ["all"], "mode": "project"},
        permission_snapshot={"mode": "read_only", "allowed_tools": []},
        max_budget_usd=1.0,
        scheduled_for=now.isoformat(),
    )
    claimed = database.claim_background_job(
        worker_id="dead-worker",
        now=now.isoformat(),
        lease_expires_at=(now - timedelta(seconds=1)).isoformat(),
    )
    assert claimed["id"] == stranded["id"]
    assert claimed["attempt"] == 1

    await service.run_cycle(now)
    await asyncio.gather(*list(service._active_tasks.values()))
    recovered = database.get_background_job(stranded["id"])
    assert recovered is not None
    assert recovered["status"] == "completed"
    assert recovered["attempt"] == 2
    assert calls.count(stranded["id"]) == 1
    await dispatcher.stop()


@pytest.mark.asyncio
async def test_background_run_rejects_unapproved_side_effect_before_execution(tmp_path):
    database = ChatDatabase(tmp_path / "background-read-only.sqlite3")
    run_service = AgentRunService(database)
    session_id = str(uuid4())
    action = {
        "content": "分析项目，并尝试写入一个文件",
        "targets": ["solo"],
        "mode": "casual",
        "_background_job_id": str(uuid4()),
        "_background_allowed_tools": [],
        "_background_read_only": True,
    }
    run = run_service.create(session_id, action)
    action["_agent_run_id"] = run["id"]

    class BackgroundProvider(_FakeProvider):
        def __init__(self):
            super().__init__("solo")
            self.calls = 0

        async def stream_chat(self, messages, tools=None, **kwargs):
            self.calls += 1
            assert "write_artifact" not in {tool.name for tool in (tools or [])}
            assert any(
                "无人值守执行" in message.content
                for message in messages
                if message.role == "system"
            )
            if self.calls == 1:
                yield {
                    "delta": "",
                    "usage": None,
                    "tool_calls": [{
                        "id": "blocked-background-write",
                        "name": "write_artifact",
                        "arguments": {
                            "path": "must-not-exist.md",
                            "content": "forbidden",
                        },
                    }],
                }
                return
            observation = next(
                message for message in reversed(messages) if message.role == "tool"
            )
            assert observation.tool_error is True
            assert "后台权限快照不允许" in observation.content
            yield {"delta": "已完成只读分析，未修改文件。", "usage": None}

    orchestrator = SessionOrchestrator(session_id, mode="casual")
    orchestrator.providers = {"solo": BackgroundProvider()}
    orchestrator.refresh_providers = lambda: None
    orchestrator.run_service = run_service
    orchestrator.begin_agent_run(run["id"])
    tool_runs_before = len(tool_registry.list_runs(session_id=session_id, limit=None))
    async for event in orchestrator.stream_user_message(action):
        run_service.observe(run["id"], event)
    completed = run_service.finish(run["id"])

    assert completed["status"] == "completed_with_warnings"
    assert len(tool_registry.list_runs(session_id=session_id, limit=None)) == tool_runs_before
    calls = database.list_agent_tool_calls(run["id"])
    assert len(calls) == 1
    assert calls[0]["status"] == "denied"
    assert calls[0]["result"]["status"] == "denied"
