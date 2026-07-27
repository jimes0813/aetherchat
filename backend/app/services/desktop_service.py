from __future__ import annotations

import asyncio
import difflib
import io
import json
import os
import secrets
import shutil
import textwrap
from collections.abc import Mapping
from datetime import UTC, date, datetime
from datetime import time as datetime_time
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".csv",
    ".tsv",
    ".docx",
    ".xlsx",
    ".pptx",
    ".pdf",
}
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_TEXT_BYTES = 5 * 1024 * 1024
MAX_DOCUMENT_PAYLOAD_BYTES = 10 * 1024 * 1024
MAX_PDF_NOTE_CHARACTERS = 200_000
MAX_SHEET_ROWS = 500
MAX_SHEET_COLUMNS = 100


class DesktopCapabilityError(ValueError):
    pass


def desktop_token_is_valid(candidate: str | None) -> bool:
    expected = os.environ.get("AETHERCHAT_DESKTOP_TOKEN", "")
    return bool(expected and candidate and secrets.compare_digest(expected, candidate))


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    return str(value)


def _json_value_type(value: Any) -> str | None:
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, datetime_time):
        return "time"
    return None


def _restore_typed_value(value: Any, value_type: Any) -> Any:
    if not isinstance(value, str):
        return value
    try:
        if value_type == "datetime":
            return datetime.fromisoformat(value)
        if value_type == "date":
            return date.fromisoformat(value)
        if value_type == "time":
            return datetime_time.fromisoformat(value)
    except ValueError as exc:
        raise DesktopCapabilityError(f"Invalid ISO {value_type} value: {value}") from exc
    return value


def _set_paragraph_text_preserving_runs(paragraph: Any, new_text: str) -> None:
    old_text = paragraph.text
    if old_text == new_text:
        return
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(new_text)
        return
    if not old_text:
        runs[0].text = new_text
        return

    old_character_runs: list[int] = []
    for run_index, run in enumerate(runs):
        old_character_runs.extend([run_index] * len(run.text))
    if not old_character_runs:
        runs[0].text = new_text
        return

    assigned: list[int | None] = [None] * len(new_text)
    matcher = difflib.SequenceMatcher(a=old_text, b=new_text, autojunk=False)
    for tag, old_start, _old_end, new_start, new_end in matcher.get_opcodes():
        if tag == "equal":
            for offset in range(new_end - new_start):
                assigned[new_start + offset] = old_character_runs[old_start + offset]
            continue
        if new_start == new_end:
            continue
        style_index = (
            old_character_runs[old_start]
            if old_start < len(old_character_runs)
            else old_character_runs[-1]
        )
        for position in range(new_start, new_end):
            assigned[position] = style_index

    buckets: list[list[str]] = [[] for _ in runs]
    for position, character in enumerate(new_text):
        assigned_index = assigned[position]
        run_index = assigned_index if assigned_index is not None else 0
        buckets[run_index].append(character)
    for run, characters in zip(runs, buckets, strict=True):
        run.text = "".join(characters)


def _decode_text_document(raw: bytes) -> tuple[str, str]:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig"), "utf-8-sig"
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "utf-16"
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        try:
            return raw.decode("gb18030"), "gb18030"
        except UnicodeDecodeError as exc:
            raise DesktopCapabilityError("Text encoding is not UTF-8, UTF-16, or GB18030") from exc


def _validated_path(path_value: str, *, must_exist: bool) -> Path:
    if not isinstance(path_value, str) or not path_value.strip() or "\x00" in path_value:
        raise DesktopCapabilityError("Invalid document path")
    raw = path_value.strip()
    if raw.startswith(("\\\\", "//")):
        raise DesktopCapabilityError("Network and device paths are not supported")
    _drive, tail = os.path.splitdrive(raw)
    if ":" in tail:
        raise DesktopCapabilityError("NTFS alternate data streams are not supported")
    path = Path(raw).expanduser()
    try:
        resolved = path.resolve(strict=must_exist)
    except OSError as exc:
        raise DesktopCapabilityError(f"Unable to resolve document path: {exc}") from exc
    if resolved.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise DesktopCapabilityError(f"Unsupported document type: {resolved.suffix or '(none)'}")
    if must_exist:
        if not resolved.is_file():
            raise DesktopCapabilityError("Selected document does not exist or is not a file")
        if resolved.stat().st_size > MAX_FILE_BYTES:
            raise DesktopCapabilityError("Selected document exceeds the 50 MB safety limit")
    else:
        resolved.parent.mkdir(parents=True, exist_ok=True)
    return resolved


class DesktopDocumentService:
    async def read(self, path_value: str) -> dict[str, Any]:
        path = _validated_path(path_value, must_exist=True)
        return await asyncio.to_thread(self._read_sync, path)

    async def write(
        self,
        path_value: str,
        data: dict[str, Any],
        *,
        source_path_value: str | None = None,
        create_backup: bool = True,
    ) -> dict[str, Any]:
        try:
            payload_size = len(json.dumps(data, ensure_ascii=False, default=str).encode("utf-8"))
        except (TypeError, ValueError) as exc:
            raise DesktopCapabilityError("Document data must be JSON-compatible") from exc
        if payload_size > MAX_DOCUMENT_PAYLOAD_BYTES:
            raise DesktopCapabilityError("Document edit payload exceeds the 10 MB safety limit")
        destination = _validated_path(path_value, must_exist=False)
        source: Path | None = None
        if source_path_value:
            source = _validated_path(source_path_value, must_exist=True)
        elif destination.exists():
            source = _validated_path(path_value, must_exist=True)
        elif destination.suffix.lower() not in {".txt", ".md", ".json", ".csv", ".tsv"}:
            raise DesktopCapabilityError(
                "Creating a new structured document requires an existing source_path"
            )
        if source is not None and destination.suffix.lower() != source.suffix.lower():
            raise DesktopCapabilityError("Save As must keep the original document extension")
        return await asyncio.to_thread(
            self._write_sync,
            source,
            destination,
            data,
            create_backup,
        )

    def _read_sync(self, path: Path) -> dict[str, Any]:
        extension = path.suffix.lower()
        base = {
            "name": path.name,
            "format": extension.removeprefix("."),
            "size_bytes": path.stat().st_size,
            "modified_at": datetime.fromtimestamp(path.stat().st_mtime, UTC).isoformat(),
        }
        if extension in {".txt", ".md", ".json", ".csv", ".tsv"}:
            if path.stat().st_size > MAX_TEXT_BYTES:
                raise DesktopCapabilityError("Text document exceeds the 5 MB editing limit")
            text, encoding = _decode_text_document(path.read_bytes())
            return {**base, "editor_mode": "text", "data": {"text": text, "encoding": encoding}}
        if extension == ".docx":
            from docx import Document

            document = Document(str(path))
            tables = [
                [[cell.text for cell in row.cells] for row in table.rows]
                for table in document.tables
            ]
            return {
                **base,
                "editor_mode": "structured",
                "data": {
                    "paragraphs": [paragraph.text for paragraph in document.paragraphs],
                    "tables": tables,
                },
            }
        if extension == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, data_only=False, read_only=False)
            sheets = []
            truncated = False
            for worksheet in workbook.worksheets:
                max_row = min(worksheet.max_row, MAX_SHEET_ROWS)
                max_column = min(worksheet.max_column, MAX_SHEET_COLUMNS)
                truncated = truncated or worksheet.max_row > max_row or worksheet.max_column > max_column
                values = [
                    [_json_value(worksheet.cell(row=row, column=column).value) for column in range(1, max_column + 1)]
                    for row in range(1, max_row + 1)
                ]
                value_types = [
                    [_json_value_type(worksheet.cell(row=row, column=column).value) for column in range(1, max_column + 1)]
                    for row in range(1, max_row + 1)
                ]
                sheets.append({"name": worksheet.title, "values": values, "value_types": value_types})
            return {**base, "editor_mode": "structured", "data": {"sheets": sheets}, "truncated": truncated}
        if extension == ".pptx":
            from pptx import Presentation

            presentation = Presentation(str(path))
            slides = []
            for slide_index, slide in enumerate(presentation.slides):
                texts = []
                for shape_index, shape in enumerate(slide.shapes):
                    if getattr(shape, "has_text_frame", False):
                        texts.append({
                            "shape_index": shape_index,
                            "shape_name": shape.name,
                            "text": shape.text,
                        })
                slides.append({"slide_index": slide_index, "texts": texts})
            return {**base, "editor_mode": "structured", "data": {"slides": slides}}
        if extension == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(path)
            extracted = []
            for page in reader.pages[:100]:
                extracted.append(page.extract_text() or "")
            metadata: Mapping[str, Any] = reader.metadata or {}
            return {
                **base,
                "editor_mode": "pdf_note",
                "read_only_text": "\n\n".join(extracted)[:1_000_000],
                "data": {
                    "metadata": {
                        "title": metadata.get("/Title", ""),
                        "author": metadata.get("/Author", ""),
                        "subject": metadata.get("/Subject", ""),
                    },
                    "append_note": "",
                },
                "page_count": len(reader.pages),
            }
        raise DesktopCapabilityError("Unsupported document type")

    def _write_sync(
        self,
        source: Path | None,
        destination: Path,
        data: dict[str, Any],
        create_backup: bool,
    ) -> dict[str, Any]:
        if not isinstance(data, dict):
            raise DesktopCapabilityError("Document data must be an object")
        backup_path: Path | None = None
        if create_backup and destination.exists():
            timestamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
            backup_path = destination.with_name(
                f"{destination.stem}.aetherchat-backup-{timestamp}{destination.suffix}"
            )
            shutil.copy2(destination, backup_path)

        temporary = destination.with_name(
            f".{destination.stem}.aetherchat-{uuid4().hex}.tmp{destination.suffix}"
        )
        try:
            extension = destination.suffix.lower()
            if extension in {".txt", ".md", ".json", ".csv", ".tsv"}:
                text = str(data.get("text", ""))
                if len(text.encode("utf-8")) > MAX_TEXT_BYTES:
                    raise DesktopCapabilityError("Text output exceeds the 5 MB editing limit")
                encoding = str(data.get("encoding") or "utf-8")
                if encoding not in {"utf-8", "utf-8-sig", "utf-16", "gb18030"}:
                    raise DesktopCapabilityError("Unsupported text output encoding")
                temporary.write_text(text, encoding=encoding)
            elif extension == ".docx":
                if source is None:
                    raise DesktopCapabilityError("DOCX output requires an existing source document")
                self._write_docx(source, temporary, data)
            elif extension == ".xlsx":
                if source is None:
                    raise DesktopCapabilityError("XLSX output requires an existing source workbook")
                self._write_xlsx(source, temporary, data)
            elif extension == ".pptx":
                if source is None:
                    raise DesktopCapabilityError("PPTX output requires an existing source presentation")
                self._write_pptx(source, temporary, data)
            elif extension == ".pdf":
                if source is None:
                    raise DesktopCapabilityError("PDF output requires an existing source document")
                self._write_pdf(source, temporary, data)
            else:
                raise DesktopCapabilityError("Unsupported document type")
            with temporary.open("rb+") as handle:
                handle.flush()
                os.fsync(handle.fileno())
            temporary.replace(destination)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise

        return {
            "saved": True,
            "name": destination.name,
            "format": destination.suffix.lower().removeprefix("."),
            "size_bytes": destination.stat().st_size,
            "backup_path": str(backup_path) if backup_path else None,
            "saved_at": datetime.now(UTC).isoformat(),
        }

    @staticmethod
    def _write_docx(source: Path, destination: Path, data: dict[str, Any]) -> None:
        from docx import Document

        document = Document(str(source))
        paragraphs = data.get("paragraphs") or []
        if not isinstance(paragraphs, list):
            raise DesktopCapabilityError("DOCX paragraphs must be an array")
        for index, text in enumerate(paragraphs):
            if index < len(document.paragraphs):
                _set_paragraph_text_preserving_runs(document.paragraphs[index], str(text))
            else:
                document.add_paragraph(str(text))

        tables = data.get("tables") or []
        if not isinstance(tables, list):
            raise DesktopCapabilityError("DOCX tables must be an array")
        for table_index, table_data in enumerate(tables[: len(document.tables)]):
            if not isinstance(table_data, list):
                continue
            table = document.tables[table_index]
            for row_index, row_data in enumerate(table_data[: len(table.rows)]):
                if not isinstance(row_data, list):
                    continue
                for column_index, value in enumerate(row_data[: len(table.rows[row_index].cells)]):
                    cell = table.cell(row_index, column_index)
                    new_text = str(value if value is not None else "")
                    if cell.text == new_text:
                        continue
                    parts = new_text.split("\n")
                    if len(parts) != len(cell.paragraphs):
                        raise DesktopCapabilityError(
                            "Changing the paragraph count inside a DOCX table cell is not supported"
                        )
                    for paragraph, part in zip(cell.paragraphs, parts, strict=True):
                        _set_paragraph_text_preserving_runs(paragraph, part)
        document.save(str(destination))

    @staticmethod
    def _write_xlsx(source: Path, destination: Path, data: dict[str, Any]) -> None:
        from openpyxl import load_workbook
        from openpyxl.cell.cell import MergedCell

        workbook = load_workbook(source, data_only=False, read_only=False)
        sheets = data.get("sheets") or []
        if not isinstance(sheets, list):
            raise DesktopCapabilityError("XLSX sheets must be an array")
        for sheet_data in sheets:
            if not isinstance(sheet_data, dict):
                continue
            name = str(sheet_data.get("name") or "Sheet")[:31]
            worksheet = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
            values = sheet_data.get("values") or []
            value_types = sheet_data.get("value_types") or []
            if not isinstance(values, list) or len(values) > MAX_SHEET_ROWS:
                raise DesktopCapabilityError("XLSX data exceeds the 500-row editing limit")
            for row_index, row in enumerate(values, start=1):
                if not isinstance(row, list) or len(row) > MAX_SHEET_COLUMNS:
                    raise DesktopCapabilityError("XLSX data exceeds the 100-column editing limit")
                for column_index, value in enumerate(row, start=1):
                    cell = worksheet.cell(row=row_index, column=column_index)
                    if isinstance(cell, MergedCell):
                        continue
                    value_type = None
                    if row_index - 1 < len(value_types) and isinstance(value_types[row_index - 1], list):
                        type_row = value_types[row_index - 1]
                        if column_index - 1 < len(type_row):
                            value_type = type_row[column_index - 1]
                    if _json_value(cell.value) == value and (
                        value_type is None or _json_value_type(cell.value) == value_type
                    ):
                        continue
                    restored = _restore_typed_value(value, value_type)
                    if cell.value == restored:
                        continue
                    cell.value = restored
        workbook.save(destination)

    @staticmethod
    def _write_pptx(source: Path, destination: Path, data: dict[str, Any]) -> None:
        from pptx import Presentation

        presentation = Presentation(str(source))
        slides = data.get("slides") or []
        if not isinstance(slides, list):
            raise DesktopCapabilityError("PPTX slides must be an array")
        for slide_data in slides:
            if not isinstance(slide_data, dict):
                continue
            slide_index = int(slide_data.get("slide_index", -1))
            if slide_index < 0 or slide_index >= len(presentation.slides):
                continue
            slide = presentation.slides[slide_index]
            for text_data in slide_data.get("texts") or []:
                if not isinstance(text_data, dict):
                    continue
                shape_index = int(text_data.get("shape_index", -1))
                if 0 <= shape_index < len(slide.shapes):
                    shape = slide.shapes[shape_index]
                    if getattr(shape, "has_text_frame", False):
                        new_text = str(text_data.get("text", ""))
                        if shape.text == new_text:
                            continue
                        new_paragraphs = new_text.split("\n")
                        paragraphs = list(shape.text_frame.paragraphs)
                        if len(new_paragraphs) != len(paragraphs):
                            raise DesktopCapabilityError(
                                "Changing the paragraph count inside a PPTX text shape is not supported"
                            )
                        for paragraph, paragraph_text in zip(paragraphs, new_paragraphs, strict=True):
                            _set_paragraph_text_preserving_runs(paragraph, paragraph_text)
        presentation.save(str(destination))

    @staticmethod
    def _write_pdf(source: Path, destination: Path, data: dict[str, Any]) -> None:
        from pypdf import PdfReader, PdfWriter

        reader = PdfReader(source)
        writer = PdfWriter()
        writer.clone_document_from_reader(reader)
        metadata = data.get("metadata") or {}
        if isinstance(metadata, dict):
            existing_metadata = {
                str(key): str(value)
                for key, value in (reader.metadata or {}).items()
                if value is not None
            }
            existing_metadata.update({
                "/Title": str(metadata.get("title") or ""),
                "/Author": str(metadata.get("author") or ""),
                "/Subject": str(metadata.get("subject") or ""),
            })
            writer.add_metadata(existing_metadata)
        note = str(data.get("append_note") or "").strip()
        if note:
            if len(note) > MAX_PDF_NOTE_CHARACTERS:
                raise DesktopCapabilityError("PDF note exceeds the 200,000-character safety limit")
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.pdfgen import canvas

            packet = io.BytesIO()
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            note_canvas = canvas.Canvas(packet, pagesize=letter)
            page_width, page_height = letter

            def start_page():
                note_canvas.setFont("STSong-Light", 12)
                text_object = note_canvas.beginText(54, page_height - 54)
                text_object.textLine("AetherChat 附加备注")
                text_object.textLine("")
                return text_object

            text = start_page()
            for paragraph in note.splitlines() or [note]:
                for line in textwrap.wrap(paragraph, width=72, replace_whitespace=False) or [""]:
                    if text.getY() < 54:
                        note_canvas.drawText(text)
                        note_canvas.showPage()
                        text = start_page()
                    text.textLine(line)
            note_canvas.drawText(text)
            note_canvas.save()
            packet.seek(0)
            note_reader = PdfReader(packet)
            for note_page in note_reader.pages:
                writer.add_page(note_page)
        with destination.open("wb") as handle:
            writer.write(handle)


class DesktopMouseService:
    async def execute(
        self,
        action: Literal["get", "move", "click"],
        *,
        x: int | None = None,
        y: int | None = None,
        button: Literal["left", "right", "middle"] = "left",
    ) -> dict[str, Any]:
        return await asyncio.to_thread(self._execute_sync, action, x, y, button)

    @staticmethod
    def _execute_sync(action: str, x: int | None, y: int | None, button: str) -> dict[str, Any]:
        if os.name != "nt":
            raise DesktopCapabilityError("Mouse control is unavailable on Linux; document and screenshot tools remain supported")
        import win32api
        import win32con

        left = win32api.GetSystemMetrics(win32con.SM_XVIRTUALSCREEN)
        top = win32api.GetSystemMetrics(win32con.SM_YVIRTUALSCREEN)
        width = win32api.GetSystemMetrics(win32con.SM_CXVIRTUALSCREEN)
        height = win32api.GetSystemMetrics(win32con.SM_CYVIRTUALSCREEN)
        bounds = {"left": left, "top": top, "width": width, "height": height}
        if action in {"move", "click"}:
            if x is None or y is None:
                raise DesktopCapabilityError("Mouse x and y coordinates are required")
            if not (left <= x < left + width and top <= y < top + height):
                raise DesktopCapabilityError("Mouse coordinates are outside the virtual desktop")
            win32api.SetCursorPos((x, y))
        if action == "click":
            flags = {
                "left": (win32con.MOUSEEVENTF_LEFTDOWN, win32con.MOUSEEVENTF_LEFTUP),
                "right": (win32con.MOUSEEVENTF_RIGHTDOWN, win32con.MOUSEEVENTF_RIGHTUP),
                "middle": (win32con.MOUSEEVENTF_MIDDLEDOWN, win32con.MOUSEEVENTF_MIDDLEUP),
            }[button]
            win32api.mouse_event(flags[0], 0, 0, 0, 0)
            win32api.mouse_event(flags[1], 0, 0, 0, 0)
        current_x, current_y = win32api.GetCursorPos()
        return {
            "action": action,
            "position": {"x": current_x, "y": current_y},
            "button": button if action == "click" else None,
            "virtual_desktop": bounds,
        }


document_service = DesktopDocumentService()
mouse_service = DesktopMouseService()
