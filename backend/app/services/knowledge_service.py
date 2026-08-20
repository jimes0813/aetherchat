from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import math
import os
import re
import shutil
import sqlite3
import zipfile
from array import array
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from threading import RLock
from typing import Any
from uuid import UUID, uuid4

from docx import Document
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pypdf import PdfReader

from ..data_dir import ensure_private_data_dir, harden_private_file
from .artifact_service import ArtifactService

MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_EXTRACTED_CHARS = 500_000
VECTOR_DIMENSIONS = 256
TEXT_EXTENSIONS = {
    ".txt", ".md", ".json", ".py", ".js", ".ts", ".tsx", ".jsx", ".css", ".html",
    ".yaml", ".yml", ".toml", ".xml", ".sql", ".sh", ".go", ".rs", ".java", ".c",
    ".cpp", ".h",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".csv", ".pdf", ".docx", ".xlsx", ".pptx"}
OFFICE_EXTENSIONS = {".docx", ".xlsx", ".pptx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED_EXTENSIONS |= IMAGE_EXTENSIONS


def _utc_now() -> str:
    return datetime.now(UTC).isoformat()


def _data_root(base_path: str | os.PathLike[str] | None = None) -> Path:
    configured = base_path or os.environ.get("AETHERCHAT_DATA_DIR")
    if configured:
        return Path(configured).expanduser().resolve()
    if os.name == "nt":
        return Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "AetherChat"
    return Path(os.environ.get("XDG_DATA_HOME", Path.home() / ".local" / "share")) / "aetherchat"


def _display_name(filename: str | None) -> str:
    value = Path(str(filename or "attachment").replace("\\", "/")).name
    value = re.sub(r"[\x00-\x1f\x7f]", "", value).strip()
    return value[:240] or "attachment"


@dataclass(frozen=True)
class ExtractedChunk:
    locator: str
    content: str


class KnowledgeError(ValueError):
    pass


class KnowledgeService:
    def __init__(self, base_path: str | os.PathLike[str] | None = None):
        self.data_root = ensure_private_data_dir(_data_root(base_path))
        self.base_path = self.data_root / "knowledge"
        self.db_path = self.data_root / "knowledge.sqlite3"
        self._lock = RLock()
        self.base_path.mkdir(parents=True, exist_ok=True)
        self._initialize()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.db_path, timeout=15)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA busy_timeout = 15000")
        harden_private_file(self.db_path)
        return connection

    def _initialize(self) -> None:
        with self._lock, self._connect() as connection:
            connection.execute("PRAGMA journal_mode = WAL")
            harden_private_file(self.db_path.with_name(f"{self.db_path.name}-wal"))
            harden_private_file(self.db_path.with_name(f"{self.db_path.name}-shm"))
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS knowledge_attachments (
                    id TEXT PRIMARY KEY,
                    session_id TEXT NOT NULL,
                    name TEXT NOT NULL,
                    extension TEXT NOT NULL,
                    content_type TEXT NOT NULL,
                    size_bytes INTEGER NOT NULL,
                    sha256 TEXT NOT NULL,
                    status TEXT NOT NULL,
                    chunk_count INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_knowledge_attachments_session
                    ON knowledge_attachments(session_id, created_at DESC);
                CREATE VIRTUAL TABLE IF NOT EXISTS knowledge_chunks USING fts5(
                    id UNINDEXED,
                    session_id UNINDEXED,
                    attachment_id UNINDEXED,
                    source_name,
                    locator,
                    content,
                    tokenize='unicode61'
                );
                CREATE TABLE IF NOT EXISTS knowledge_vectors (
                    id TEXT PRIMARY KEY,
                    session_id TEXT NOT NULL,
                    attachment_id TEXT NOT NULL,
                    vector BLOB NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_knowledge_vectors_session
                    ON knowledge_vectors(session_id, attachment_id);
                """
            )
            missing = connection.execute(
                """SELECT id, session_id, attachment_id, content FROM knowledge_chunks
                   WHERE id NOT IN (SELECT id FROM knowledge_vectors)"""
            ).fetchall()
            connection.executemany(
                "INSERT OR IGNORE INTO knowledge_vectors (id, session_id, attachment_id, vector) VALUES (?, ?, ?, ?)",
                [
                    (row["id"], row["session_id"], row["attachment_id"], self._vector_blob(row["content"]))
                    for row in missing
                ],
            )

    @staticmethod
    def _session_id(session_id: str) -> str:
        return ArtifactService().normalize_session_id(session_id)

    def _attachment_root(self, session_id: str, attachment_id: str) -> Path:
        safe_session = self._session_id(session_id)
        try:
            safe_attachment = str(UUID(str(attachment_id)))
        except ValueError as exc:
            raise KnowledgeError("附件 ID 无效") from exc
        root = self.base_path.resolve()
        path = (root / safe_session / safe_attachment).resolve()
        try:
            path.relative_to(root)
        except ValueError as exc:
            raise KnowledgeError("附件路径越界") from exc
        return path

    @staticmethod
    def validate_file(name: str, data: bytes) -> str:
        extension = Path(name).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise KnowledgeError("不支持此文件类型。支持 PNG/JPEG/WebP、常见文本/代码、PDF 和 Office 文件。")
        if not data:
            raise KnowledgeError("不能上传空文件")
        if len(data) > MAX_UPLOAD_BYTES:
            raise KnowledgeError("文件超过 25 MB 限制")
        if extension == ".pdf" and not data.startswith(b"%PDF-"):
            raise KnowledgeError("文件扩展名与 PDF 内容不匹配")
        if extension in OFFICE_EXTENSIONS:
            if not data.startswith(b"PK"):
                raise KnowledgeError("文件扩展名与 Office Open XML 内容不匹配")
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    expanded = sum(item.file_size for item in archive.infolist())
                    if expanded > MAX_ARCHIVE_EXPANDED_BYTES:
                        raise KnowledgeError("Office 文件解压后超过 100 MB 安全限制")
            except zipfile.BadZipFile as exc:
                raise KnowledgeError("Office 文件已损坏或格式无效") from exc
        if extension in IMAGE_EXTENSIONS:
            try:
                with Image.open(io.BytesIO(data)) as image:
                    image.verify()
                with Image.open(io.BytesIO(data)) as image:
                    if image.width * image.height > 40_000_000:
                        raise KnowledgeError("图片超过 4000 万像素安全限制")
                    expected = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP"}[extension]
                    if image.format != expected:
                        raise KnowledgeError("图片扩展名与实际格式不匹配")
            except KnowledgeError:
                raise
            except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
                raise KnowledgeError("图片无法解析、已损坏或尺寸不安全") from exc
        if extension in TEXT_EXTENSIONS | {".csv"} and b"\x00" in data[:8192]:
            raise KnowledgeError("文本文件包含二进制内容")
        return extension

    def create_attachment(
        self,
        session_id: str,
        filename: str | None,
        content_type: str | None,
        data: bytes,
    ) -> dict[str, Any]:
        import hashlib

        safe_session = self._session_id(session_id)
        name = _display_name(filename)
        extension = self.validate_file(name, data)
        if extension in IMAGE_EXTENSIONS:
            data = self._sanitize_image(data, extension)
        attachment_id = str(uuid4())
        root = self._attachment_root(safe_session, attachment_id)
        root.mkdir(parents=True, exist_ok=False)
        source_path = root / f"source{extension}"
        try:
            source_path.write_bytes(data)
            chunks = self._extract(source_path, extension)
            created_at = _utc_now()
            with self._lock, self._connect() as connection:
                connection.execute(
                    """INSERT INTO knowledge_attachments
                       (id, session_id, name, extension, content_type, size_bytes, sha256, status, chunk_count, created_at)
                       VALUES (?, ?, ?, ?, ?, ?, ?, 'ready', ?, ?)""",
                    (
                        attachment_id, safe_session, name, extension,
                        str(content_type or "application/octet-stream")[:200], len(data),
                        hashlib.sha256(data).hexdigest(), len(chunks), created_at,
                    ),
                )
                indexed_chunks = [
                    (str(uuid4()), safe_session, attachment_id, name, chunk.locator, chunk.content)
                    for chunk in chunks
                ]
                connection.executemany(
                    """INSERT INTO knowledge_chunks
                       (id, session_id, attachment_id, source_name, locator, content)
                       VALUES (?, ?, ?, ?, ?, ?)""",
                    indexed_chunks,
                )
                connection.executemany(
                    "INSERT INTO knowledge_vectors (id, session_id, attachment_id, vector) VALUES (?, ?, ?, ?)",
                    [
                        (chunk_id, chunk_session, chunk_attachment, self._vector_blob(content))
                        for chunk_id, chunk_session, chunk_attachment, _name, _locator, content in indexed_chunks
                    ],
                )
            return self.get_attachment(safe_session, attachment_id)
        except Exception:
            shutil.rmtree(root, ignore_errors=True)
            raise

    def list_attachments(self, session_id: str) -> list[dict[str, Any]]:
        safe_session = self._session_id(session_id)
        with self._lock, self._connect() as connection:
            rows = connection.execute(
                "SELECT * FROM knowledge_attachments WHERE session_id = ? ORDER BY created_at DESC",
                (safe_session,),
            ).fetchall()
        return [dict(row) for row in rows]

    def get_attachment(self, session_id: str, attachment_id: str) -> dict[str, Any]:
        safe_session = self._session_id(session_id)
        with self._lock, self._connect() as connection:
            row = connection.execute(
                "SELECT * FROM knowledge_attachments WHERE session_id = ? AND id = ?",
                (safe_session, str(attachment_id)),
            ).fetchone()
        if row is None:
            raise KeyError("附件不存在或不属于当前会话")
        return dict(row)

    def preview_attachment(self, session_id: str, attachment_id: str, limit: int = 20) -> dict[str, Any]:
        item = self.get_attachment(session_id, attachment_id)
        if item["extension"] in IMAGE_EXTENSIONS:
            return {"attachment": item, "kind": "image", "chunks": [], "truncated": False}
        bounded = max(1, min(50, int(limit)))
        with self._lock, self._connect() as connection:
            rows = connection.execute(
                """SELECT locator, content FROM knowledge_chunks
                   WHERE session_id = ? AND attachment_id = ? LIMIT ?""",
                (self._session_id(session_id), str(attachment_id), bounded + 1),
            ).fetchall()
        return {
            "attachment": item,
            "kind": "text",
            "chunks": [dict(row) for row in rows[:bounded]],
            "truncated": len(rows) > bounded,
        }

    def image_content(self, session_id: str, attachment_id: str) -> tuple[bytes, str]:
        item = self.get_attachment(session_id, attachment_id)
        extension = item["extension"]
        if extension not in IMAGE_EXTENSIONS:
            raise KnowledgeError("仅安全重编码后的图片可以直接预览")
        data = (self._attachment_root(session_id, attachment_id) / f"source{extension}").read_bytes()
        mime = "image/jpeg" if extension in {".jpg", ".jpeg"} else f"image/{extension[1:]}"
        return data, mime

    def image_path(self, session_id: str, attachment_id: str) -> str:
        """Return the validated native path for a session-owned image.

        Model providers only receive image pixels, never a trustworthy local
        filename.  Server-side tools such as ComfyUI can use this resolver to
        bridge an attachment without asking the model to invent a path.
        """

        item = self.get_attachment(session_id, attachment_id)
        extension = item["extension"]
        if extension not in IMAGE_EXTENSIONS:
            raise KnowledgeError("附件不是可用的 PNG、JPEG 或 WebP 图片")
        path = self._attachment_root(session_id, attachment_id) / f"source{extension}"
        if not path.is_file() or path.is_symlink():
            raise KnowledgeError("附件图片文件不存在或不安全")
        return str(path.resolve(strict=True))

    def image_paths(
        self,
        session_id: str,
        attachment_ids: Iterable[str],
    ) -> list[str]:
        paths: list[str] = []
        for reference in self.attachment_refs(session_id, attachment_ids):
            item = self.get_attachment(session_id, reference["id"])
            if item["extension"] in IMAGE_EXTENSIONS:
                paths.append(self.image_path(session_id, reference["id"]))
        return paths

    def attachment_refs(self, session_id: str, attachment_ids: Iterable[str]) -> list[dict[str, Any]]:
        refs = []
        seen: set[str] = set()
        for attachment_id in attachment_ids:
            value = str(attachment_id)
            if value in seen:
                continue
            seen.add(value)
            item = self.get_attachment(session_id, value)
            refs.append({
                "id": item["id"], "name": item["name"], "content_type": item["content_type"],
                "size_bytes": item["size_bytes"], "chunk_count": item["chunk_count"],
            })
        return refs

    def delete_attachment(self, session_id: str, attachment_id: str) -> None:
        safe_session = self._session_id(session_id)
        self.get_attachment(safe_session, attachment_id)
        with self._lock, self._connect() as connection:
            connection.execute(
                "DELETE FROM knowledge_vectors WHERE session_id = ? AND attachment_id = ?",
                (safe_session, str(attachment_id)),
            )
            connection.execute(
                "DELETE FROM knowledge_chunks WHERE session_id = ? AND attachment_id = ?",
                (safe_session, str(attachment_id)),
            )
            connection.execute(
                "DELETE FROM knowledge_attachments WHERE session_id = ? AND id = ?",
                (safe_session, str(attachment_id)),
            )
        shutil.rmtree(self._attachment_root(safe_session, attachment_id), ignore_errors=True)

    def delete_session(self, session_id: str) -> dict[str, int]:
        safe_session = self._session_id(session_id)
        with self._lock, self._connect() as connection:
            attachment_count = int(connection.execute(
                "SELECT COUNT(*) FROM knowledge_attachments WHERE session_id = ?",
                (safe_session,),
            ).fetchone()[0])
            connection.execute(
                "DELETE FROM knowledge_vectors WHERE session_id = ?",
                (safe_session,),
            )
            connection.execute(
                "DELETE FROM knowledge_chunks WHERE session_id = ?",
                (safe_session,),
            )
            connection.execute(
                "DELETE FROM knowledge_attachments WHERE session_id = ?",
                (safe_session,),
            )
        session_root = (self.base_path.resolve() / safe_session).resolve()
        try:
            session_root.relative_to(self.base_path.resolve())
        except ValueError as exc:
            raise KnowledgeError("会话资料路径越界") from exc
        if session_root.is_dir() and not session_root.is_symlink():
            shutil.rmtree(session_root)
        return {"attachments": attachment_count}

    def search(self, session_id: str, query: str, limit: int = 8) -> list[dict[str, Any]]:
        safe_session = self._session_id(session_id)
        clean_query = " ".join(str(query).strip().split())[:500]
        terms = re.findall(r"[\w\u3400-\u9fff]+", clean_query, flags=re.UNICODE)
        if not terms:
            return []
        bounded = max(1, min(30, int(limit)))
        match = " OR ".join(f'"{term.replace(chr(34), chr(34) * 2)}"' for term in terms[:12])
        with self._lock, self._connect() as connection:
            try:
                lexical_rows = connection.execute(
                    """SELECT id, attachment_id, source_name, locator, content, bm25(knowledge_chunks) AS rank
                       FROM knowledge_chunks
                       WHERE knowledge_chunks MATCH ? AND session_id = ?
                       ORDER BY rank LIMIT 60""",
                    (match, safe_session),
                ).fetchall()
            except sqlite3.OperationalError:
                lexical_rows = []
            if not lexical_rows:
                clauses = " OR ".join("content LIKE ?" for _ in terms[:8])
                lexical_rows = connection.execute(
                    f"""SELECT id, attachment_id, source_name, locator, content, 0 AS rank
                        FROM knowledge_chunks WHERE session_id = ? AND ({clauses}) LIMIT 60""",
                    (safe_session, *(f"%{term}%" for term in terms[:8])),
                ).fetchall()
            vector_rows = connection.execute(
                """SELECT k.id, k.attachment_id, k.source_name, k.locator, k.content, v.vector
                   FROM knowledge_vectors v JOIN knowledge_chunks k ON k.id = v.id
                   WHERE v.session_id = ? LIMIT 5000""",
                (safe_session,),
            ).fetchall()

        query_vector = self._vector(clean_query)
        vector_ranked = sorted(
            ((self._cosine_blob(query_vector, row["vector"]), row) for row in vector_rows),
            key=lambda item: item[0],
            reverse=True,
        )[:60]
        candidates: dict[str, dict[str, Any]] = {}
        for index, row in enumerate(lexical_rows):
            candidates[row["id"]] = {
                "row": row, "lexical": 1.0 / (1.0 + index), "vector": 0.0,
                "match_types": {"全文"},
            }
        for similarity, row in vector_ranked:
            candidate = candidates.setdefault(row["id"], {
                "row": row, "lexical": 0.0, "vector": 0.0, "match_types": set(),
            })
            candidate["vector"] = max(0.0, float(similarity))

        lowered_query = clean_query.casefold()
        query_terms = [term.casefold() for term in terms[:12]]
        compact_query = re.sub(r"\s+", "", lowered_query)
        query_bigrams = {compact_query[index:index + 2] for index in range(max(0, len(compact_query) - 1))}
        ranked: list[tuple[float, dict[str, Any]]] = []
        for candidate in candidates.values():
            row = candidate["row"]
            content = str(row["content"])
            lowered = content.casefold()
            coverage = sum(1 for term in query_terms if term in lowered) / max(1, len(query_terms))
            exact = 1.0 if lowered_query and lowered_query in lowered else 0.0
            compact_content = re.sub(r"\s+", "", lowered)
            bigram_coverage = (
                sum(1 for gram in query_bigrams if gram in compact_content) / len(query_bigrams)
                if query_bigrams else 0.0
            )
            semantic = max(candidate["vector"], bigram_coverage * 0.85)
            if semantic > 0.05:
                candidate["match_types"].add("哈希相似")
            if candidate["lexical"] <= 0 and semantic <= 0.05:
                continue
            score = candidate["lexical"] * 0.50 + semantic * 0.35 + coverage * 0.10 + exact * 0.05
            ranked.append((score, candidate))
        ranked.sort(key=lambda item: item[0], reverse=True)

        results: list[dict[str, Any]] = []
        for score, candidate in ranked[:bounded]:
            row = candidate["row"]
            content = str(row["content"])
            lowered = content.casefold()
            coverage = sum(1 for term in query_terms if term in lowered) / max(1, len(query_terms))
            exact = 1.0 if lowered_query and lowered_query in lowered else 0.0
            results.append({
                "attachment_id": row["attachment_id"],
                "source_name": row["source_name"],
                "locator": row["locator"],
                "content": row["content"],
                "citation": f"{row['source_name']} · {row['locator']}",
                "score": round(max(0.0, min(1.0, score)), 4),
                "match_types": sorted(candidate["match_types"]),
                "rank": len(results) + 1,
                "layers": {
                    "lexical": round(float(candidate["lexical"]), 4),
                    "vector": round(float(candidate["vector"]), 4),
                    "coverage": round(float(coverage), 4),
                    "exact": round(float(exact), 4),
                },
            })
        return results

    def session_stats(self, session_id: str) -> dict[str, Any]:
        """Summarize the local knowledge index for a session (UI status strip)."""
        safe_session = self._session_id(session_id)
        with self._lock, self._connect() as connection:
            attachments = connection.execute(
                """SELECT extension, status, size_bytes, chunk_count
                   FROM knowledge_attachments WHERE session_id = ?""",
                (safe_session,),
            ).fetchall()
            chunk_count = connection.execute(
                "SELECT COUNT(*) AS n FROM knowledge_chunks WHERE session_id = ?",
                (safe_session,),
            ).fetchone()["n"]
            vector_count = connection.execute(
                "SELECT COUNT(*) AS n FROM knowledge_vectors WHERE session_id = ?",
                (safe_session,),
            ).fetchone()["n"]
        by_ext: dict[str, int] = {}
        total_bytes = 0
        ready = 0
        for row in attachments:
            ext = str(row["extension"] or "").lower() or "(none)"
            by_ext[ext] = by_ext.get(ext, 0) + 1
            total_bytes += int(row["size_bytes"] or 0)
            if str(row["status"] or "") == "ready":
                ready += 1
        return {
            "session_id": safe_session,
            "attachment_count": len(attachments),
            "ready_count": ready,
            "chunk_count": int(chunk_count or 0),
            "vector_count": int(vector_count or 0),
            "total_bytes": total_bytes,
            "extensions": by_ext,
            "retrieval": {
                "lexical": "fts5+bm25",
                "vector": "local-hash-256d",
                "hybrid": "lexical*0.5 + semantic*0.35 + coverage*0.1 + exact*0.05",
                "offline": True,
            },
        }

    @staticmethod
    def _vector_features(text: str) -> list[str]:
        normalized = " ".join(str(text).casefold().split())
        words = re.findall(r"[a-z0-9_]{2,}|[\u3400-\u9fff]", normalized)
        compact = re.sub(r"\s+", "", normalized)
        ngrams = [compact[index:index + 3] for index in range(max(0, len(compact) - 2))]
        return words + ngrams[:6000]

    @classmethod
    def _vector(cls, text: str) -> list[float]:
        values = [0.0] * VECTOR_DIMENSIONS
        for feature in cls._vector_features(text):
            digest = hashlib.blake2b(feature.encode("utf-8"), digest_size=8).digest()
            number = int.from_bytes(digest, "little")
            index = number % VECTOR_DIMENSIONS
            values[index] += 1.0 if number & (1 << 63) else -1.0
        norm = math.sqrt(sum(value * value for value in values)) or 1.0
        return [value / norm for value in values]

    @classmethod
    def _vector_blob(cls, text: str) -> bytes:
        return array("f", cls._vector(text)).tobytes()

    @staticmethod
    def _cosine_blob(query_vector: list[float], blob: bytes) -> float:
        stored = array("f")
        stored.frombytes(blob)
        if len(stored) != len(query_vector):
            return 0.0
        return sum(left * right for left, right in zip(query_vector, stored, strict=True))

    def context_for_attachments(
        self, session_id: str, attachment_ids: Iterable[str], query: str, max_chars: int = 12_000
    ) -> str:
        ids = [ref["id"] for ref in self.attachment_refs(session_id, attachment_ids)]
        if not ids:
            return ""
        candidates = [item for item in self.search(session_id, query, limit=18) if item["attachment_id"] in ids]
        if not candidates:
            placeholders = ",".join("?" for _ in ids)
            with self._lock, self._connect() as connection:
                rows = connection.execute(
                    f"""SELECT attachment_id, source_name, locator, content FROM knowledge_chunks
                        WHERE session_id = ? AND attachment_id IN ({placeholders}) LIMIT 12""",
                    (self._session_id(session_id), *ids),
                ).fetchall()
            candidates = [
                {**dict(row), "citation": f"{row['source_name']} · {row['locator']}"} for row in rows
            ]
        parts = ["以下是用户本条消息附带的项目资料。回答时仅据此资料支持相关事实，并用 [来源: 文件名 · 位置] 标注引用："]
        used = len(parts[0])
        for item in candidates:
            section = f"\n\n[来源: {item['citation']}]\n{item['content']}"
            if used + len(section) > max_chars:
                break
            parts.append(section)
            used += len(section)
        return "".join(parts)

    def explicit_attachment_context(
        self,
        session_id: str,
        attachment_ids: Iterable[str],
        max_chars: int = 8_000,
    ) -> str:
        """Return ordered chunks for files the user explicitly cited."""

        ids: list[str] = []
        seen: set[str] = set()
        for attachment_id in attachment_ids:
            try:
                item = self.get_attachment(session_id, str(attachment_id))
            except (KeyError, KnowledgeError, ValueError):
                continue
            if item["id"] in seen:
                continue
            seen.add(item["id"])
            ids.append(item["id"])
        if not ids:
            return ""
        placeholders = ",".join("?" for _ in ids)
        with self._lock, self._connect() as connection:
            rows = connection.execute(
                f"""SELECT source_name, locator, content FROM knowledge_chunks
                    WHERE session_id = ? AND attachment_id IN ({placeholders})
                    ORDER BY attachment_id, locator LIMIT 40""",
                (self._session_id(session_id), *ids),
            ).fetchall()
        if not rows:
            return ""
        parts: list[str] = []
        used = 0
        for row in rows:
            section = f"[来源: {row['source_name']} · {row['locator']}]\n{row['content']}"
            if used + len(section) + 2 > max_chars:
                break
            parts.append(section)
            used += len(section) + 2
        return "\n\n".join(parts)

    def context_for_query(
        self,
        session_id: str,
        query: str,
        max_chars: int = 8_000,
    ) -> str:
        """Retrieve relevant evidence from all local session attachments.

        This is intentionally local-only: it uses the existing FTS/hash-vector
        index and never sends project documents to an embedding service.
        """

        candidates = [
            item for item in self.search(session_id, query, limit=12)
            if float(item.get("score") or 0) >= 0.18
        ]
        if not candidates:
            return ""
        heading = (
            "以下内容由 AetherChat 从当前会话的本地知识库自动检索。"
            "只在与问题相关时使用；事实结论必须用 [来源: 文件名 · 位置] 标注，"
            "若证据不足应明确说明而不是猜测："
        )
        parts = [heading]
        used = len(heading)
        seen: set[tuple[str, str]] = set()
        for item in candidates:
            identity = (str(item["attachment_id"]), str(item["locator"]))
            if identity in seen:
                continue
            seen.add(identity)
            section = f"\n\n[来源: {item['citation']}]\n{item['content']}"
            if used + len(section) > max_chars:
                break
            parts.append(section)
            used += len(section)
        return "".join(parts) if len(parts) > 1 else ""

    def image_data_urls(self, session_id: str, attachment_ids: Iterable[str]) -> list[str]:
        urls: list[str] = []
        total = 0
        for reference in self.attachment_refs(session_id, attachment_ids):
            item = self.get_attachment(session_id, reference["id"])
            extension = item["extension"]
            if extension not in IMAGE_EXTENSIONS:
                continue
            data = (self._attachment_root(session_id, item["id"]) / f"source{extension}").read_bytes()
            total += len(data)
            if total > 20 * 1024 * 1024 or len(urls) >= 4:
                break
            mime = "image/jpeg" if extension in {".jpg", ".jpeg"} else f"image/{extension[1:]}"
            urls.append(f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}")
        return urls

    def _extract(self, path: Path, extension: str) -> list[ExtractedChunk]:
        if extension in IMAGE_EXTENSIONS:
            return []
        if extension in TEXT_EXTENSIONS:
            text = self._decode_text(path.read_bytes())
            if extension == ".json":
                try:
                    text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
                except json.JSONDecodeError as exc:
                    raise KnowledgeError(f"JSON 格式无效：{exc.msg}") from exc
            return self._chunk_text(text, "全文")
        if extension == ".csv":
            return self._extract_csv(path)
        if extension == ".pdf":
            return self._extract_pdf(path)
        if extension == ".docx":
            return self._extract_docx(path)
        if extension == ".xlsx":
            return self._extract_xlsx(path)
        if extension == ".pptx":
            return self._extract_pptx(path)
        raise KnowledgeError("不支持此文件类型")

    @staticmethod
    def _sanitize_image(data: bytes, extension: str) -> bytes:
        output = io.BytesIO()
        with Image.open(io.BytesIO(data)) as image:
            image.load()
            image.thumbnail((8192, 8192))
            image_format = "JPEG" if extension in {".jpg", ".jpeg"} else extension[1:].upper()
            safe_image = (
                image.convert("RGB")
                if image_format == "JPEG" and image.mode not in {"RGB", "L"}
                else image.copy()
            )
            save_options = {"quality": 92, "optimize": True} if image_format in {"JPEG", "WEBP"} else {"optimize": True}
            safe_image.save(output, format=image_format, **save_options)
        sanitized = output.getvalue()
        if len(sanitized) > MAX_UPLOAD_BYTES:
            raise KnowledgeError("规范化后的图片仍超过 25 MB 限制")
        return sanitized

    @staticmethod
    def _decode_text(data: bytes) -> str:
        for encoding in ("utf-8-sig", "gb18030", "utf-16", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise KnowledgeError("无法识别文本编码")

    @staticmethod
    def _chunk_text(text: str, locator: str, size: int = 1400, overlap: int = 120) -> list[ExtractedChunk]:
        clean = text.replace("\x00", "").strip()[:MAX_EXTRACTED_CHARS]
        if not clean:
            return []
        chunks = []
        start = 0
        part = 1
        while start < len(clean):
            end = min(len(clean), start + size)
            if end < len(clean):
                boundary = max(clean.rfind("\n", start + size // 2, end), clean.rfind("。", start + size // 2, end))
                if boundary > start:
                    end = boundary + 1
            label = locator if len(clean) <= size else f"{locator} · 片段 {part}"
            chunks.append(ExtractedChunk(label, clean[start:end].strip()))
            if end >= len(clean):
                break
            start = max(start + 1, end - overlap)
            part += 1
        return chunks

    def _extract_csv(self, path: Path) -> list[ExtractedChunk]:
        text = self._decode_text(path.read_bytes())
        rows = list(csv.reader(io.StringIO(text)))[:10_000]
        chunks = []
        for start in range(0, len(rows), 50):
            output = io.StringIO()
            csv.writer(output).writerows(rows[start:start + 50])
            chunks.extend(self._chunk_text(output.getvalue(), f"行 {start + 1}-{min(start + 50, len(rows))}"))
        return chunks

    def _extract_pdf(self, path: Path) -> list[ExtractedChunk]:
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted and not reader.decrypt(""):
                raise KnowledgeError("暂不支持有密码的 PDF")
            chunks = []
            for index, page in enumerate(reader.pages[:250]):
                chunks.extend(self._chunk_text(page.extract_text() or "", f"第 {index + 1} 页"))
            return chunks
        except KnowledgeError:
            raise
        except Exception as exc:
            raise KnowledgeError("PDF 无法解析或已损坏") from exc

    def _extract_docx(self, path: Path) -> list[ExtractedChunk]:
        document = Document(str(path))
        chunks = []
        for index, paragraph in enumerate(document.paragraphs[:10_000]):
            chunks.extend(self._chunk_text(paragraph.text, f"段落 {index + 1}"))
        for table_index, table in enumerate(document.tables[:200]):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows[:2000]]
            chunks.extend(self._chunk_text("\n".join(rows), f"表格 {table_index + 1}"))
        return chunks

    def _extract_xlsx(self, path: Path) -> list[ExtractedChunk]:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        chunks = []
        total_cells = 0
        try:
            for sheet in workbook.worksheets[:50]:
                batch: list[str] = []
                batch_start = 1
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    values = ["" if value is None else str(value) for value in row[:200]]
                    total_cells += len(values)
                    if total_cells > 100_000 or row_index > 10_000:
                        break
                    batch.append("\t".join(values).rstrip())
                    if len(batch) == 50:
                        chunks.extend(self._chunk_text("\n".join(batch), f"工作表 {sheet.title} · 行 {batch_start}-{row_index}"))
                        batch, batch_start = [], row_index + 1
                if batch:
                    chunks.extend(self._chunk_text("\n".join(batch), f"工作表 {sheet.title} · 行 {batch_start}-{batch_start + len(batch) - 1}"))
                if total_cells > 100_000:
                    break
        finally:
            workbook.close()
        return chunks

    def _extract_pptx(self, path: Path) -> list[ExtractedChunk]:
        presentation = Presentation(str(path))
        chunks = []
        for index, slide in enumerate(presentation.slides[:500]):
            texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            chunks.extend(self._chunk_text("\n".join(texts), f"第 {index + 1} 张幻灯片"))
        return chunks


knowledge_service = KnowledgeService()
