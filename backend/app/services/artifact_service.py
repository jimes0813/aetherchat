from __future__ import annotations

import base64
import builtins
import csv
import hashlib
import io
import json
import mimetypes
import os
import re
import shutil
import zipfile
from copy import deepcopy
from datetime import UTC, datetime
from difflib import unified_diff
from pathlib import Path, PureWindowsPath
from threading import RLock
from typing import Any
from uuid import NAMESPACE_URL, uuid4, uuid5

HISTORY_FILENAME = ".aether_history.jsonl"
MERGE_RULES = {"two_votes", "facilitator", "manual_confirmation"}
MAX_ARTIFACT_BYTES = 5 * 1024 * 1024
MAX_GENERATED_FILE_BYTES = 512 * 1024 * 1024
MAX_HISTORY_SNAPSHOT_BYTES = 5 * 1024 * 1024
MAX_PREVIEW_CHARS = 300_000
EDITABLE_SUFFIXES = {
    ".md", ".markdown", ".txt", ".html", ".htm", ".svg", ".json", ".csv",
    ".tsv", ".py", ".js", ".jsx", ".ts", ".tsx", ".css", ".xml", ".yaml",
    ".yml", ".toml", ".ini", ".sh", ".sql", ".prompt",
}
OFFICE_SUFFIXES = {".docx", ".xlsx", ".pptx"}
_locks: dict[str, RLock] = {}
_locks_guard = RLock()


def _utc_now() -> str:
    return datetime.now(UTC).isoformat()


def _data_root(base_path: str | os.PathLike[str] | None = None) -> Path:
    configured = base_path or os.environ.get("AETHERCHAT_DATA_DIR")
    if configured:
        return Path(configured).expanduser().resolve()
    return (Path(os.environ.get("LOCALAPPDATA", ".")) / "AetherChat").resolve()


def _lock_for(path: Path) -> RLock:
    key = str(path.resolve())
    with _locks_guard:
        return _locks.setdefault(key, RLock())


class ArtifactPathError(ValueError):
    """Raised when an artifact/session path could escape its workspace."""


class ArtifactService:
    """Desktop-friendly, session-isolated artifact storage.

    Every public path operation goes through the same resolver.  Apart from
    preventing ``..``/absolute-path traversal, the resolver rejects symlinked
    path components so a file cannot be redirected outside the session after
    it has been validated.
    """

    def __init__(self, base_path: str | os.PathLike[str] | None = None):
        self.data_root = _data_root(base_path)
        self.base_path = self.data_root / "artifacts"

    def normalize_session_id(self, session_id: str) -> str:
        value = str(session_id or "").strip()
        if (
            not value
            or value in {".", ".."}
            or "/" in value
            or "\\" in value
            or ":" in value
            or "\x00" in value
        ):
            raise ArtifactPathError("Invalid session_id")
        return value

    def normalize_path(self, path: str) -> str:
        if not isinstance(path, str):
            raise ArtifactPathError("Artifact path must be a string")
        raw = path.strip().replace("\\", "/")
        windows_path = PureWindowsPath(path)
        if (
            not raw
            or raw.startswith("/")
            or windows_path.is_absolute()
            or bool(windows_path.drive)
            or "\x00" in raw
        ):
            raise ArtifactPathError("Artifact path must be relative to the session workspace")

        parts = [part for part in raw.split("/") if part not in {"", "."}]
        if not parts or any(part == ".." for part in parts):
            raise ArtifactPathError("Artifact path traversal is not allowed")
        if any(":" in part for part in parts):
            raise ArtifactPathError("Artifact path contains an invalid component")

        normalized = "/".join(parts)
        if normalized == HISTORY_FILENAME:
            raise ArtifactPathError("Artifact metadata files are reserved")
        return normalized

    def _session_root(self, session_id: str) -> Path:
        safe_session_id = self.normalize_session_id(session_id)
        base = self.base_path.resolve()
        candidate = (base / safe_session_id).resolve()
        try:
            candidate.relative_to(base)
        except ValueError as exc:  # defence in depth for unusual filesystem paths
            raise ArtifactPathError("Session path escapes the artifact root") from exc
        return candidate

    def _artifact_path(self, session_id: str, path: str) -> tuple[Path, str, Path]:
        root = self._session_root(session_id)
        normalized = self.normalize_path(path)

        cursor = root
        for component in normalized.split("/"):
            cursor = cursor / component
            if cursor.exists() and cursor.is_symlink():
                raise ArtifactPathError("Symlinked artifact paths are not allowed")

        candidate = (root / Path(*normalized.split("/"))).resolve()
        try:
            candidate.relative_to(root.resolve())
        except ValueError as exc:
            raise ArtifactPathError("Artifact path escapes the session workspace") from exc
        return candidate, normalized, root

    async def prepare_checkout(self, session_id: str) -> str:
        repo_path = self._session_root(session_id)
        if not repo_path.exists():
            repo_path.mkdir(parents=True, exist_ok=True)
            (repo_path / "README.md").write_text(
                "# AetherChat Project\n\nAetherChat desktop workspace initialized.\n",
                encoding="utf-8",
            )
            (repo_path / "docs").mkdir(exist_ok=True)
            (repo_path / "docs" / "PRD.md").write_text(
                "# Product Notes\n\n- Multi-model collaboration\n- Task claiming and review\n- Artifact editing\n",
                encoding="utf-8",
            )
            (repo_path / "src").mkdir(exist_ok=True)
            (repo_path / "src" / "demo.ts").write_text(
                "export const greeting = 'hello from AetherChat';\n",
                encoding="utf-8",
            )
            self._append_history(repo_path, "README.md", "Initial commit", "system")
            self._append_history(repo_path, "docs/PRD.md", "Initial commit", "system")
            self._append_history(repo_path, "src/demo.ts", "Initial commit", "system")
        elif not repo_path.is_dir():
            raise ArtifactPathError("Session workspace is not a directory")
        return str(repo_path)

    async def write(
        self,
        session_id: str,
        path: str,
        content: str,
        commit_message: str,
        author: str,
        *,
        source_model: str | None = None,
        status: str = "ready",
        parent_version: int | None = None,
    ) -> dict[str, Any]:
        encoded = content.encode("utf-8")
        if len(encoded) > MAX_ARTIFACT_BYTES:
            raise ArtifactPathError("Artifact exceeds the 5 MB editing limit")
        if status not in {"draft", "ready", "archived"}:
            raise ArtifactPathError("Invalid artifact status")
        await self.prepare_checkout(session_id)
        full_path, normalized, repo_path = self._artifact_path(session_id, path)
        full_path.parent.mkdir(parents=True, exist_ok=True)
        # Preserve the exact UTF-8 snapshot on every platform. Text-mode
        # writes would translate LF to CRLF on Windows and could translate a
        # restored CRLF snapshot a second time, producing extra blank lines.
        full_path.write_bytes(encoded)
        entry = self._append_history(
            repo_path,
            normalized,
            commit_message,
            author,
            source_model=source_model,
            status=status,
            parent_version=parent_version,
        )
        return {
            "hash": entry["hash"],
            "short_hash": entry["short_hash"],
            "path": normalized,
            "version": entry["version"],
            "updated_at": entry["created_at"],
            "mime_type": entry["mime_type"],
            "status": entry["status"],
            "source_model": entry["source_model"],
            "parent_version": entry["parent_version"],
            "size_bytes": entry["size_bytes"],
        }

    async def import_file(
        self,
        session_id: str,
        path: str,
        source_path: str | os.PathLike[str],
        commit_message: str,
        author: str,
        *,
        source_model: str | None = None,
        status: str = "ready",
    ) -> dict[str, Any]:
        source = Path(source_path).expanduser().resolve(strict=True)
        if not source.is_file() or source.is_symlink():
            raise ArtifactPathError("Generated file source must be a regular file")
        size_bytes = source.stat().st_size
        if size_bytes > MAX_GENERATED_FILE_BYTES:
            raise ArtifactPathError("Generated file exceeds the 512 MB import limit")
        if status not in {"draft", "ready", "archived"}:
            raise ArtifactPathError("Invalid artifact status")
        await self.prepare_checkout(session_id)
        full_path, normalized, repo_path = self._artifact_path(session_id, path)
        full_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = full_path.with_name(f".{full_path.name}.{uuid4().hex}.tmp")
        try:
            shutil.copyfile(source, temporary)
            os.replace(temporary, full_path)
        finally:
            temporary.unlink(missing_ok=True)
        entry = self._append_history(
            repo_path,
            normalized,
            commit_message,
            author,
            source_model=source_model,
            status=status,
        )
        return {
            "hash": entry["hash"],
            "short_hash": entry["short_hash"],
            "path": normalized,
            "version": entry["version"],
            "updated_at": entry["created_at"],
            "mime_type": entry["mime_type"],
            "status": entry["status"],
            "source_model": entry["source_model"],
            "parent_version": entry["parent_version"],
            "size_bytes": entry["size_bytes"],
        }

    async def native_location(self, session_id: str, path: str) -> dict[str, Any]:
        await self.prepare_checkout(session_id)
        full_path, normalized, _ = self._artifact_path(session_id, path)
        if not full_path.is_file() or full_path.is_symlink():
            raise FileNotFoundError(path)
        return {
            "path": str(full_path),
            "artifact_path": normalized,
            "display_name": full_path.name,
            "mime_type": self._mime_for(normalized),
            "size_bytes": full_path.stat().st_size,
        }

    def delete_session(self, session_id: str) -> bool:
        root = self._session_root(session_id)
        if not root.exists():
            return False
        if not root.is_dir() or root.is_symlink():
            raise ArtifactPathError("Session artifact root is not a safe directory")
        shutil.rmtree(root)
        return True

    async def read(self, session_id: str, path: str) -> str:
        await self.prepare_checkout(session_id)
        full_path, _, _ = self._artifact_path(session_id, path)
        if not full_path.is_file():
            raise FileNotFoundError(str(path))
        return full_path.read_text(encoding="utf-8")

    async def list_files(self, session_id: str) -> list[dict[str, Any]]:
        repo_path = Path(await self.prepare_checkout(session_id))
        files: list[dict[str, Any]] = []
        history_rows = self._read_history(repo_path)
        history_by_path: dict[str, list[dict[str, Any]]] = {}
        for item in history_rows:
            try:
                history_path = self.normalize_path(str(item.get("path") or ""))
            except ArtifactPathError:
                continue
            history_by_path.setdefault(history_path, []).append(item)
        current_paths: set[str] = set()
        for full_path in repo_path.rglob("*"):
            if (
                not full_path.is_file()
                or full_path.is_symlink()
                or full_path.name == HISTORY_FILENAME
            ):
                continue
            try:
                full_path.resolve().relative_to(repo_path.resolve())
            except ValueError:
                continue
            file_path = full_path.relative_to(repo_path).as_posix()
            current_paths.add(file_path)
            files.append(
                {
                    "id": str(uuid5(NAMESPACE_URL, f"aetherchat:{session_id}:{file_path}")),
                    "session_id": str(session_id),
                    "path": file_path,
                    "kind": self._kind_for(file_path),
                    "current_hash": self._hash_file(full_path),
                    "content_hash": self._hash_file(full_path),
                    "mime_type": self._mime_for(file_path),
                    "size_bytes": full_path.stat().st_size,
                    "updated_at": full_path.stat().st_mtime,
                    "availability": "current",
                    "recoverable": True,
                }
            )
            matching = history_by_path.get(file_path, [])
            if matching:
                latest = matching[-1]
                files[-1].update({
                    "current_version": latest.get("version", len(matching)),
                    "status": latest.get("status", "ready"),
                    "source_model": latest.get("source_model"),
                    "parent_version": latest.get("parent_version"),
                })
            else:
                files[-1].update({"current_version": 1, "status": "ready", "source_model": None, "parent_version": None})

        # Keep deleted/moved files discoverable when a safe snapshot still
        # exists in the append-only history. This turns a stale 404 into an
        # explicit, user-controlled restore path.
        for file_path, matching in history_by_path.items():
            if file_path in current_paths or not matching:
                continue
            latest = matching[-1]
            files.append({
                "id": str(uuid5(NAMESPACE_URL, f"aetherchat:{session_id}:{file_path}")),
                "session_id": str(session_id),
                "path": file_path,
                "kind": self._kind_for(file_path),
                "current_hash": str(latest.get("hash") or ""),
                "content_hash": str(latest.get("hash") or ""),
                "mime_type": self._mime_for(file_path),
                "size_bytes": int(latest.get("size_bytes") or 0),
                "updated_at": latest.get("created_at"),
                "current_version": int(latest.get("version") or len(matching)),
                "status": latest.get("status", "archived"),
                "source_model": latest.get("source_model"),
                "parent_version": latest.get("parent_version"),
                "availability": "history_only",
                "recoverable": latest.get("content_encoding") == "utf-8",
            })
        return sorted(files, key=lambda item: item["path"])

    async def preview(self, session_id: str, path: str) -> dict[str, Any]:
        await self.prepare_checkout(session_id)
        full_path, normalized, _ = self._artifact_path(session_id, path)
        if not full_path.is_file():
            versions = await self.history(session_id, normalized, limit=1)
            if not versions:
                raise FileNotFoundError(path)
            latest = versions[0]
            suffix = Path(normalized).suffix.lower()
            mime_type = self._mime_for(normalized)
            recoverable_version = (
                int(latest.get("version") or 0)
                if latest.get("content_encoding") == "utf-8"
                else None
            )
            if recoverable_version is None or suffix not in EDITABLE_SUFFIXES:
                return self._preview_result(
                    normalized,
                    mime_type,
                    "download_only",
                    "",
                    availability="history_only",
                    recoverable_version=recoverable_version,
                )
            content = self._history_text(latest)
            mode = (
                "markdown"
                if suffix in {".md", ".markdown"}
                else "isolated_html"
                if suffix in {".html", ".htm", ".svg"}
                else "code"
            )
            if mode == "isolated_html":
                content = self._isolate_markup(content)
            return self._preview_result(
                normalized,
                mime_type,
                mode,
                content,
                availability="history_only",
                recoverable_version=recoverable_version,
            )
        suffix = full_path.suffix.lower()
        mime_type = self._mime_for(normalized)
        if suffix in OFFICE_SUFFIXES:
            self._validate_office_archive(full_path)
            content = self._office_summary(full_path, suffix)
            return self._preview_result(normalized, mime_type, "office_summary", content)
        if suffix not in EDITABLE_SUFFIXES:
            return self._preview_result(normalized, mime_type, "download_only", "")
        if full_path.stat().st_size > MAX_ARTIFACT_BYTES:
            raise ArtifactPathError("Artifact exceeds the 5 MB preview limit")
        try:
            content = full_path.read_text(encoding="utf-8")
        except UnicodeDecodeError as exc:
            raise ArtifactPathError("Artifact is not valid UTF-8 text") from exc
        mode = "markdown" if suffix in {".md", ".markdown"} else "isolated_html" if suffix in {".html", ".htm", ".svg"} else "code"
        if mode == "isolated_html":
            content = self._isolate_markup(content)
        return self._preview_result(normalized, mime_type, mode, content)

    def _isolate_markup(self, content: str) -> str:
        # The iframe CSP and empty sandbox are the primary boundary. Removing
        # navigation/resource attributes here also prevents user-triggered or
        # meta-refresh network attempts inside the opaque-origin frame.
        content = re.sub(r"<\s*script\b[^>]*>.*?<\s*/\s*script\s*>", "", content, flags=re.IGNORECASE | re.DOTALL)
        content = re.sub(r"<\s*(?:meta|base|link)\b[^>]*>", "", content, flags=re.IGNORECASE)
        content = re.sub(
            r"\s+(?:href|src|srcset|action|formaction|poster|data|xlink:href)\s*=\s*(?:\"[^\"]*\"|'[^']*'|[^\s>]+)",
            "",
            content,
            flags=re.IGNORECASE,
        )
        content = re.sub(
            r"\s+on[a-z0-9_-]+\s*=\s*(?:\"[^\"]*\"|'[^']*'|[^\s>]+)",
            "",
            content,
            flags=re.IGNORECASE,
        )
        return content

    def _validate_office_archive(self, path: Path) -> None:
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > 10_000 or sum(item.file_size for item in members) > 50 * 1024 * 1024:
                    raise ArtifactPathError("Office artifact expands beyond the 50 MB safety limit")
                if any(item.file_size > 20 * 1024 * 1024 for item in members):
                    raise ArtifactPathError("Office artifact contains an oversized entry")
        except zipfile.BadZipFile as exc:
            raise ArtifactPathError("Office artifact is not a valid package") from exc

    def _preview_result(
        self,
        path: str,
        mime_type: str,
        mode: str,
        content: str,
        *,
        availability: str = "current",
        recoverable_version: int | None = None,
    ) -> dict[str, Any]:
        truncated = len(content) > MAX_PREVIEW_CHARS
        return {
            "path": path,
            "mime_type": mime_type,
            "mode": mode,
            "content": content[:MAX_PREVIEW_CHARS],
            "truncated": truncated,
            "sandbox": "allow-none" if mode == "isolated_html" else None,
            "availability": availability,
            "recoverable_version": recoverable_version,
        }

    async def diff(self, session_id: str, path: str, from_version: int, to_version: int | None = None) -> dict[str, Any]:
        normalized = self.normalize_path(path)
        versions = await self.history(session_id, normalized)
        by_version = {int(item["version"]): item for item in versions}
        target_version = to_version or max(by_version, default=0)
        if from_version not in by_version or target_version not in by_version:
            raise ArtifactPathError("Artifact version not found")
        before = self._history_text(by_version[from_version])
        after = self._history_text(by_version[target_version])
        changes = "".join(unified_diff(
            before.splitlines(keepends=True), after.splitlines(keepends=True),
            fromfile=f"{normalized}@v{from_version}", tofile=f"{normalized}@v{target_version}",
        ))
        return {"path": normalized, "from_version": from_version, "to_version": target_version, "diff": changes}

    async def restore(self, session_id: str, path: str, version: int, author: str = "user") -> dict[str, Any]:
        normalized = self.normalize_path(path)
        versions = await self.history(session_id, normalized)
        selected = next((item for item in versions if int(item.get("version", 0)) == version), None)
        if selected is None:
            raise ArtifactPathError("Artifact version not found")
        content = self._history_text(selected)
        return await self.write(
            session_id, normalized, content, f"Restore {normalized} to v{version}", author,
            source_model=selected.get("source_model"), status="ready", parent_version=version,
        )

    def _history_text(self, item: dict[str, Any]) -> str:
        if item.get("content_encoding") != "utf-8":
            raise ArtifactPathError("Binary artifact versions cannot be edited or restored")
        return str(item.get("content", ""))

    async def export(self, session_id: str, path: str, output_format: str) -> tuple[bytes, str, str]:
        normalized = self.normalize_path(path)
        fmt = output_format.lower().lstrip(".")
        stem = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(normalized).stem)[:80] or "artifact"
        full_path, _, _ = self._artifact_path(session_id, normalized)
        if not full_path.is_file():
            raise FileNotFoundError(path)
        source_suffix = full_path.suffix.lower()
        if fmt == "original" or source_suffix == f".{fmt}" and source_suffix in OFFICE_SUFFIXES:
            if full_path.stat().st_size > MAX_ARTIFACT_BYTES:
                raise ArtifactPathError("Artifact exceeds the 5 MB download limit")
            safe_suffix = source_suffix.lstrip(".") or "bin"
            return full_path.read_bytes(), f"{stem}.{safe_suffix}", self._mime_for(normalized)
        content = await self.read(session_id, normalized)
        if fmt == "docx":
            data = self._export_docx(content, Path(normalized).name)
            return data, f"{stem}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if fmt == "xlsx":
            data = self._export_xlsx(content, Path(normalized).suffix.lower())
            return data, f"{stem}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if fmt == "pptx":
            data = self._export_pptx(content, Path(normalized).name)
            return data, f"{stem}.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        if fmt in {"md", "txt", "html", "svg", "json", "csv"}:
            return content.encode("utf-8"), f"{stem}.{fmt}", self._mime_for(f"x.{fmt}")
        raise ArtifactPathError("Unsupported export format")

    def _export_docx(self, content: str, title: str) -> bytes:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.shared import Inches, Pt
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.header.paragraphs[0].text = title
        footer = section.footer.paragraphs[0]
        footer.text = "Generated by AetherChat"
        document.styles["Normal"].font.size = Pt(10.5)
        lines = content.splitlines()
        index = 0
        while index < len(lines):
            line = lines[index]
            if line.strip() == "---page---":
                document.add_section(WD_SECTION.NEW_PAGE)
            elif image_match := re.fullmatch(r"!\[([^\]]*)\]\((data:image/(?:png|jpeg|webp);base64,[A-Za-z0-9+/=\s]+)\)", line.strip(), flags=re.IGNORECASE):
                image_data = self._decode_data_image(image_match.group(2))
                paragraph = document.add_paragraph()
                run = paragraph.add_run()
                run.add_picture(io.BytesIO(image_data), width=Inches(5.8))
                if image_match.group(1):
                    document.add_paragraph(image_match.group(1), style="Caption")
            elif re.fullmatch(r"\[citation:\s*.+\]", line.strip(), flags=re.IGNORECASE):
                document.add_paragraph(line.strip()[1:-1], style="Quote")
            elif re.match(r"^#{1,6}\s+", line):
                hashes, text = line.split(" ", 1)
                document.add_heading(text.strip(), level=min(len(hashes), 6))
            elif line.lstrip().startswith(("- ", "* ")):
                document.add_paragraph(line.lstrip()[2:], style="List Bullet")
            elif "|" in line and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[index + 1]):
                headers = [cell.strip() for cell in line.strip(" |").split("|")]
                rows: list[list[str]] = []
                index += 2
                while index < len(lines) and "|" in lines[index]:
                    rows.append([cell.strip() for cell in lines[index].strip(" |").split("|")])
                    index += 1
                table = document.add_table(rows=1, cols=max(1, len(headers)))
                table.style = "Table Grid"
                for col, value in enumerate(headers):
                    table.rows[0].cells[col].text = value
                for values in rows:
                    cells = table.add_row().cells
                    for col, value in enumerate(values[:len(cells)]):
                        cells[col].text = value
                continue
            elif line.strip():
                document.add_paragraph(line)
            index += 1
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def _export_xlsx(self, content: str, suffix: str) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheets: list[tuple[str, list[list[Any]]]] = []
        if suffix == ".json":
            parsed = json.loads(content)
            if isinstance(parsed, dict) and isinstance(parsed.get("sheets"), list):
                for item in parsed["sheets"][:20]:
                    if isinstance(item, dict) and isinstance(item.get("rows"), list):
                        sheets.append((str(item.get("name", "Sheet")), item["rows"][:5000]))
            elif isinstance(parsed, list):
                if parsed and all(isinstance(row, dict) for row in parsed):
                    keys = list(dict.fromkeys(key for row in parsed for key in row))[:100]
                    sheets.append(("Data", [keys] + [[row.get(key) for key in keys] for row in parsed[:4999]]))
                else:
                    sheets.append(("Data", [[value] if not isinstance(value, list) else value for value in parsed[:5000]]))
        if not sheets:
            delimiter = "\t" if suffix == ".tsv" else ","
            sheets = [("Data", [row[:100] for row in csv.reader(io.StringIO(content), delimiter=delimiter)][:5000])]
        for sheet_index, (name, rows) in enumerate(sheets):
            safe_name = re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or f"Sheet{sheet_index + 1}"
            worksheet = workbook.create_sheet(safe_name)
            worksheet.freeze_panes = "A2"
            for row_index, row in enumerate(rows, 1):
                if not isinstance(row, list):
                    row = [row]
                for col_index, value in enumerate(row[:100], 1):
                    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
                        value = "'" + value
                    worksheet.cell(row_index, col_index, value)
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="5D3B20")
            for column in range(1, min(worksheet.max_column, 100) + 1):
                width = max((len(str(worksheet.cell(row, column).value or "")) for row in range(1, min(worksheet.max_row, 200) + 1)), default=8)
                worksheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), 48)
        if not workbook.worksheets:
            workbook.create_sheet("Data")
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _export_pptx(self, content: str, title: str) -> bytes:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        sections = [part.strip() for part in re.split(r"^---\s*$", content, flags=re.MULTILINE) if part.strip()]
        if not sections:
            sections = [content]
        if len(sections) > 100:
            raise ArtifactPathError("Presentation exceeds the 100-slide limit")
        for slide_index, section in enumerate(sections):
            lines = section.splitlines()
            heading = re.sub(r"^#+\s*", "", lines[0]).strip() if lines else f"Slide {slide_index + 1}"
            notes_lines: list[str] = []
            image_data: bytes | None = None
            body_lines: list[str] = []
            in_notes = False
            for line in lines[1:]:
                if line.strip().lower() in {":::notes", "notes:"}:
                    in_notes = True
                    continue
                image_match = re.fullmatch(r"!\[[^\]]*\]\((data:image/(?:png|jpeg|webp);base64,[A-Za-z0-9+/=\s]+)\)", line.strip(), flags=re.IGNORECASE)
                if image_match and image_data is None:
                    image_data = self._decode_data_image(image_match.group(1))
                    continue
                if not line.strip():
                    continue
                target = notes_lines if in_notes else body_lines
                target.append(re.sub(r"^[-*]\s+", "", line).strip())
            if sum(len(line) for line in body_lines) > 1800:
                raise ArtifactPathError(f"Slide {slide_index + 1} text may overflow")
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = heading or title
            frame = slide.placeholders[1].text_frame
            frame.clear()
            for line_index, line in enumerate(body_lines or [""]):
                paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.size = Pt(24 if len(body_lines) < 8 else 18)
            if image_data is not None:
                slide.shapes.add_picture(io.BytesIO(image_data), Inches(9.1), Inches(2.0), width=Inches(3.5))
            if notes_lines:
                slide.notes_slide.notes_text_frame.text = "\n".join(notes_lines)
        output = io.BytesIO()
        presentation.save(output)
        return output.getvalue()

    def _decode_data_image(self, value: str) -> bytes:
        from PIL import Image
        try:
            encoded = value.split(",", 1)[1]
            raw = base64.b64decode(encoded, validate=True)
            if len(raw) > 5 * 1024 * 1024:
                raise ArtifactPathError("Embedded image exceeds the 5 MB limit")
            with Image.open(io.BytesIO(raw)) as image:
                image.load()
                if image.width * image.height > 20_000_000:
                    raise ArtifactPathError("Embedded image exceeds the 20 megapixel limit")
                safe = image.convert("RGBA" if image.mode in {"RGBA", "LA"} else "RGB")
                output = io.BytesIO()
                safe.save(output, format="PNG", optimize=True)
                return output.getvalue()
        except ArtifactPathError:
            raise
        except Exception as exc:
            raise ArtifactPathError("Embedded image is invalid") from exc

    def _office_summary(self, path: Path, suffix: str) -> str:
        if path.stat().st_size > MAX_ARTIFACT_BYTES:
            raise ArtifactPathError("Office artifact exceeds the 5 MB preview limit")
        if suffix == ".docx":
            from docx import Document
            document = Document(str(path))
            return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text)[:MAX_PREVIEW_CHARS]
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            lines: list[str] = []
            for worksheet in workbook.worksheets[:20]:
                lines.append(f"## {worksheet.title}")
                for row in worksheet.iter_rows(max_row=200, max_col=50, values_only=True):
                    lines.append("\t".join("" if value is None else str(value) for value in row))
            workbook.close()
            return "\n".join(lines)[:MAX_PREVIEW_CHARS]
        from pptx import Presentation
        presentation = Presentation(str(path))
        return "\n\n".join(
            f"## Slide {index + 1}\n" + "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text)
            for index, slide in enumerate(presentation.slides[:100])
        )[:MAX_PREVIEW_CHARS]

    async def history(
        self,
        session_id: str,
        path: str,
        limit: int | None = None,
    ) -> list[dict[str, Any]]:
        repo_path = Path(await self.prepare_checkout(session_id))
        normalized = self.normalize_path(path)
        rows = [
            item
            for item in self._read_history(repo_path)
            if item.get("path") == normalized
        ]
        rows.reverse()
        return rows if limit is None else rows[: max(0, limit)]

    async def all_history(self, session_id: str) -> list[dict[str, Any]]:
        repo_path = Path(await self.prepare_checkout(session_id))
        rows = self._read_history(repo_path)
        rows.reverse()
        return rows

    def _read_history(self, repo_path: Path) -> list[dict[str, Any]]:
        history_file = repo_path / HISTORY_FILENAME
        if not history_file.exists():
            return []
        rows: list[dict[str, Any]] = []
        with _lock_for(history_file):
            for line in history_file.read_text(encoding="utf-8").splitlines():
                try:
                    item = json.loads(line)
                except (json.JSONDecodeError, TypeError):
                    continue
                if isinstance(item, dict):
                    rows.append(item)
        return rows

    def _append_history(
        self, repo_path: Path, path: str, message: str, author: str, *,
        source_model: str | None = None, status: str = "ready", parent_version: int | None = None,
    ) -> dict[str, Any]:
        full_path = repo_path / Path(*path.split("/"))
        file_hash = self._hash_file(full_path)
        raw = full_path.read_bytes() if full_path.exists() else b""
        if len(raw) > MAX_HISTORY_SNAPSHOT_BYTES:
            snapshot = ""
            content_encoding = "external"
        else:
            try:
                snapshot = raw.decode("utf-8")
                content_encoding = "utf-8"
            except UnicodeDecodeError:
                snapshot = base64.b64encode(raw).decode("ascii")
                content_encoding = "base64"

        history_file = repo_path / HISTORY_FILENAME
        with _lock_for(history_file):
            existing = self._read_history(repo_path)
            versions = [item for item in existing if item.get("path") == path]
            previous_hash = versions[-1].get("hash") if versions else None
            created_at = _utc_now()
            item = {
                "id": str(uuid4()),
                "version": len(versions) + 1,
                "hash": file_hash,
                "short_hash": file_hash[:7],
                "previous_hash": previous_hash,
                "message": str(message or "Artifact update"),
                "commit_message": str(message or "Artifact update"),
                "author": str(author or "unknown"),
                "path": path,
                "kind": self._kind_for(path),
                "mime_type": self._mime_for(path),
                "status": status,
                "source_model": source_model,
                "parent_version": parent_version,
                "size_bytes": len(raw),
                "content_encoding": content_encoding,
                "content": snapshot,
                "created_at": created_at,
            }
            with history_file.open("a", encoding="utf-8") as handle:
                handle.write(json.dumps(item, ensure_ascii=False) + "\n")
        return item

    def _hash_file(self, path: Path) -> str:
        if not path.exists():
            return "0" * 64
        return hashlib.sha256(path.read_bytes()).hexdigest()

    def _kind_for(self, path: str) -> str:
        suffix = Path(path).suffix.lower()
        if suffix in {".md", ".markdown", ".txt", ".html", ".htm"}:
            return "doc"
        if suffix in {".json", ".csv", ".tsv", ".xlsx"}:
            return "dataset"
        if suffix == ".prompt":
            return "prompt"
        if suffix in {".png", ".jpg", ".jpeg", ".gif", ".pdf", ".docx", ".pptx"}:
            return "binary"
        return "code"

    def _mime_for(self, path: str) -> str:
        overrides = {
            ".md": "text/markdown",
            ".markdown": "text/markdown",
            ".svg": "image/svg+xml",
            # Several Linux mime databases reserve .ts for Qt Linguist files.
            # AetherChat artifacts use .ts/.tsx for TypeScript source code.
            ".ts": "text/typescript",
            ".tsx": "text/typescript",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }
        return overrides.get(Path(path).suffix.lower()) or mimetypes.guess_type(path)[0] or "application/octet-stream"


class MergeProposalStore:
    """Small durable store shared by HTTP routes, tools and exports."""

    def __init__(self, base_path: str | os.PathLike[str] | None = None):
        self.data_root = _data_root(base_path)
        self.storage_path = self.data_root / "merge_proposals.json"
        self.storage_path.parent.mkdir(parents=True, exist_ok=True)

    def list(self, session_id: str) -> list[dict[str, Any]]:
        session = str(session_id)
        with _lock_for(self.storage_path):
            return [
                deepcopy(item)
                for item in self._load_unlocked()
                if str(item.get("session_id")) == session
            ]

    def get(self, proposal_id: str) -> dict[str, Any] | None:
        with _lock_for(self.storage_path):
            for item in self._load_unlocked():
                if item.get("id") == proposal_id:
                    return deepcopy(item)
        return None

    def create(
        self,
        *,
        session_id: str,
        artifact_path: str,
        proposer: str,
        rationale: str,
        from_hash: str | None = None,
        to_hash: str | None = None,
        merge_rule: str = "two_votes",
    ) -> dict[str, Any]:
        rule = self.normalize_rule(merge_rule)
        now = _utc_now()
        thresholds: dict[str, dict[str, Any]] = {
            "two_votes": {"approvals": 2, "rejections": 1},
            "facilitator": {"facilitator_votes": 1},
            "manual_confirmation": {"manual_confirmations": 1},
        }
        proposal = {
            "id": str(uuid4()),
            "session_id": str(session_id),
            "artifact_path": artifact_path,
            "proposer": proposer,
            "from_hash": from_hash,
            "to_hash": to_hash,
            "rationale": rationale,
            "status": "open",
            "merge_rule": rule,
            "merge_threshold": thresholds[rule],
            "votes_needed": 2 if rule == "two_votes" else 1,
            "votes": {},
            "vote_metadata": {},
            "created_at": now,
            "updated_at": now,
            "resolved_at": None,
        }
        with _lock_for(self.storage_path):
            proposals = self._load_unlocked()
            proposals.append(proposal)
            self._save_unlocked(proposals)
        return deepcopy(proposal)

    def delete_session(self, session_id: str) -> int:
        session = str(session_id)
        with _lock_for(self.storage_path):
            proposals = self._load_unlocked()
            retained = [
                item for item in proposals
                if str(item.get("session_id")) != session
            ]
            removed = len(proposals) - len(retained)
            if removed:
                self._save_unlocked(retained)
        return removed

    def vote(
        self,
        proposal_id: str,
        *,
        voter: str,
        approve: bool,
        role: str | None = None,
        manual_confirmation: bool = False,
    ) -> tuple[dict[str, Any] | None, bool]:
        with _lock_for(self.storage_path):
            proposals = self._load_unlocked()
            for proposal in proposals:
                if proposal.get("id") != proposal_id:
                    continue

                # Terminal proposals are immutable: no late vote can alter the
                # decision or even rewrite the audit trail.
                if proposal.get("status") in {"merged", "rejected"}:
                    return deepcopy(proposal), False

                now = _utc_now()
                voter_name = str(voter or "unknown")
                proposal.setdefault("votes", {})[voter_name] = bool(approve)
                proposal.setdefault("vote_metadata", {})[voter_name] = {
                    "role": role,
                    "manual_confirmation": bool(manual_confirmation),
                    "created_at": now,
                }

                rule = self.normalize_rule(proposal.get("merge_rule", "two_votes"))
                normalized_role = str(role or "").strip().lower()
                normalized_voter = voter_name.strip().lower()
                is_facilitator = normalized_role in {"facilitator", "host", "moderator"} or normalized_voter in {
                    "facilitator",
                    "host",
                    "moderator",
                } or normalized_voter.endswith(":facilitator")
                is_manual_actor = bool(manual_confirmation) or normalized_role in {
                    "user",
                    "owner",
                    "human",
                    "manual",
                } or normalized_voter in {"user", "owner", "human", "manual"} or normalized_voter.startswith("user:")

                if rule == "two_votes":
                    approvals = sum(value is True for value in proposal["votes"].values())
                    rejections = sum(value is False for value in proposal["votes"].values())
                    if approvals >= 2:
                        proposal["status"] = "merged"
                    elif rejections >= 1:
                        proposal["status"] = "rejected"
                elif rule == "facilitator" and is_facilitator or rule == "manual_confirmation" and is_manual_actor:
                    proposal["status"] = "merged" if approve else "rejected"

                proposal["updated_at"] = now
                if proposal["status"] in {"merged", "rejected"}:
                    proposal["resolved_at"] = now
                    proposal["resolved_by"] = voter_name
                self._save_unlocked(proposals)
                return deepcopy(proposal), True
        return None, False

    @staticmethod
    def normalize_rule(rule: str) -> str:
        aliases = {
            "two_approvals": "two_votes",
            "two-votes": "two_votes",
            "2_votes": "two_votes",
            "两票通过": "two_votes",
            "host": "facilitator",
            "moderator": "facilitator",
            "主持人通过": "facilitator",
            "manual": "manual_confirmation",
            "human_confirmation": "manual_confirmation",
            "人工确认": "manual_confirmation",
        }
        normalized = aliases.get(str(rule or "two_votes").strip().lower(), str(rule or "two_votes").strip().lower())
        if normalized not in MERGE_RULES:
            raise ValueError(f"Unsupported merge_rule: {rule}")
        return normalized

    def _load_unlocked(self) -> builtins.list[dict[str, Any]]:
        if not self.storage_path.exists():
            return []
        try:
            payload = json.loads(self.storage_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError, TypeError):
            return []
        if isinstance(payload, dict):
            payload = payload.get("proposals", [])
        return [item for item in payload if isinstance(item, dict)] if isinstance(payload, list) else []

    def _save_unlocked(self, proposals: builtins.list[dict[str, Any]]) -> None:
        payload = {"version": 1, "proposals": proposals}
        temp_path = self.storage_path.with_suffix(self.storage_path.suffix + ".tmp")
        temp_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        temp_path.replace(self.storage_path)
